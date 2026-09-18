"""
app.exporters.consolidado_exporter

Exportador especializado para la Matriz 1: Consolidado de Actividades (Fase 9).
Consume directamente los modelos EstadisticaActividad producidos por el Motor Estadístico (Fase 7).
"""

from pathlib import Path
from typing import Any, Dict, List, Optional
import openpyxl

from app.audit.audit_logger import get_logger
from app.exporters.base_exporter import BaseExcelExporter, calcular_sha256
from app.exporters.exceptions import (
    ExportacionBloqueadaError,
    ModificacionPlantillaBaseError,
)
from app.exporters.models import (
    EstatusPlantilla,
    EstrategiaExportacion,
    ModoExportacionCatalogo,
    PoliticaCatalogo,
    ResultadoMatrizExportada,
)
from app.statistics.models import EstadisticaActividad
from app.templates_analysis.models import EsquemaPlantilla, MappingManifesto, TipoMapeoColumna

logger = get_logger(__name__)


class ConsolidadoExporter(BaseExcelExporter):
    """
    Genera el libro Excel de Matriz 1: Consolidado de Actividades.
    Escribe una fila agregada por cada actividad institucional respetando las fórmulas
    de sumatoria inferior o la Tabla1 de la plantilla oficial.
    """

    def exportar(
        self,
        actividades: List[EstadisticaActividad],
        carpeta_destino: Path,
    ) -> ResultadoMatrizExportada:
        """
        Exporta las estadísticas de actividades a la Matriz 1.

        Args:
            actividades: Lista de EstadisticaActividad calculadas en Fase 7.
            carpeta_destino: Directorio donde se guardará el archivo generado.

        Returns:
            ResultadoMatrizExportada con la auditoría de celdas, fórmulas y hashes.
        """
        self.validar_precondiciones()

        ruta_plantilla = Path(self.esquema.ruta_archivo)
        sha_antes = calcular_sha256(ruta_plantilla)

        carpeta_destino.mkdir(parents=True, exist_ok=True)
        ruta_salida = carpeta_destino / self.nombre_archivo_salida

        # Cargar libro en modo lectura/escritura preservando fórmulas
        wb = openpyxl.load_workbook(str(ruta_plantilla), data_only=False)
        ws = wb[self.esquema.hoja_inspeccionada]

        es_oficial = self.esquema.es_oficial or (ws.max_column >= 40)
        fila_inicio = 2 if es_oficial else (self.esquema.fila_inicio_datos or 5)
        fila_actual = fila_inicio
        filas_escritas = 0

        columnas_sin_fuente: List[str] = []
        columnas_no_homologadas: List[str] = []
        advertencias: List[str] = []

        # Celda prototipo de estilo (fila_inicio)
        fila_prototipo = fila_inicio

        for idx_act, act in enumerate(actividades):
            fila_actual = fila_inicio + idx_act
            filas_escritas += 1

            for item in self.manifesto.items:
                c_idx = item.columna_indice
                val = self._extraer_valor_actividad(
                    act, item, idx_act + 1, fila_actual, columnas_no_homologadas, columnas_sin_fuente, es_oficial
                )

                if self._autorizar_escritura(ws, fila_actual, c_idx, item):
                    celda_destino = ws.cell(row=fila_actual, column=c_idx)
                    celda_proto = ws.cell(row=fila_prototipo, column=c_idx)
                    self._clonar_estilo_celda(celda_proto, celda_destino)
                    celda_destino.value = val

        # Sincronizar Tabla1 nativa en plantillas oficiales
        if es_oficial:
            self._sincronizar_tabla_excel(ws, "Tabla1", fila_actual if filas_escritas > 0 else 2)

        wb.save(str(ruta_salida))
        wb.close()

        # Validar integridad SHA-256 de la plantilla original
        sha_despues = calcular_sha256(ruta_plantilla)
        if sha_antes != sha_despues:
            raise ModificacionPlantillaBaseError(
                f"La plantilla base {ruta_plantilla.name} fue alterada durante la exportación de {self.id_matriz}."
            )

        # Verificar fórmulas en el archivo generado
        formulas_verificadas = self.verificar_formulas_post_escritura(
            ruta_salida,
            fila_inicio_datos=fila_inicio if filas_escritas > 0 else None,
            fila_fin_datos=fila_actual if filas_escritas > 0 else None,
        )

        estatus = (
            EstatusPlantilla.OFICIAL_REAL
            if es_oficial
            else EstatusPlantilla.GOLDEN_TEST_FIXTURE_TEST_ONLY
        )

        logger.info(
            f"Matriz 1 exportada exitosamente: {filas_escritas} fila(s), "
            f"{len(formulas_verificadas)} fórmula(s) preservada(s) en {ruta_salida.name}"
        )

        return ResultadoMatrizExportada(
            id_matriz=self.id_matriz,
            nombre_matriz=self.nombre_matriz,
            nombre_archivo=self.nombre_archivo_salida,
            ruta_salida=str(ruta_salida),
            plantilla_utilizada=ruta_plantilla.name,
            estatus_plantilla=estatus,
            sha256_plantilla_antes=sha_antes,
            sha256_plantilla_despues=sha_despues,
            plantilla_inalterada=True,
            total_filas_escritas=filas_escritas,
            total_registros_historicos_preservados=0,
            estrategia_exportacion=EstrategiaExportacion.GENERACION_DESDE_PLANTILLA,
            fila_inicio=fila_inicio if filas_escritas > 0 else None,
            fila_fin=fila_actual if filas_escritas > 0 else None,
            celdas_con_formula_preservadas=formulas_verificadas,
            columnas_sin_fuente_respetadas=sorted(list(set(columnas_sin_fuente))),
            columnas_catalogo_no_homologadas=sorted(list(set(columnas_no_homologadas))),
            advertencias=advertencias,
            exitosa=True,
        )

    def _extraer_valor_actividad(
        self,
        act: EstadisticaActividad,
        item,
        secuencial: int,
        fila_actual: int,
        cols_no_homologadas: List[str],
        cols_sin_fuente: List[str],
        es_oficial: bool = False,
    ) -> Any:
        """Extrae el valor del modelo para la columna del Consolidado."""
        if item.tipo_mapeo == TipoMapeoColumna.PROTEGIDO:
            return None

        if item.tipo_mapeo == TipoMapeoColumna.SIN_FUENTE:
            cols_sin_fuente.append(item.encabezado_original)
            return None

        attr = (item.campo_ssot_fuente or "").lower()
        enc_lower = item.encabezado_original.lower().strip()
        c_idx = item.columna_indice

        # 1. Correlativo / No
        if "numero_secuencial" in attr or enc_lower in ("no", "no.", "num", "n°"):
            if es_oficial:
                return 1 if fila_actual == 2 else f"=+A{fila_actual - 1}+1"
            return secuencial

        # 2. Plantilla Oficial Real (42 columnas)
        if es_oficial:
            # Metadatos territoriales y de actividad (cols 2..24)
            if c_idx == 2 or enc_lower == "sede":
                return act.sede
            if c_idx == 3 or enc_lower == "dep_sede":
                return act.departamento
            if c_idx == 4 or enc_lower == "mun_sede":
                return act.municipio
            if c_idx == 8 or enc_lower == "ámbito" or enc_lower == "ambito":
                return act.eje_linea
            if c_idx == 9 or enc_lower == "actividad":
                return act.nombre_actividad
            if c_idx == 13 or enc_lower == "evento":
                return act.tipo_evento
            if c_idx == 15 or enc_lower == "fecha_evento":
                return act.fecha
            if c_idx == 20 or enc_lower == "región_atendida" or enc_lower == "region_atendida":
                return act.departamento
            if c_idx == 21 or enc_lower == "departamento":
                return act.departamento
            if c_idx == 22 or enc_lower == "municipio":
                return act.municipio

            # Desglose cuantitativo agregado (cols 25..42)
            if c_idx == 25:  # Total, Atención, M
                return act.desglose_sexo_global.masculino
            if c_idx == 26:  # Total, Atención, F
                return act.desglose_sexo_global.femenino
            if c_idx == 27:  # Total, estud_M_grado
                return act.estudiantes.conteo_sexo.masculino
            if c_idx == 28:  # Total, estud_F_grado
                return act.estudiantes.conteo_sexo.femenino
            if c_idx == 29:  # Total, estud_M_posgrado
                return 0
            if c_idx == 30:  # Total, estud_F_posgrado
                return 0
            if c_idx == 31:  # Total, docente_M
                return act.docentes.conteo_sexo.masculino
            if c_idx == 32:  # Total, docente_F
                return act.docentes.conteo_sexo.femenino
            if c_idx == 33:  # Total, Administrativos_M
                return act.administrativos_no_docentes.conteo_sexo.masculino
            if c_idx == 34:  # Total, Administrativos_F
                return act.administrativos_no_docentes.conteo_sexo.femenino
            if c_idx == 35:  # Total, partic_M_inst_pública
                return act.colaboradores.conteo_sexo.masculino
            if c_idx == 36:  # Total, partic_F_inst_pública
                return act.colaboradores.conteo_sexo.femenino
            if c_idx in (37, 38, 39, 40):  # Privada y ONG
                return 0
            if c_idx == 41:  # Total, Protagonistas_M
                return act.beneficiados.conteo_sexo.masculino
            if c_idx == 42:  # Total, Protagonistas_F
                return act.beneficiados.conteo_sexo.femenino

        # 3. Mapeo para Fixtures Sintéticos (17 columnas)
        if "nombre_actividad" in attr or enc_lower == "nombre de la actividad":
            return act.nombre_actividad
        if "sede" in attr or enc_lower == "sede":
            return act.sede
        if "municipio" in attr or enc_lower == "municipio":
            politica = self.politicas_catalogos.get("municipio")
            if politica:
                if politica.modo_exportacion == ModoExportacionCatalogo.BLOQUEAR and politica.requerido:
                    raise ExportacionBloqueadaError("Exportación bloqueada: catálogo de municipios requerido y no homologado.")
                cols_no_homologadas.append(item.encabezado_original)
            return act.municipio
        if "departamento" in attr or enc_lower == "departamento":
            return act.departamento
        if "fecha" in attr or enc_lower == "fecha":
            return act.fecha
        if "eje" in attr or "linea" in attr or enc_lower == "eje":
            return act.eje_linea
        if "tipo_evento" in attr or enc_lower == "tipo de evento":
            return act.tipo_evento
        if "total_asistencias" in attr or enc_lower == "total asistencias":
            return act.total_asistencias
        if "personas_unicas" in attr or enc_lower == "personas unicas":
            return act.total_personas_unicas
        if "femenino" in attr or enc_lower == "femenino":
            return act.desglose_sexo_global.femenino
        if "masculino" in attr or enc_lower == "masculino":
            return act.desglose_sexo_global.masculino
        if "estudiantes" in attr or enc_lower == "estudiantes":
            return act.estudiantes.conteo_sexo.total
        if "docentes" in attr or enc_lower == "docentes":
            return act.docentes.conteo_sexo.total
        if "administrativos" in attr or "no_docentes" in attr or enc_lower == "administrativos":
            return act.administrativos_no_docentes.conteo_sexo.total
        if "beneficiados" in attr or enc_lower == "beneficiados":
            return act.beneficiados.conteo_sexo.total
        if "colaboradores" in attr or enc_lower == "colaboradores":
            return act.colaboradores.conteo_sexo.total

        return None

