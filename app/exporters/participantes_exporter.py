"""
app.exporters.participantes_exporter

Exportador especializado para las matrices oficiales de detalle (Fase 9):
- Matriz 2: Estudiantes
- Matriz 3: Académicos y Administrativos
- Matriz 4: Colaboradores
- Matriz 5: Protagonistas Beneficiados

Soporta casos con N > 0 y N = 0 (plantilla limpia sin inventar filas ficticias).
"""

import copy
import datetime
from pathlib import Path
import re
from typing import Any, Dict, List, Optional
import openpyxl

from app.audit.audit_logger import get_logger
from app.exporters.base_exporter import BaseExcelExporter, calcular_sha256
from app.exporters.exceptions import (
    ExportacionBloqueadaError,
    ModificacionPlantillaBaseError,
)
from app.exporters.models import (
    EstatusPlantilla,
    EstrategiaExportacion,
    ModoExportacionCatalogo,
    PoliticaCatalogo,
    ResultadoMatrizExportada,
)
from app.routing.enums import SubtipoInstitucional
from app.routing.models import EnrutamientoRegistro
from app.templates_analysis.models import EsquemaPlantilla, MappingManifesto, TipoMapeoColumna

logger = get_logger(__name__)


class ParticipantesExporter(BaseExcelExporter):
    """
    Genera libros Excel para matrices de detalle de participantes a partir de EnrutamientoRegistro.
    Soporta matrices oficiales reales (57, 58, 57, 53 columnas) y fixtures sintéticos.
    Implementa estrategia APPEND para M5 (conservación de 32 históricos) y GENERACION_DESDE_PLANTILLA para M2..M4.
    """

    def exportar(
        self,
        registros: List[EnrutamientoRegistro],
        carpeta_destino: Path,
    ) -> ResultadoMatrizExportada:
        """
        Exporta los registros asignados a la matriz de destino.

        Args:
            registros: Lista de EnrutamientoRegistro procedentes del routing de Fase 6.
            carpeta_destino: Directorio donde se guardará el archivo generado.

        Returns:
            ResultadoMatrizExportada con auditoría detallada.
        """
        self.validar_precondiciones()

        ruta_plantilla = Path(self.esquema.ruta_archivo)
        sha_antes = calcular_sha256(ruta_plantilla)

        carpeta_destino.mkdir(parents=True, exist_ok=True)
        ruta_salida = carpeta_destino / self.nombre_archivo_salida

        wb = openpyxl.load_workbook(str(ruta_plantilla), data_only=False)
        ws = wb[self.esquema.hoja_inspeccionada]

        es_oficial = self.esquema.es_oficial or (ws.max_column >= 50)
        mid_norm = self.id_matriz.lower()
        es_m5 = "matriz_5" in mid_norm or "m5" in mid_norm or "benefici" in mid_norm

        if es_m5 and es_oficial:
            estrategia = EstrategiaExportacion.APPEND
        else:
            estrategia = EstrategiaExportacion.GENERACION_DESDE_PLANTILLA

        columnas_sin_fuente: List[str] = []
        columnas_no_homologadas: List[str] = []
        advertencias: List[str] = []

        total_historicos_preservados = 0
        filas_escritas = 0
        fila_inicio_escrita = None
        fila_fin_escrita = None

        if estrategia == EstrategiaExportacion.APPEND:
            # Estrategia APPEND para M5 oficial: detectar registros históricos existentes en filas 2..33
            ultima_fila_existente = 1
            for r in range(2, ws.max_row + 1):
                # Verificar presencia de datos en columnas clave de M5
                val_ced = ws.cell(row=r, column=36).value
                val_nom = ws.cell(row=r, column=34).value
                val_sede = ws.cell(row=r, column=2).value
                if val_ced or val_nom or (val_sede and not str(val_sede).startswith("=")):
                    ultima_fila_existente = r

            total_historicos_preservados = max(0, ultima_fila_existente - 1)
            fila_prototipo = 2

            if len(registros) == 0:
                logger.info(
                    f"Matriz {self.id_matriz} (APPEND): 0 nuevos registros. Se preservan intactos los {total_historicos_preservados} registros históricos."
                )
                self._sincronizar_tabla_excel(ws, "Tabla1", ultima_fila=ultima_fila_existente)
            else:
                fila_inicio_escrita = ultima_fila_existente + 1
                for idx_reg, reg in enumerate(registros):
                    fila_actual = fila_inicio_escrita + idx_reg
                    filas_escritas += 1
                    fila_fin_escrita = fila_actual

                    self._escribir_fila_oficial_participante(
                        ws=ws,
                        reg=reg,
                        fila_actual=fila_actual,
                        idx_reg=idx_reg,
                        fila_prototipo=fila_prototipo,
                        cols_no_homologadas=columnas_no_homologadas,
                        cols_sin_fuente=columnas_sin_fuente,
                    )

                self._sincronizar_tabla_excel(ws, "Tabla1", ultima_fila=fila_fin_escrita)

        else:
            # Estrategia GENERACION_DESDE_PLANTILLA (M2, M3, M4 o Fixtures)
            if es_oficial:
                fila_inicio = 2
                fila_prototipo = 2
                if len(registros) == 0:
                    logger.info(f"Matriz {self.id_matriz} cuenta con 0 participantes. Se genera plantilla base limpia.")
                    self._sincronizar_tabla_excel(ws, "Tabla1", ultima_fila=2)
                else:
                    fila_inicio_escrita = fila_inicio
                    for idx_reg, reg in enumerate(registros):
                        fila_actual = fila_inicio + idx_reg
                        filas_escritas += 1
                        fila_fin_escrita = fila_actual

                        self._escribir_fila_oficial_participante(
                            ws=ws,
                            reg=reg,
                            fila_actual=fila_actual,
                            idx_reg=idx_reg,
                            fila_prototipo=fila_prototipo,
                            cols_no_homologadas=columnas_no_homologadas,
                            cols_sin_fuente=columnas_sin_fuente,
                        )

                    self._sincronizar_tabla_excel(ws, "Tabla1", ultima_fila=fila_fin_escrita)

            else:
                # Modo Fixture sintético (retrocompatibilidad para tests existentes)
                fila_inicio = self.esquema.fila_inicio_datos or 5
                fila_prototipo = fila_inicio
                if len(registros) == 0:
                    logger.info(f"Matriz {self.id_matriz} cuenta con 0 participantes. Se genera plantilla base limpia.")
                else:
                    fila_inicio_escrita = fila_inicio
                    for idx_reg, reg in enumerate(registros):
                        fila_actual = fila_inicio + idx_reg
                        filas_escritas += 1
                        fila_fin_escrita = fila_actual

                        for item in self.manifesto.items:
                            c_idx = item.columna_indice
                            val = self._extraer_valor_participante(
                                reg, item, idx_reg + 1, columnas_no_homologadas, columnas_sin_fuente
                            )

                            if self._autorizar_escritura(ws, fila_actual, c_idx, item):
                                celda_destino = ws.cell(row=fila_actual, column=c_idx)
                                celda_proto = ws.cell(row=fila_prototipo, column=c_idx)
                                self._clonar_estilo_celda(celda_proto, celda_destino)
                                celda_destino.value = val

        wb.save(str(ruta_salida))
        wb.close()

        # Validar integridad SHA-256 de la plantilla base original
        sha_despues = calcular_sha256(ruta_plantilla)
        if sha_antes != sha_despues:
            raise ModificacionPlantillaBaseError(
                f"La plantilla base {ruta_plantilla.name} fue alterada durante la exportación de {self.id_matriz}."
            )

        # Verificar fórmulas en el archivo resultante
        formulas_verificadas = self.verificar_formulas_post_escritura(
            ruta_salida,
            fila_inicio_datos=fila_inicio_escrita,
            fila_fin_datos=fila_fin_escrita,
        )

        estatus = (
            EstatusPlantilla.OFICIAL_REAL
            if es_oficial
            else EstatusPlantilla.GOLDEN_TEST_FIXTURE_TEST_ONLY
        )

        logger.info(
            f"Matriz {self.id_matriz} exportada: {filas_escritas} registro(s) escrito(s), "
            f"{total_historicos_preservados} histórico(s) preservado(s) en {ruta_salida.name}"
        )

        return ResultadoMatrizExportada(
            id_matriz=self.id_matriz,
            nombre_matriz=self.nombre_matriz,
            nombre_archivo=self.nombre_archivo_salida,
            ruta_salida=str(ruta_salida),
            plantilla_utilizada=ruta_plantilla.name,
            estatus_plantilla=estatus,
            sha256_plantilla_antes=sha_antes,
            sha256_plantilla_despues=sha_despues,
            plantilla_inalterada=True,
            total_filas_escritas=filas_escritas,
            total_registros_historicos_preservados=total_historicos_preservados,
            estrategia_exportacion=estrategia,
            fila_inicio=fila_inicio_escrita,
            fila_fin=fila_fin_escrita,
            celdas_con_formula_preservadas=formulas_verificadas,
            columnas_sin_fuente_respetadas=sorted(list(set(columnas_sin_fuente))),
            columnas_catalogo_no_homologadas=sorted(list(set(columnas_no_homologadas))),
            advertencias=advertencias,
            exitosa=True,
        )

    def _escribir_fila_oficial_participante(
        self,
        ws,
        reg: EnrutamientoRegistro,
        fila_actual: int,
        idx_reg: int,
        fila_prototipo: int,
        cols_no_homologadas: List[str],
        cols_sin_fuente: List[str],
    ) -> None:
        """Escribe una fila completa conforme al Mapping Manifesto real de Fase 9.2."""
        mid = self.id_matriz.lower()
        is_m2 = "matriz_2" in mid or "estudiant" in mid
        is_m3 = "matriz_3" in mid or "academic" in mid
        is_m4 = "matriz_4" in mid or "colaborad" in mid
        is_m5 = "matriz_5" in mid or "benefici" in mid

        max_cols = 57 if (is_m2 or is_m4) else (58 if is_m3 else (53 if is_m5 else ws.max_column))

        # 1. Correlativo en columna 1
        celda_a = ws.cell(row=fila_actual, column=1)
        celda_proto_a = ws.cell(row=fila_prototipo, column=1)
        self._clonar_estilo_celda(celda_proto_a, celda_a)
        if fila_actual == 2:
            celda_a.value = 1
        else:
            celda_a.value = f"=+A{fila_actual - 1}+1"

        # 2. Metadatos de Actividad (Cols 2..32)
        act = reg.actividad
        mapa_actividad = {
            2: act.sede,
            3: act.departamento,
            4: act.municipio_evento,
            5: act.programa,
            6: None,
            7: None,
            8: act.ambito,
            9: act.nombre_actividad_oficial or act.nombre_actividad_original,
            10: None,
            11: act.tipo_evento,
            12: None,
            13: act.tipo_evento,
            14: getattr(reg.persona, "area", None) or act.eje_linea_estrategica,
            15: act.eje_linea_estrategica,
            16: None,
            17: None,
            18: None,
            19: None,
            20: None,
            21: None,
            22: None,
            23: None,
            24: None,
            25: None,
            26: None,
            27: None,
            28: None,
            29: None,
            30: None,
            31: None,
            32: None,
        }
        for c_idx in range(2, 33):
            val = mapa_actividad.get(c_idx)
            celda = ws.cell(row=fila_actual, column=c_idx)
            celda_proto = ws.cell(row=fila_prototipo, column=c_idx)
            self._clonar_estilo_celda(celda_proto, celda)
            if val is not None:
                celda.value = val

        # 3. Columnas intermedias y demografía según matriz
        p = reg.persona
        sexo_val = getattr(p.sexo_normalizado, "value", p.sexo_normalizado) if p.sexo_normalizado else None
        nombres_completos = p.nombre_completo or f"{p.nombres or ''} {p.apellidos or ''}".strip()
        carrera_val = p.carrera_oficial or p.carrera_normalizada or p.carrera_original
        if carrera_val:
            cols_no_homologadas.append("Nombre_carrera")

        if is_m2:
            # M2 (Estudiantes - 57 columnas)
            col_edad = 40
            col_fecha_nac = 39
            mapa_especifico = {
                33: nombres_completos,
                34: p.numero_unico,
                35: p.cedula,
                36: None,
                37: None,
                38: sexo_val,
                39: p.fecha_nacimiento,
                40: '=DATEDIF(Tabla1[[#This Row],[Fecha_nacimiento]], TODAY(), "Y")',
                41: "1. Natural de Nicaragua" if p.cedula else None,
                42: "Nicaragua" if p.cedula else None,
                43: p.etnia,
                44: p.departamento_persona,
                45: p.municipio_persona,
                46: None,
                47: None,
                48: None,
                49: p.discapacidad,
                50: None,
                51: "Grado",
                52: carrera_val,
                53: getattr(p, "area", None),
                54: None,
                55: None,
                56: None,
                57: None,
            }
        elif is_m3:
            # M3 (Académicos y Administrativos - 58 columnas)
            col_edad = 42
            col_fecha_nac = 41
            subtipo = "DOCENTE" if reg.subtipo_institucional == SubtipoInstitucional.ACADEMICO else "ADMINISTRATIVO"
            mapa_especifico = {
                33: None,  # Mentoría
                34: subtipo,  # Tipo_protagonistas
                35: nombres_completos,
                36: p.numero_unico,
                37: p.cedula,
                38: None,
                39: None,
                40: sexo_val,
                41: p.fecha_nacimiento,
                42: '=DATEDIF(Tabla1[[#This Row],[Fecha_nacimiento]], TODAY(), "Y")',
                43: "1. Natural de Nicaragua" if p.cedula else None,
                44: "Nicaragua" if p.cedula else None,
                45: p.etnia,
                46: p.departamento_persona,
                47: p.municipio_persona,
                48: None,
                49: None,
                50: None,
                51: p.discapacidad,
                52: None,
                53: p.carrera_original,
                54: getattr(p, "area", None),
                55: p.departamento_persona,
                56: getattr(p, "cargo", None),
                57: getattr(p, "cargo", None),
                58: None,
            }
        elif is_m4:
            # M4 (Colaboradores - 57 columnas)
            col_edad = 42
            col_fecha_nac = 41
            mapa_especifico = {
                33: None,  # Mentoría
                34: "COLABORADOR",  # Tipo_protagonistas
                35: nombres_completos,
                36: p.numero_unico,
                37: p.cedula,
                38: None,
                39: None,
                40: sexo_val,
                41: p.fecha_nacimiento,
                42: '=DATEDIF(Tabla1[[#This Row],[Fecha_nacimiento]], TODAY(), "Y")',
                43: "1. Natural de Nicaragua" if p.cedula else None,
                44: "Nicaragua" if p.cedula else None,
                45: p.etnia,
                46: p.departamento_persona,
                47: p.municipio_persona,
                48: None,
                49: None,
                50: None,
                51: p.discapacidad,
                52: None,
                53: None,
                54: getattr(p, "cargo", None),
                55: getattr(p, "cargo", None),
                56: getattr(p, "institucion_procedencia", None),
                57: getattr(p, "institucion_procedencia", None),
            }
        elif is_m5:
            # M5 (Beneficiarios - 53 columnas)
            col_edad = 41
            col_fecha_nac = 40
            mapa_especifico = {
                33: None,  # Mentoría
                34: nombres_completos,
                35: p.numero_unico,
                36: p.cedula,
                37: None,
                38: None,
                39: sexo_val,
                40: p.fecha_nacimiento,
                41: '=DATEDIF(Tabla1[[#This Row],[Fecha_nacimiento]], TODAY(), "Y")',
                42: "1. Natural de Nicaragua" if p.cedula else None,
                43: "Nicaragua" if p.cedula else None,
                44: p.etnia,
                45: p.departamento_persona,
                46: p.municipio_persona,
                47: None,
                48: None,
                49: None,
                50: p.discapacidad,
                51: None,
                52: None,
                53: getattr(p, "beneficio", None) or "Capacitación / Acompañamiento",
            }
        else:
            mapa_especifico = {}
            col_edad = None
            col_fecha_nac = None

        # Escribir columnas específicas
        for c_idx in range(33, max_cols + 1):
            celda = ws.cell(row=fila_actual, column=c_idx)
            celda_proto = ws.cell(row=fila_prototipo, column=c_idx)
            self._clonar_estilo_celda(celda_proto, celda)

            val = mapa_especifico.get(c_idx)
            if c_idx == col_edad:
                # FÓRMULA ESTRUCTURADA DE EDAD - NUNCA ESCRIBIR ENTERO ESTÁTICO
                celda.value = '=DATEDIF(Tabla1[[#This Row],[Fecha_nacimiento]], TODAY(), "Y")'
            elif c_idx == col_fecha_nac:
                # Escribir fecha formateada
                self._escribir_fecha_celda(celda, val)
            elif val is not None:
                celda.value = val

    def _escribir_fecha_celda(self, celda, val: Any) -> None:
        """Escribe una fecha en celda Excel con formato nativo de fecha."""
        if val is None:
            return
        if isinstance(val, (datetime.date, datetime.datetime)):
            celda.value = val
            celda.number_format = "yyyy-mm-dd"
            return
        val_str = str(val).strip()
        if not val_str:
            return
        try:
            if "-" in val_str and len(val_str) == 10:
                parts = val_str.split("-")
                d = datetime.date(int(parts[0]), int(parts[1]), int(parts[2]))
                celda.value = d
                celda.number_format = "yyyy-mm-dd"
                return
            if "/" in val_str and len(val_str) == 10:
                parts = val_str.split("/")
                d = datetime.date(int(parts[2]), int(parts[1]), int(parts[0]))
                celda.value = d
                celda.number_format = "yyyy-mm-dd"
                return
        except Exception:
            pass
        celda.value = val_str

    def _extraer_valor_participante(
        self,
        reg: EnrutamientoRegistro,
        item,
        secuencial: int,
        cols_no_homologadas: List[str],
        cols_sin_fuente: List[str],
    ) -> Any:
        """Extrae el valor del registro para fixtures sintéticos (retrocompatibilidad)."""
        if item.tipo_mapeo == TipoMapeoColumna.PROTEGIDO:
            return None

        if item.tipo_mapeo == TipoMapeoColumna.SIN_FUENTE:
            cols_sin_fuente.append(item.encabezado_original)
            return None

        attr = (item.campo_ssot_fuente or "").lower()
        enc_lower = item.encabezado_original.lower()

        # 1. Secuencial correlativo
        if "numero_secuencial" in attr or enc_lower in ("no.", "no", "n°", "num"):
            return secuencial

        # 2. Cédula
        if "cedula" in attr:
            return reg.persona.cedula

        # 3. Nombres y Apellidos
        if "nombre_completo" in attr or ("nombre" in enc_lower and "apellido" in enc_lower):
            nombres = reg.persona.nombres or ""
            apellidos = reg.persona.apellidos or ""
            completo = f"{nombres} {apellidos}".strip()
            return completo or reg.persona.apellidos or None

        if "nombres" in attr or enc_lower == "nombres":
            return reg.persona.nombres

        if "apellidos" in attr or enc_lower == "apellidos":
            return reg.persona.apellidos

        # 4. Sexo
        if "sexo" in attr:
            sn = reg.persona.sexo_normalizado
            return getattr(sn, "value", sn) if sn else None

        # 5. Edad
        if "edad" in attr:
            return reg.persona.edad

        # 6. Carrera / Programa Académico
        if "carrera" in attr or "carrera" in enc_lower:
            politica = self.politicas_catalogos.get("carrera")
            if politica:
                if politica.modo_exportacion == ModoExportacionCatalogo.BLOQUEAR and politica.requerido:
                    raise ExportacionBloqueadaError(
                        f"Exportación bloqueada: catálogo de carreras requerido y no homologado para '{reg.persona.apellidos}'."
                    )
                cols_no_homologadas.append(item.encabezado_original)
            return reg.persona.carrera_original or reg.persona.carrera_normalizada

        # 7. Rol o Subtipo Institucional (Matriz 3)
        if "rol" in attr or "tipo" in attr or "rol" in enc_lower:
            if reg.subtipo_institucional == SubtipoInstitucional.ACADEMICO:
                return "DOCENTE"
            elif reg.subtipo_institucional == SubtipoInstitucional.ADMINISTRATIVO:
                return "ADMINISTRATIVO"
            cat = reg.participacion.categoria_participacion
            return getattr(cat, "value", cat) if cat else None

        # 8. Área o Dependencia
        if "area" in attr or "facultad" in attr or "area" in enc_lower:
            area_val = getattr(reg.persona, "area", None)
            if not area_val and reg.participacion.observaciones:
                m = re.search(r"Área:\s*([^|;,\n]+)", reg.participacion.observaciones, re.IGNORECASE)
                if m:
                    area_val = m.group(1).strip()
            return area_val

        # 9. Cargo
        if "cargo" in attr or "cargo" in enc_lower:
            return getattr(reg.persona, "cargo", None)

        # 10. Institución de Procedencia (Matriz 4)
        if "institucion" in attr or "organizacion" in attr or "institucion" in enc_lower:
            return getattr(reg.persona, "institucion_procedencia", None)

        # 11. Comunidad (Matriz 5)
        if "comunidad" in attr or "comunidad" in enc_lower:
            return getattr(reg.persona, "comunidad", None)

        # 12. Actividad y Metadatos de Evento
        if "actividad" in attr or "actividad" in enc_lower:
            return reg.actividad.nombre_actividad_oficial or reg.actividad.nombre_actividad

        if "sede" in attr or "sede" in enc_lower:
            return reg.actividad.sede

        if "fecha" in attr or "fecha" in enc_lower:
            return reg.actividad.fecha_evento

        # 13. Observaciones / Control
        if "observaciones" in attr or "observaciones" in enc_lower:
            if reg.estado_operativo == "EN_REVISION":
                obs: List[str] = []
                for h in reg.hallazgos_asociados:
                    obs.append(h.mensaje)
                return "; ".join(obs) if obs else "EN_REVISION"
            return None

        return None
