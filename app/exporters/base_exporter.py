"""
app.exporters.base_exporter

Clase base no destructiva para exportadores Excel con openpyxl (Fase 9).
Asegura la preservación de fórmulas, formato, compuertas de seguridad en ZonaEscribible
y cálculo criptográfico de hashes SHA-256.
"""

import copy
import hashlib
from pathlib import Path
import re
from typing import Any, Dict, List, Optional, Tuple
import openpyxl
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from app.audit.audit_logger import get_logger
from app.exporters.exceptions import (
    ExportacionBloqueadaError,
    ModificacionPlantillaBaseError,
    ViolacionZonaEscribibleError,
)
from app.exporters.models import (
    EstatusPlantilla,
    FormulaVerificada,
    ModoExportacionCatalogo,
    PoliticaCatalogo,
    ResultadoMatrizExportada,
    TipoSinFuente,
)
from app.templates_analysis.models import (
    EsquemaPlantilla,
    ItemMappingManifesto,
    MappingManifesto,
    TipoMapeoColumna,
)

logger = get_logger(__name__)


def calcular_sha256(ruta_archivo: Path) -> str:
    """Calcula el hash criptográfico SHA-256 de un archivo en disco."""
    if not ruta_archivo.exists():
        return ""
    sha = hashlib.sha256()
    with open(ruta_archivo, "rb") as f:
        while chunk := f.read(65536):
            sha.update(chunk)
    return sha.hexdigest()


class BaseExcelExporter:
    """
    Exportador base no destructivo.
    Trabaja sobre una copia de la plantilla base en staging, preservando
    la plantilla original intacta y verificando fórmulas y estilos.
    """

    def __init__(
        self,
        id_matriz: str,
        nombre_matriz: str,
        nombre_archivo_salida: str,
        esquema: EsquemaPlantilla,
        manifesto: MappingManifesto,
        politicas_catalogos: Optional[Dict[str, PoliticaCatalogo]] = None,
        configuracion: Optional[Dict[str, Any]] = None,
    ):
        self.id_matriz = id_matriz
        self.nombre_matriz = nombre_matriz
        self.nombre_archivo_salida = nombre_archivo_salida
        self.esquema = esquema
        self.manifesto = manifesto
        self.politicas_catalogos = politicas_catalogos or {}
        self.configuracion = configuracion or {}

    def validar_precondiciones(self) -> None:
        """Valida que la plantilla base exista y que no haya columnas obligatorias sin fuente."""
        ruta_plantilla = Path(self.esquema.ruta_archivo)
        if not ruta_plantilla.exists():
            raise FileNotFoundError(f"Plantilla no encontrada: {ruta_plantilla}")

        if not self.esquema.fila_inicio_datos:
            raise ExportacionBloqueadaError(
                f"No se pudo determinar la fila segura de inicio de datos para {self.id_matriz}."
            )

        # Verificar si hay columnas marcadas como SIN_FUENTE que sean obligatorias
        for item in self.manifesto.items:
            if item.tipo_mapeo == TipoMapeoColumna.SIN_FUENTE:
                # Consultar configuración o convención de obligatoriedad
                regla_sin_fuente = getattr(item, "clasificacion_sin_fuente", TipoSinFuente.SIN_FUENTE_OPCIONAL)
                if regla_sin_fuente == TipoSinFuente.SIN_FUENTE_OBLIGATORIO:
                    raise ExportacionBloqueadaError(
                        f"Columna obligatoria '{item.encabezado_original}' en {self.id_matriz} carece de fuente de datos."
                    )

    def _autorizar_escritura(self, ws: Worksheet, fila: int, col: int, item: ItemMappingManifesto) -> bool:
        """
        Aplica las 4 compuertas de seguridad para autorizar la escritura en una celda:
        1. Polígono de ZonaEscribible.
        2. Celda original no contiene fórmula.
        3. Columna con tipo escribible (no PROTEGIDO ni NO_APLICA).
        4. No sobrescribe filas de encabezado.
        """
        # Compuerta 0: Fila debe ser >= fila_inicio_datos
        if fila < (self.esquema.fila_inicio_datos or 1):
            return False

        # Compuerta 1: Polígono de ZonaEscribible
        if self.esquema.zonas_escribibles:
            dentro_zona = False
            for zona in self.esquema.zonas_escribibles:
                col_ok = zona.columna_inicio <= col <= zona.columna_fin
                fila_ok = fila >= zona.fila_inicio
                if zona.fila_fin_estimada:
                    fila_ok = fila_ok and (fila <= zona.fila_fin_estimada)
                if col_ok and fila_ok:
                    dentro_zona = True
                    break
            if not dentro_zona and self.esquema.es_oficial:
                # En plantillas oficiales con tablas estructuradas, las filas se expanden dinámicamente
                dentro_zona = (
                    1 <= col <= (self.esquema.max_columnas or 100)
                    and fila >= (self.esquema.fila_inicio_datos or 2)
                )
            if not dentro_zona:
                logger.warning(f"Intento de escritura fuera de ZonaEscribible en {self.id_matriz}: ({fila}, {col})")
                return False

        # Compuerta 2: Preservar fórmulas preexistentes
        celda_actual = ws.cell(row=fila, column=col)
        val_str = str(celda_actual.value).strip() if celda_actual.value is not None else ""
        if val_str.startswith("="):
            logger.info(f"Preservando fórmula en celda ({fila}, {col}): {val_str}")
            return False

        # Compuerta 3: Tipo de mapeo escribible
        if item.tipo_mapeo in (TipoMapeoColumna.PROTEGIDO, TipoMapeoColumna.NO_APLICA):
            return False

        return True

    def _sincronizar_tabla_excel(self, ws: Worksheet, nombre_tabla: str = "Tabla1", ultima_fila: int = 2) -> None:
        """
        Sincroniza dinámicamente el rango de la tabla nativa de Excel (Tabla1) y su autoFilter.
        Garantiza que la tabla siga siendo un objeto OpenXML nativo válido.
        No encoge la tabla por debajo de su tamaño original para evitar desvincular celdas.
        """
        if hasattr(ws, "tables") and nombre_tabla in ws.tables:
            tabla = ws.tables[nombre_tabla]
            ref_actual = str(tabla.ref)
            m = re.search(r":([A-Za-z]+)(\d+)", ref_actual)
            orig_max_col = m.group(1) if m else get_column_letter(ws.max_column or len(tabla.tableColumns))
            orig_max_row = int(m.group(2)) if m else 22

            max_c = ws.max_column or len(tabla.tableColumns)
            col_letra = get_column_letter(max_c)

            nueva_fin = max(ultima_fila, orig_max_row, 2)
            nueva_ref = f"A1:{col_letra}{nueva_fin}"
            tabla.ref = nueva_ref
            if getattr(tabla, "autoFilter", None):
                tabla.autoFilter.ref = nueva_ref
            logger.info(f"Tabla {nombre_tabla} en {self.id_matriz} sincronizada a {nueva_ref}")

    def _clonar_estilo_celda(self, celda_origen, celda_destino) -> None:
        """Clona de forma segura el formato visual de una celda prototipo."""
        if celda_origen.has_style:
            celda_destino.font = copy.copy(celda_origen.font)
            celda_destino.border = copy.copy(celda_origen.border)
            celda_destino.fill = copy.copy(celda_origen.fill)
            celda_destino.number_format = celda_origen.number_format
            celda_destino.protection = copy.copy(celda_origen.protection)
            celda_destino.alignment = copy.copy(celda_origen.alignment)

    def verificar_formulas_post_escritura(
        self,
        ruta_archivo: Path,
        fila_inicio_datos: Optional[int] = None,
        fila_fin_datos: Optional[int] = None,
    ) -> List[FormulaVerificada]:
        """
        Re-inspecciona el archivo generado en disco para constatar que todas las
        fórmulas originales de la plantilla siguen presentes e intactas en sus coordenadas.
        Diferencia explícitamente PRESERVACIÓN ESTRUCTURAL de RECALCULACIÓN DEL RESULTADO.
        """
        wb_check = openpyxl.load_workbook(str(ruta_archivo), data_only=False)
        ws_check = wb_check[self.esquema.hoja_inspeccionada]
        formulas_verificadas: List[FormulaVerificada] = []

        for cp in self.esquema.celdas_protegidas:
            if cp.tipo_proteccion == "FORMULA" and cp.formula:
                celda_check = ws_check[cp.coordenada]
                val_encontrado = str(celda_check.value).strip() if celda_check.value is not None else ""
                preservada = (val_encontrado == cp.formula)
                no_reemplazada = preservada

                # Extraer rango referenciado (ej. 'I5:I9' en '=SUM(I5:I9)')
                match_rango = re.search(r"([A-Za-z]+[0-9]+:[A-Za-z]+[0-9]+)", cp.formula)
                rango_ref = match_rango.group(1) if match_rango else None

                # Extraer fila de la coordenada de fórmula (ej. 10 de 'I10')
                match_fila = re.search(r"\d+", cp.coordenada)
                fila_formula = int(match_fila.group(0)) if match_fila else None

                # Validar que las filas de datos escritas no invadan la celda de fórmula
                no_invade = True
                if fila_formula is not None and fila_fin_datos is not None:
                    es_totales = "SUM(" in cp.formula.upper()
                    if es_totales and fila_fin_datos >= fila_formula:
                        no_invade = False

                if not preservada:
                    logger.error(
                        f"VIOLACIÓN DE FÓRMULA en {self.id_matriz} ({cp.coordenada}): "
                        f"Esperada='{cp.formula}', Encontrada='{val_encontrado}'"
                    )

                formulas_verificadas.append(
                    FormulaVerificada(
                        coordenada=cp.coordenada,
                        formula=cp.formula,
                        formula_original=cp.formula,
                        formula_posterior=val_encontrado,
                        rango_referenciado=rango_ref,
                        preservada=preservada,
                        tipo_preservacion="PRESERVACION_ESTRUCTURAL",
                        recalculacion_nota=(
                            "openpyxl preserva la definición sintáctica de la fórmula pero carece de motor "
                            "de evaluación en tiempo de ejecución. La recalculación matemática es realizada por "
                            "Microsoft Excel al abrir el archivo."
                        ),
                        no_reemplazada=no_reemplazada,
                        filas_no_invaden_rango=no_invade,
                        valor_evaluado_en_celda=val_encontrado or None,
                    )
                )

        wb_check.close()
        return formulas_verificadas
