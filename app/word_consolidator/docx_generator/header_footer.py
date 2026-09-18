"""
app.word_consolidator.docx_generator.header_footer

Gestor de encabezados y pies de página institucionales para python-docx.
Maneja primera página diferenciada (portada limpia) y campos dinámicos
de numeración de página nativos de Word.
"""

from typing import Optional

import docx
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn
from docx.shared import Pt

from app.word_consolidator.docx_generator.style_manager import ColorPalette, StyleManager


class HeaderFooterManager:
    """
    Configura encabezados y pies de página institucionales.
    """

    @staticmethod
    def _add_page_number_field(run: docx.text.run.Run) -> None:
        """Inserta el campo dinámico OpenXML 'PAGE' de Microsoft Word."""
        fldChar1 = parse_xml(r'<w:fldChar %s w:fldCharType="begin"/>' % nsdecls('w'))
        instrText = parse_xml(r'<w:instrText %s xml:space="preserve"> PAGE </w:instrText>' % nsdecls('w'))
        fldChar2 = parse_xml(r'<w:fldChar %s w:fldCharType="separate"/>' % nsdecls('w'))
        fldChar3 = parse_xml(r'<w:fldChar %s w:fldCharType="end"/>' % nsdecls('w'))
        run._r.append(fldChar1)
        run._r.append(instrText)
        run._r.append(fldChar2)
        run._r.append(fldChar3)

    @classmethod
    def setup_headers_and_footers(
        cls,
        doc: docx.Document,
        institucion: Optional[str] = None,
        periodo_texto: Optional[str] = None,
        area_texto: Optional[str] = None,
    ) -> None:
        """
        Configura los encabezados y pies de página en las secciones del documento.
        La portada (primera página) no muestra encabezado ni pie.
        """
        if not doc.sections:
            return

        for section in doc.sections:
            # 1. Portada sin encabezado ni pie
            section.different_first_page_header_footer = True
            
            # 2. Encabezado de páginas normales
            header = section.header
            if not header.paragraphs:
                header_p = header.add_paragraph()
            else:
                header_p = header.paragraphs[0]
            
            header_p.text = ""  # limpiar
            header_p.paragraph_format.space_before = Pt(0)
            header_p.paragraph_format.space_after = Pt(2)
            
            # Tabla de 1 fila, 2 columnas para alineación izquierda / derecha perfecta
            header_tbl = header.add_table(rows=1, cols=2, width=section.page_width - section.left_margin - section.right_margin)
            header_tbl.autofit = False
            
            # Columna izquierda: Institución
            c_left = header_tbl.cell(0, 0)
            p_left = c_left.paragraphs[0]
            p_left.paragraph_format.space_after = Pt(0)
            r_left = p_left.add_run(institucion or "INFORME CONSOLIDADO INSTITUCIONAL")
            r_left.font.name = StyleManager.FONT_NAME
            r_left.font.size = StyleManager.SIZE_HEADER_FOOTER
            r_left.font.color.rgb = StyleManager.hex_to_rgb(ColorPalette.ACCENT_SLATE)
            r_left.bold = True
            
            # Columna derecha: Período
            c_right = header_tbl.cell(0, 1)
            p_right = c_right.paragraphs[0]
            p_right.alignment = WD_ALIGN_PARAGRAPH.RIGHT
            p_right.paragraph_format.space_after = Pt(0)
            r_right = p_right.add_run(periodo_texto or "")
            r_right.font.name = StyleManager.FONT_NAME
            r_right.font.size = StyleManager.SIZE_HEADER_FOOTER
            r_right.font.color.rgb = StyleManager.hex_to_rgb(ColorPalette.ACCENT_SLATE)
            
            # Borde inferior sutil al encabezado
            tblPr = header_tbl._tbl.tblPr
            borders = parse_xml(
                f'<w:tblBorders {nsdecls("w")}>'
                f'<w:bottom w:val="single" w:sz="6" w:space="0" w:color="{ColorPalette.BORDER_STRONG}"/>'
                f'<w:top w:val="none"/>'
                f'<w:left w:val="none"/>'
                f'<w:right w:val="none"/>'
                f'<w:insideH w:val="none"/>'
                f'<w:insideV w:val="none"/>'
                f'</w:tblBorders>'
            )
            tblPr.append(borders)
            
            # 3. Pie de página de páginas normales
            footer = section.footer
            if not footer.paragraphs:
                footer_p = footer.add_paragraph()
            else:
                footer_p = footer.paragraphs[0]
            
            footer_p.text = ""
            footer_p.paragraph_format.space_before = Pt(2)
            footer_p.paragraph_format.space_after = Pt(0)
            
            footer_tbl = footer.add_table(rows=1, cols=2, width=section.page_width - section.left_margin - section.right_margin)
            footer_tbl.autofit = False
            
            # Borde superior sutil al pie
            f_tblPr = footer_tbl._tbl.tblPr
            f_borders = parse_xml(
                f'<w:tblBorders {nsdecls("w")}>'
                f'<w:top w:val="single" w:sz="6" w:space="0" w:color="{ColorPalette.BORDER_STRONG}"/>'
                f'<w:bottom w:val="none"/>'
                f'<w:left w:val="none"/>'
                f'<w:right w:val="none"/>'
                f'<w:insideH w:val="none"/>'
                f'<w:insideV w:val="none"/>'
                f'</w:tblBorders>'
            )
            f_tblPr.append(f_borders)
            
            # Columna izquierda: Área responsable o nota institucional
            fc_left = footer_tbl.cell(0, 0)
            fp_left = fc_left.paragraphs[0]
            fp_left.paragraph_format.space_after = Pt(0)
            fr_left = fp_left.add_run(area_texto or "Sistema de Automatización Estadístico de Asistencias")
            fr_left.font.name = StyleManager.FONT_NAME
            fr_left.font.size = StyleManager.SIZE_HEADER_FOOTER
            fr_left.font.color.rgb = StyleManager.hex_to_rgb(ColorPalette.TEXT_MUTED)
            
            # Columna derecha: Numeración de página dinámica
            fc_right = footer_tbl.cell(0, 1)
            fp_right = fc_right.paragraphs[0]
            fp_right.alignment = WD_ALIGN_PARAGRAPH.RIGHT
            fp_right.paragraph_format.space_after = Pt(0)
            fr_label = fp_right.add_run("Página ")
            fr_label.font.name = StyleManager.FONT_NAME
            fr_label.font.size = StyleManager.SIZE_HEADER_FOOTER
            fr_label.font.color.rgb = StyleManager.hex_to_rgb(ColorPalette.TEXT_MUTED)
            
            fr_num = fp_right.add_run()
            fr_num.font.name = StyleManager.FONT_NAME
            fr_num.font.size = StyleManager.SIZE_HEADER_FOOTER
            fr_num.font.color.rgb = StyleManager.hex_to_rgb(ColorPalette.TEXT_MUTED)
            cls._add_page_number_field(fr_num)
