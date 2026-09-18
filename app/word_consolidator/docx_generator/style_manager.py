"""
app.word_consolidator.docx_generator.style_manager

Administrador centralizado de estilos tipográficos, cromáticos y espaciado
para la generación de documentos Word institucionales.
Garantiza un diseño institucional, sobrio, elegante y legible.
"""

from dataclasses import dataclass
from typing import Optional, Tuple

import docx
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls
from docx.shared import Inches, Mm, Pt, RGBColor


@dataclass(frozen=True)
class ColorPalette:
    """Paleta de colores sobria e institucional para el informe."""
    # Azules institucionales
    PRIMARY_NAVY: str = "002D62"       # Azul institucional universitario
    SECONDARY_BLUE: str = "1E3A8A"     # Azul secundario para subtítulos
    ACCENT_SLATE: str = "4A5568"       # Pizarra para metadatos y etiquetas
    
    # Texto
    TEXT_DARK: str = "1A202C"          # Texto principal (alto contraste)
    TEXT_MUTED: str = "718096"         # Texto secundario / notas al pie
    TEXT_WHITE: str = "FFFFFF"         # Texto sobre fondos oscuros
    
    # Fondos y bandas
    BG_LIGHT_GRAY: str = "F7FAFC"      # Fondo suave para filas alternas
    BG_CARD: str = "F8FAFC"            # Fondo de tarjetas y recuadros
    BG_HEADER: str = "002D62"          # Fondo de encabezados de tabla primarios
    BG_SUBHEADER: str = "EDF2F7"       # Fondo de encabezados secundarios
    
    # Bordes
    BORDER_LIGHT: str = "E2E8F0"       # Bordes sutiles
    BORDER_STRONG: str = "CBD5E0"      # Bordes principales de tabla
    
    # Estados de auditoría (sobrios, no estridentes)
    CONCORDANTE_BG: str = "DEF7EC"     # Verde suave institucional
    CONCORDANTE_TEXT: str = "03543F"   # Verde oscuro legible
    REVISION_BG: str = "FDE8E8"        # Rojo/ámbar suave de auditoría
    REVISION_TEXT: str = "9B1C1C"      # Rojo oscuro legible


class StyleManager:
    """
    Administrador centralizado de estilos del documento.
    Aplica tipografía consistente, espaciados equilibrados y reglas de diseño institucional.
    """

    FONT_NAME = "Arial"
    
    # Tamaños de fuente
    SIZE_COVER_TITLE = Pt(20)
    SIZE_COVER_SUBTITLE = Pt(12)
    SIZE_H1 = Pt(14)
    SIZE_H2 = Pt(11.5)
    SIZE_H3 = Pt(10)
    SIZE_BODY = Pt(9.5)
    SIZE_TABLE_HEADER = Pt(8.5)
    SIZE_TABLE_BODY = Pt(8.0)
    SIZE_NOTE = Pt(7.5)
    SIZE_HEADER_FOOTER = Pt(8.0)

    @classmethod
    def hex_to_rgb(cls, hex_str: str) -> RGBColor:
        """Convierte cadena hexadecimal en objeto RGBColor de docx."""
        hex_clean = hex_str.lstrip("#")
        return RGBColor(
            int(hex_clean[0:2], 16),
            int(hex_clean[2:4], 16),
            int(hex_clean[4:6], 16),
        )

    @classmethod
    def apply_document_styles(cls, doc: docx.Document) -> None:
        """Configura los estilos base en el documento Word."""
        styles = doc.styles

        # Estilo Normal
        if "Normal" in styles:
            normal = styles["Normal"]
            normal.font.name = cls.FONT_NAME
            normal.font.size = cls.SIZE_BODY
            normal.font.color.rgb = cls.hex_to_rgb(ColorPalette.TEXT_DARK)
            normal.paragraph_format.line_spacing = 1.15
            normal.paragraph_format.space_after = Pt(4)
            normal.paragraph_format.space_before = Pt(0)

    @classmethod
    def add_heading_1(cls, doc: docx.Document, text: str, space_before: int = 12) -> docx.text.paragraph.Paragraph:
        """Añade un encabezado de Sección principal (H1) con estilo institucional."""
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(space_before)
        p.paragraph_format.space_after = Pt(6)
        p.paragraph_format.keep_with_next = True
        
        run = p.add_run(text)
        run.font.name = cls.FONT_NAME
        run.font.size = cls.SIZE_H1
        run.font.bold = True
        run.font.color.rgb = cls.hex_to_rgb(ColorPalette.PRIMARY_NAVY)
        return p

    @classmethod
    def add_heading_2(cls, doc: docx.Document, text: str, space_before: int = 8) -> docx.text.paragraph.Paragraph:
        """Añade un encabezado de Subsección (H2)."""
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(space_before)
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.keep_with_next = True
        
        run = p.add_run(text)
        run.font.name = cls.FONT_NAME
        run.font.size = cls.SIZE_H2
        run.font.bold = True
        run.font.color.rgb = cls.hex_to_rgb(ColorPalette.SECONDARY_BLUE)
        return p

    @classmethod
    def add_heading_3(cls, doc: docx.Document, text: str, space_before: int = 6) -> docx.text.paragraph.Paragraph:
        """Añade un encabezado de Nivel 3 (H3) para fichas o detalles."""
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(space_before)
        p.paragraph_format.space_after = Pt(2)
        p.paragraph_format.keep_with_next = True
        
        run = p.add_run(text)
        run.font.name = cls.FONT_NAME
        run.font.size = cls.SIZE_H3
        run.font.bold = True
        run.font.color.rgb = cls.hex_to_rgb(ColorPalette.ACCENT_SLATE)
        return p

    @classmethod
    def add_paragraph(
        cls,
        doc: docx.Document,
        text: str = "",
        bold: bool = False,
        italic: bool = False,
        color_hex: Optional[str] = None,
        align: WD_ALIGN_PARAGRAPH = WD_ALIGN_PARAGRAPH.LEFT,
        space_after: int = 4,
    ) -> docx.text.paragraph.Paragraph:
        """Añade un párrafo estándar con formato tipográfico uniforme."""
        p = doc.add_paragraph()
        p.alignment = align
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(space_after)
        p.paragraph_format.line_spacing = 1.15
        
        if text:
            run = p.add_run(text)
            run.font.name = cls.FONT_NAME
            run.font.size = cls.SIZE_BODY
            run.bold = bold
            run.italic = italic
            if color_hex:
                run.font.color.rgb = cls.hex_to_rgb(color_hex)
            else:
                run.font.color.rgb = cls.hex_to_rgb(ColorPalette.TEXT_DARK)
        return p

    @classmethod
    def add_callout(
        cls,
        doc: docx.Document,
        text: str,
        title: Optional[str] = None,
        bg_hex: str = ColorPalette.BG_SUBHEADER,
        border_hex: str = ColorPalette.PRIMARY_NAVY,
    ) -> None:
        """Crea un recuadro o callout institucional sobrio para notas o alertas."""
        table = doc.add_table(rows=1, cols=1)
        table.autofit = False
        cell = table.cell(0, 0)
        
        # Sombreado de fondo
        shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{bg_hex}"/>')
        cell._tc.get_or_add_tcPr().append(shd)
        
        # Borde izquierdo destacado
        borders = parse_xml(
            f'<w:tcBorders {nsdecls("w")}>'
            f'<w:top w:val="none"/>'
            f'<w:left w:val="single" w:sz="24" w:space="0" w:color="{border_hex}"/>'
            f'<w:bottom w:val="none"/>'
            f'<w:right w:val="none"/>'
            f'</w:tcBorders>'
        )
        cell._tc.get_or_add_tcPr().append(borders)
        
        p = cell.paragraphs[0]
        p.paragraph_format.space_before = Pt(3)
        p.paragraph_format.space_after = Pt(3)
        p.paragraph_format.left_indent = Pt(6)
        
        if title:
            r_title = p.add_run(f"{title}: ")
            r_title.font.name = cls.FONT_NAME
            r_title.font.size = cls.SIZE_BODY
            r_title.bold = True
            r_title.font.color.rgb = cls.hex_to_rgb(border_hex)
            
        r_text = p.add_run(text)
        r_text.font.name = cls.FONT_NAME
        r_text.font.size = cls.SIZE_BODY
        r_text.font.color.rgb = cls.hex_to_rgb(ColorPalette.TEXT_DARK)
        
        # Espacio posterior al callout
        spacer = doc.add_paragraph()
        spacer.paragraph_format.space_before = Pt(0)
        spacer.paragraph_format.space_after = Pt(4)
