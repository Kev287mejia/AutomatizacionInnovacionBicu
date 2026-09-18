"""
app.word_consolidator.docx_generator.table_builder

Constructor reutilizable y de alta fidelidad para tablas institucionales en python-docx.
Maneja:
- Repetición automática de encabezados en saltos de página (w:tblHeader)
- Prevención de corte vertical de filas (w:cantSplit)
- Anchos explícitos en celdas para evitar deformaciones en Word
- Sombreado de celdas (w:shd) y márgenes internos (padding / w:tcMar)
- Insignias de auditoría (Badges) para estados de concordancia
"""

from typing import List, Optional, Sequence, Union

import docx
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn
from docx.shared import Inches, Mm, Pt

from app.word_consolidator.docx_generator.style_manager import ColorPalette, StyleManager
from app.word_consolidator.models import EstadoDiscrepancia


class TableBuilder:
    """
    Constructor y estilista de tablas Word para informes técnicos y ejecutivos.
    """

    @classmethod
    def set_cell_margins(
        cls,
        cell: docx.table._Cell,
        top_dxa: int = 120,    # ~6 pt
        bottom_dxa: int = 120, # ~6 pt
        left_dxa: int = 160,   # ~8 pt
        right_dxa: int = 160,  # ~8 pt
    ) -> None:
        """Configura el relleno interno (padding) de la celda en dxa."""
        tcPr = cell._tc.get_or_add_tcPr()
        tcMar = OxmlElement("w:tcMar")
        for margin_name, val in [
            ("top", top_dxa),
            ("bottom", bottom_dxa),
            ("left", left_dxa),
            ("right", right_dxa),
        ]:
            node = OxmlElement(f"w:{margin_name}")
            node.set(qn("w:w"), str(val))
            node.set(qn("w:type"), "dxa")
            tcMar.append(node)
        tcPr.append(tcMar)

    @classmethod
    def set_cell_background(cls, cell: docx.table._Cell, fill_hex: str) -> None:
        """Aplica color de fondo (shading) a la celda."""
        shading_elm = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex.lstrip("#")}"/>')
        cell._tc.get_or_add_tcPr().append(shading_elm)

    @classmethod
    def repeat_header_on_every_page(cls, row: docx.table._Row) -> None:
        """Asegura que la fila de encabezado se repita al saltar de página en Word."""
        trPr = row._tr.get_or_add_trPr()
        tblHeader = OxmlElement("w:tblHeader")
        trPr.append(tblHeader)

    @classmethod
    def prevent_row_split(cls, row: docx.table._Row) -> None:
        """Impide que una fila se parta entre dos páginas en Word."""
        trPr = row._tr.get_or_add_trPr()
        cantSplit = OxmlElement("w:cantSplit")
        trPr.append(cantSplit)

    @classmethod
    def apply_table_borders(
        cls,
        table: docx.table.Table,
        color: str = ColorPalette.BORDER_STRONG,
        sz: str = "4",
    ) -> None:
        """Aplica bordes horizontales limpios y sutiles estilo editorial moderno."""
        tblPr = table._tbl.tblPr
        borders = parse_xml(
            f'<w:tblBorders {nsdecls("w")}>'
            f'<w:top w:val="single" w:sz="{sz}" w:space="0" w:color="{color}"/>'
            f'<w:bottom w:val="single" w:sz="{sz}" w:space="0" w:color="{color}"/>'
            f'<w:left w:val="none"/>'
            f'<w:right w:val="none"/>'
            f'<w:insideH w:val="single" w:sz="{sz}" w:space="0" w:color="{color}"/>'
            f'<w:insideV w:val="none"/>'
            f'</w:tblBorders>'
        )
        tblPr.append(borders)

    @classmethod
    def create_table(
        cls,
        doc: docx.Document,
        headers: Sequence[str],
        col_widths: Sequence[float],
        alignment: WD_TABLE_ALIGNMENT = WD_TABLE_ALIGNMENT.CENTER,
    ) -> docx.table.Table:
        """
        Crea una tabla con encabezados estilizados y anchos fijos de columna.
        col_widths: secuencia de anchos en pulgadas (Inches) o milímetros (Mm).
        """
        table = doc.add_table(rows=1, cols=len(headers))
        table.alignment = alignment
        table.autofit = False
        cls.apply_table_borders(table)

        # Fila de encabezados
        header_row = table.rows[0]
        cls.repeat_header_on_every_page(header_row)
        cls.prevent_row_split(header_row)

        for i, header_text in enumerate(headers):
            cell = header_row.cells[i]
            cls.set_cell_background(cell, ColorPalette.BG_HEADER)
            cls.set_cell_margins(cell, top_dxa=140, bottom_dxa=140, left_dxa=160, right_dxa=160)
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(0)
            
            run = p.add_run(header_text)
            run.font.name = StyleManager.FONT_NAME
            run.font.size = StyleManager.SIZE_TABLE_HEADER
            run.font.bold = True
            run.font.color.rgb = StyleManager.hex_to_rgb(ColorPalette.TEXT_WHITE)

        # Fijar anchos de columnas
        for i, width in enumerate(col_widths):
            header_row.cells[i].width = width

        return table

    @classmethod
    def add_row(
        cls,
        table: docx.table.Table,
        values: Sequence[str],
        col_widths: Sequence[float],
        alignments: Optional[Sequence[WD_ALIGN_PARAGRAPH]] = None,
        is_even: bool = False,
        is_total_row: bool = False,
        bold: bool = False,
    ) -> docx.table._Row:
        """
        Añade una fila con valores formateados, padding y anchos garantizados.
        """
        row = table.add_row()
        cls.prevent_row_split(row)

        bg_color = (
            ColorPalette.BG_SUBHEADER
            if is_total_row
            else (ColorPalette.BG_LIGHT_GRAY if is_even else "FFFFFF")
        )

        for i, val in enumerate(values):
            cell = row.cells[i]
            cell.width = col_widths[i]
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            
            if bg_color != "FFFFFF":
                cls.set_cell_background(cell, bg_color)
            cls.set_cell_margins(cell, top_dxa=100, bottom_dxa=100, left_dxa=140, right_dxa=140)

            p = cell.paragraphs[0]
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(0)
            
            if alignments and i < len(alignments):
                p.alignment = alignments[i]
            else:
                p.alignment = WD_ALIGN_PARAGRAPH.LEFT

            run = p.add_run(str(val) if val is not None else "—")
            run.font.name = StyleManager.FONT_NAME
            run.font.size = StyleManager.SIZE_TABLE_BODY
            run.bold = bold or is_total_row
            run.font.color.rgb = StyleManager.hex_to_rgb(ColorPalette.TEXT_DARK)

        return row

    @classmethod
    def add_status_cell(
        cls,
        cell: docx.table._Cell,
        estado: Union[EstadoDiscrepancia, str],
        width: Optional[float] = None,
    ) -> None:
        """
        Formatea una celda como insignia institucional (Badge) para el estado de auditoría.
        """
        if width:
            cell.width = width
        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        cls.set_cell_margins(cell, top_dxa=80, bottom_dxa=80, left_dxa=120, right_dxa=120)

        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(0)

        estado_str = estado.value if isinstance(estado, EstadoDiscrepancia) else str(estado)

        if estado_str == EstadoDiscrepancia.CONCORDANTE.value or estado_str == "APTO":
            bg = ColorPalette.CONCORDANTE_BG
            text_color = ColorPalette.CONCORDANTE_TEXT
            label = "CONCORDANTE" if estado_str != "APTO" else "APTO"
        else:
            bg = ColorPalette.REVISION_BG
            text_color = ColorPalette.REVISION_TEXT
            label = "REQUIERE REVISIÓN" if "REVISION" in estado_str else estado_str

        cls.set_cell_background(cell, bg)

        run = p.add_run(label)
        run.font.name = StyleManager.FONT_NAME
        run.font.size = Pt(7.5)
        run.font.bold = True
        run.font.color.rgb = StyleManager.hex_to_rgb(text_color)
