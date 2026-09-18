"""
app.word_consolidator.docx_generator.sections.quality_audit

Renderizador de la Sección 8: Salud de Datos y Auditoría de Discrepancias.
Sección CRÍTICA: Presenta de forma transparente y exhaustiva todas las diferencias
detectadas entre fuentes (DIMENSIÓN | FUENTE A | FUENTE B | DELTA | ESTADO).
El documento Word NO oculta ni maquilla discrepancias bajo ninguna circunstancia.
"""

import docx
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt

from app.word_consolidator.document.models import AuditoriaDocumental
from app.word_consolidator.docx_generator.style_manager import ColorPalette, StyleManager
from app.word_consolidator.docx_generator.table_builder import TableBuilder
from app.word_consolidator.models import EstadoDiscrepancia


def render_quality_audit(
    doc: docx.Document,
    auditoria: AuditoriaDocumental,
) -> None:
    """
    Construye la sección de Salud de Datos y Auditoría de Discrepancias.
    """
    StyleManager.add_heading_1(doc, "8. SALUD DE DATOS Y AUDITORÍA DE DISCREPANCIAS")

    StyleManager.add_paragraph(
        doc,
        "Esta sección expone los hallazgos de calidad, cotejos dimensionales y discrepancias matemáticas "
        "detectadas de forma automatizada por el sistema. El objetivo institucional de este análisis "
        "es garantizar la total transparencia sobre la concordancia entre los reportes agregados y "
        "las listas nominales físicas sin ocultar ninguna divergencia.",
        space_after=8,
    )

    # 1. Tabla de Resumen Cuantitativo de Auditoría
    StyleManager.add_heading_2(doc, "8.1. Balance General de Concordancia Institucional")

    headers_bal = ["Parámetro de Calidad y Auditoría", "Cifra Verificada", "Dictamen Técnico"]
    widths_bal = [Inches(4.2), Inches(2.0), Inches(3.2)]
    table_bal = TableBuilder.create_table(doc, headers_bal, widths_bal)

    invariante_txt = "CONSERVADA (100% verificado)" if auditoria.invariante_verificada else "NO VERIFICADA"

    filas_bal = [
        ("Total de Actividades Auditadas", str(auditoria.total_actividades_evaluadas), "Evaluación completa"),
        ("Actividades con Coincidencia Plena (0 discrepancias)", str(auditoria.actividades_concordantes), "CONCORDANTE"),
        ("Actividades con Discrepancias / Observaciones", str(auditoria.actividades_con_discrepancias), "REQUIERE_REVISION" if auditoria.actividades_con_discrepancias > 0 else "CONCORDANTE"),
        ("Total de Discrepancias Activas que Requieren Revisión", str(auditoria.total_discrepancias_activas), "REQUIERE_REVISION" if auditoria.total_discrepancias_activas > 0 else "CONCORDANTE"),
        ("Dimensiones Institucionales Concordantes", str(auditoria.total_concordancias), "CONCORDANTE"),
        ("Fuentes Nominales Pendientes / Faltantes", str(auditoria.total_fuentes_faltantes), "REQUIERE_REVISION" if auditoria.total_fuentes_faltantes > 0 else "CONCORDANTE"),
        ("Invariante de Conservación de Asistencias", invariante_txt, "CONCORDANTE" if auditoria.invariante_verificada else "REQUIERE_REVISION"),
    ]

    for i, (param, cifra, dictam) in enumerate(filas_bal):
        row = table_bal.add_row()
        TableBuilder.prevent_row_split(row)
        
        # Celda 0
        c0 = row.cells[0]
        c0.width = widths_bal[0]
        TableBuilder.set_cell_margins(c0)
        p0 = c0.paragraphs[0]
        p0.add_run(param).font.name = StyleManager.FONT_NAME
        p0.runs[0].font.size = StyleManager.SIZE_TABLE_BODY
        if i in (1, 2, 3):
            p0.runs[0].bold = True

        # Celda 1
        c1 = row.cells[1]
        c1.width = widths_bal[1]
        TableBuilder.set_cell_margins(c1)
        p1 = c1.paragraphs[0]
        p1.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        p1.add_run(cifra).font.name = StyleManager.FONT_NAME
        p1.runs[0].font.size = StyleManager.SIZE_TABLE_BODY
        p1.runs[0].bold = True

        # Celda 2 (Badge de estado)
        c2 = row.cells[2]
        TableBuilder.add_status_cell(c2, dictam, width=widths_bal[2])

    p_s1 = doc.add_paragraph()
    p_s1.paragraph_format.space_before = Pt(4)
    p_s1.paragraph_format.space_after = Pt(4)

    # 2. Alertas de Auditoría Emitidas
    if auditoria.alertas:
        StyleManager.add_heading_2(doc, "8.2. Alertas y Observaciones Operativas")
        for alerta in auditoria.alertas:
            StyleManager.add_callout(
                doc,
                text=alerta,
                title="ALERTA DE AUDITORÍA",
                bg_hex=ColorPalette.REVISION_BG,
                border_hex=ColorPalette.REVISION_TEXT,
            )

    # 3. TABLA FORMAL COMPLETA DE DISCREPANCIAS
    StyleManager.add_heading_2(doc, "8.3. Tabla Maestra de Discrepancias Identificadas")

    StyleManager.add_paragraph(
        doc,
        "A continuación se presenta el detalle exhaustivo de cada discrepancia identificada, "
        "especificando la actividad de origen, dimensión evaluada, fuentes contrastadas, valores "
        "registrados, diferencia matemática exacta (Delta = Fuente B - Fuente A) y dictamen de auditoría:",
        space_after=6,
    )

    headers_disc = [
        "Actividad",
        "Categoría",
        "Dimensión",
        "Fuente A (Valor)",
        "Fuente B (Valor)",
        "Delta",
        "Estado",
        "Descripción / Observación",
    ]

    widths_disc = [
        Inches(1.8),  # Actividad
        Inches(1.0),  # Categoría
        Inches(1.0),  # Dimensión
        Inches(1.3),  # Fuente A (Val)
        Inches(1.3),  # Fuente B (Val)
        Inches(0.6),  # Delta
        Inches(1.1),  # Estado
        Inches(1.3),  # Descripción
    ]

    table_disc = TableBuilder.create_table(doc, headers_disc, widths_disc)

    if auditoria.tabla_discrepancias:
        for idx, item in enumerate(auditoria.tabla_discrepancias):
            row_d = table_disc.add_row()
            TableBuilder.prevent_row_split(row_d)
            is_even = (idx % 2 == 1)

            # Formatear Fuente A y B con valores
            fa_txt = f"{item.fuente_a} ({item.valor_fuente_a})" if item.valor_fuente_a is not None else item.fuente_a
            fb_txt = f"{item.fuente_b} ({item.valor_fuente_b})" if item.valor_fuente_b is not None else item.fuente_b

            # Formatear Delta
            if item.delta is not None:
                delta_txt = f"+{item.delta}" if item.delta > 0 else str(item.delta)
            else:
                delta_txt = "—"

            vals = [
                item.nombre_actividad,
                item.categoria,
                item.dimension,
                fa_txt,
                fb_txt,
                delta_txt,
            ]

            aligns = [
                WD_ALIGN_PARAGRAPH.LEFT,
                WD_ALIGN_PARAGRAPH.CENTER,
                WD_ALIGN_PARAGRAPH.LEFT,
                WD_ALIGN_PARAGRAPH.LEFT,
                WD_ALIGN_PARAGRAPH.LEFT,
                WD_ALIGN_PARAGRAPH.RIGHT,
            ]

            for c_i, (v, al) in enumerate(zip(vals, aligns)):
                c = row_d.cells[c_i]
                c.width = widths_disc[c_i]
                TableBuilder.set_cell_margins(c)
                if is_even:
                    TableBuilder.set_cell_background(c, ColorPalette.BG_LIGHT_GRAY)
                p = c.paragraphs[0]
                p.alignment = al
                r = p.add_run(v)
                r.font.name = StyleManager.FONT_NAME
                r.font.size = StyleManager.SIZE_TABLE_BODY
                if c_i == 5:
                    r.bold = True

            # Columna 6: Estado (Badge)
            c_est = row_d.cells[6]
            if is_even:
                TableBuilder.set_cell_background(c_est, ColorPalette.BG_LIGHT_GRAY)
            TableBuilder.add_status_cell(c_est, item.estado, width=widths_disc[6])

            # Columna 7: Descripción
            c_desc = row_d.cells[7]
            c_desc.width = widths_disc[7]
            TableBuilder.set_cell_margins(c_desc)
            if is_even:
                TableBuilder.set_cell_background(c_desc, ColorPalette.BG_LIGHT_GRAY)
            p_desc = c_desc.paragraphs[0]
            p_desc.alignment = WD_ALIGN_PARAGRAPH.LEFT
            r_desc = p_desc.add_run(item.descripcion or "—")
            r_desc.font.name = StyleManager.FONT_NAME
            r_desc.font.size = StyleManager.SIZE_TABLE_BODY
    else:
        # Caso sin discrepancias
        row_none = table_disc.add_row()
        TableBuilder.prevent_row_split(row_none)
        c_none = row_none.cells[0]
        c_none.width = widths_disc[0]
        p_none = c_none.paragraphs[0]
        p_none.add_run("No se registraron discrepancias activas en el período evaluado.")
        # Llenar el resto de celdas
        for ci in range(1, 8):
            row_none.cells[ci].width = widths_disc[ci]
            row_none.cells[ci].paragraphs[0].add_run("—")
