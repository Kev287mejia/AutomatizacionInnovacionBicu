"""
app.word_consolidator.docx_generator.sections.territory

Renderizador de la Sección 7: Cobertura Territorial y Sedes Universitarias.
Presenta fielmente la distribución geográfica soportada en las fuentes.
Cumple estrictamente con la regla de NO INVENTAR comunidades ni inferir territorios ausentes.
"""

import docx
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt

from app.word_consolidator.document.models import TerritorioDocumental
from app.word_consolidator.docx_generator.style_manager import StyleManager
from app.word_consolidator.docx_generator.table_builder import TableBuilder


def render_territory(
    doc: docx.Document,
    territorio: TerritorioDocumental,
) -> None:
    """
    Construye la sección de Cobertura Territorial y Sedes.
    """
    StyleManager.add_heading_1(doc, "7. COBERTURA TERRITORIAL Y SEDES UNIVERSITARIAS")

    StyleManager.add_paragraph(
        doc,
        "Esta sección resume la distribución geográfica de las actividades institucionales y su impacto "
        "en las diversas sedes, departamentos y municipios. Los datos reflejan exclusivamente la "
        "información territorial explícitamente reportada en las fuentes procesadas.",
        space_after=8,
    )

    # 1. Resumen por Sedes Universitarias
    if territorio.desglose_sedes:
        StyleManager.add_heading_2(doc, "7.1. Cobertura por Sede o Recinto Universitario")
        headers_sedes = ["Sede Universitaria", "Actividades", "Asistencias Femeninas", "Asistencias Masculinas", "Total Asistencias"]
        widths_sedes = [Inches(3.2), Inches(1.4), Inches(1.6), Inches(1.6), Inches(1.6)]
        table_sedes = TableBuilder.create_table(doc, headers_sedes, widths_sedes)

        for i, s in enumerate(territorio.desglose_sedes):
            TableBuilder.add_row(
                table_sedes,
                values=[s.sede, str(s.total_actividades), str(s.femenino), str(s.masculino), str(s.total_asistencias)],
                col_widths=widths_sedes,
                alignments=[
                    WD_ALIGN_PARAGRAPH.LEFT,
                    WD_ALIGN_PARAGRAPH.RIGHT,
                    WD_ALIGN_PARAGRAPH.RIGHT,
                    WD_ALIGN_PARAGRAPH.RIGHT,
                    WD_ALIGN_PARAGRAPH.RIGHT,
                ],
                is_even=(i % 2 == 1),
            )

        p_s1 = doc.add_paragraph()
        p_s1.paragraph_format.space_before = Pt(4)
        p_s1.paragraph_format.space_after = Pt(4)

    # 2. Resumen por Departamento o Región Autónoma
    if territorio.desglose_departamentos:
        StyleManager.add_heading_2(doc, "7.2. Cobertura Departamental y Regional")
        headers_dept = ["Departamento / Región", "Municipios Atendidos", "Total Actividades", "Total Asistencias"]
        widths_dept = [Inches(2.8), Inches(3.6), Inches(1.5), Inches(1.5)]
        table_dept = TableBuilder.create_table(doc, headers_dept, widths_dept)

        for j, d in enumerate(territorio.desglose_departamentos):
            munis_txt = ", ".join(d.municipios) if d.municipios else "—"
            TableBuilder.add_row(
                table_dept,
                values=[d.departamento, munis_txt, str(d.total_actividades), str(d.total_asistencias)],
                col_widths=widths_dept,
                alignments=[
                    WD_ALIGN_PARAGRAPH.LEFT,
                    WD_ALIGN_PARAGRAPH.LEFT,
                    WD_ALIGN_PARAGRAPH.RIGHT,
                    WD_ALIGN_PARAGRAPH.RIGHT,
                ],
                is_even=(j % 2 == 1),
            )

        p_s2 = doc.add_paragraph()
        p_s2.paragraph_format.space_before = Pt(4)
        p_s2.paragraph_format.space_after = Pt(4)

    # 3. Resumen por Municipio
    if territorio.desglose_municipios:
        StyleManager.add_heading_2(doc, "7.3. Cobertura por Municipio")
        headers_muni = ["Municipio", "Departamento / Región", "Actividades", "Total Asistencias"]
        widths_muni = [Inches(3.2), Inches(3.2), Inches(1.5), Inches(1.5)]
        table_muni = TableBuilder.create_table(doc, headers_muni, widths_muni)

        for k, m in enumerate(territorio.desglose_municipios):
            TableBuilder.add_row(
                table_muni,
                values=[m.municipio, m.departamento or "—", str(m.total_actividades), str(m.total_asistencias)],
                col_widths=widths_muni,
                alignments=[
                    WD_ALIGN_PARAGRAPH.LEFT,
                    WD_ALIGN_PARAGRAPH.LEFT,
                    WD_ALIGN_PARAGRAPH.RIGHT,
                    WD_ALIGN_PARAGRAPH.RIGHT,
                ],
                is_even=(k % 2 == 1),
            )

    # 4. Comunidades explícitamente reportadas (sólo si existen)
    if territorio.comunidades:
        StyleManager.add_heading_2(doc, "7.4. Comunidades Específicas Reportadas", space_before=8)
        comunidades_str = ", ".join(territorio.comunidades)
        StyleManager.add_paragraph(doc, f"Comunidades registradas en las fuentes: {comunidades_str}.")
