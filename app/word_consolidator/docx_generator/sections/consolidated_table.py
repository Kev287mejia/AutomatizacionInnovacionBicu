"""
app.word_consolidator.docx_generator.sections.consolidated_table

Renderizador de la Sección 3: Matriz Consolidada General de Actividades.
Muestra simultáneamente los valores declarados en Matriz 1, los valores nominales
calculados de M2–M5, la diferencia matemática (Delta = Nominal - M1) y el estado de auditoría.
"""

import docx
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt

from app.word_consolidator.document.models import MatrizGeneralDocumental
from app.word_consolidator.docx_generator.style_manager import StyleManager
from app.word_consolidator.docx_generator.table_builder import TableBuilder
from app.word_consolidator.models import EstadoDiscrepancia


def render_consolidated_table(
    doc: docx.Document,
    matriz: MatrizGeneralDocumental,
) -> None:
    """
    Construye la Matriz Consolidada General de Actividades.
    """
    StyleManager.add_heading_1(doc, "3. MATRIZ CONSOLIDADA GENERAL DE ACTIVIDADES")

    StyleManager.add_paragraph(
        doc,
        "La siguiente matriz presenta la consolidación general de las actividades ejecutadas en el período. "
        "Permite contrastar los valores declarados formalmente en Matriz 1 (M1) frente a las asistencias "
        "reales contabilizadas nominalmente celda por celda a partir de las listas de participantes (M2–M5).",
        space_after=8,
    )

    headers = [
        "#",
        "Nombre de la Actividad",
        "Sede",
        "Territorio",
        "Fecha",
        "Eje Estratégico",
        "Total M1",
        "Total Nom.",
        "Delta",
        "Estado",
    ]

    col_widths = [
        Inches(0.35),  # #
        Inches(2.70),  # Actividad
        Inches(0.85),  # Sede
        Inches(1.10),  # Territorio
        Inches(0.75),  # Fecha
        Inches(1.10),  # Eje
        Inches(0.65),  # M1
        Inches(0.65),  # Nom
        Inches(0.50),  # Delta
        Inches(0.95),  # Estado
    ]

    table = TableBuilder.create_table(doc, headers, col_widths)

    tot_m1_acum = 0
    tot_nom_acum = 0

    for i, fila in enumerate(matriz.filas):
        num_str = str(i + 1)
        nombre_str = fila.nombre_actividad
        sede_str = fila.sede

        # Territorio: Depto / Municipio si existen
        partes_terr = []
        if fila.departamento:
            partes_terr.append(fila.departamento)
        if fila.municipio:
            partes_terr.append(fila.municipio)
        terr_str = " / ".join(partes_terr) if partes_terr else "—"

        # Fecha
        fecha_str = str(fila.fecha) if fila.fecha else "—"

        # Eje estratégico si existe
        eje_str = fila.eje_estrategico or "—"

        # Cifras M1 y Nominal
        m1_val = fila.total_m1
        m1_str = str(m1_val) if m1_val is not None else "—"
        if m1_val is not None:
            tot_m1_acum += m1_val

        nom_val = fila.total_nominal
        nom_str = str(nom_val)
        tot_nom_acum += nom_val

        # Delta exacto: Nominal - M1
        if m1_val is not None:
            delta_val = nom_val - m1_val
            delta_str = f"+{delta_val}" if delta_val > 0 else str(delta_val)
        else:
            delta_str = "—"

        # Fila en tabla
        row = table.add_row()
        TableBuilder.prevent_row_split(row)
        is_even = (i % 2 == 1)

        valores_texto = [num_str, nombre_str, sede_str, terr_str, fecha_str, eje_str, m1_str, nom_str, delta_str]
        alignments = [
            WD_ALIGN_PARAGRAPH.CENTER,  # #
            WD_ALIGN_PARAGRAPH.LEFT,    # Actividad
            WD_ALIGN_PARAGRAPH.LEFT,    # Sede
            WD_ALIGN_PARAGRAPH.LEFT,    # Territorio
            WD_ALIGN_PARAGRAPH.CENTER,  # Fecha
            WD_ALIGN_PARAGRAPH.LEFT,    # Eje
            WD_ALIGN_PARAGRAPH.RIGHT,   # M1
            WD_ALIGN_PARAGRAPH.RIGHT,   # Nom
            WD_ALIGN_PARAGRAPH.RIGHT,   # Delta
        ]

        for col_idx, (val, align) in enumerate(zip(valores_texto, alignments)):
            cell = row.cells[col_idx]
            cell.width = col_widths[col_idx]
            TableBuilder.set_cell_margins(cell, top_dxa=100, bottom_dxa=100, left_dxa=120, right_dxa=120)
            if is_even:
                TableBuilder.set_cell_background(cell, "F8FAFC")
            
            p = cell.paragraphs[0]
            p.alignment = align
            run = p.add_run(val)
            run.font.name = StyleManager.FONT_NAME
            run.font.size = StyleManager.SIZE_TABLE_BODY
            run.font.color.rgb = StyleManager.hex_to_rgb("1A202C")
            if col_idx in (6, 7):
                run.bold = True

        # Celda de Estado (Badge)
        cell_est = row.cells[9]
        cell_est.width = col_widths[9]
        if is_even:
            TableBuilder.set_cell_background(cell_est, "F8FAFC")
        TableBuilder.add_status_cell(cell_est, fila.estado_auditoria, width=col_widths[9])

    # Fila de Totales Generales
    totales = matriz.totales_generales
    m1_tot_str = str(totales.total_m1) if totales.total_m1 is not None else str(tot_m1_acum)
    nom_tot_str = str(totales.total_nominal) if totales.total_nominal > 0 else str(tot_nom_acum)
    
    if totales.total_m1 is not None:
        delta_tot = totales.total_nominal - totales.total_m1
        delta_tot_str = f"+{delta_tot}" if delta_tot > 0 else str(delta_tot)
    else:
        delta_tot_str = "—"

    row_tot = table.add_row()
    TableBuilder.prevent_row_split(row_tot)

    tot_vals = [
        "—",
        f"TOTALES CONSOLIDADOS ({matriz.total_actividades} Actividades)",
        "—",
        "—",
        "—",
        "—",
        m1_tot_str,
        nom_tot_str,
        delta_tot_str,
    ]

    for col_idx, val in enumerate(tot_vals):
        c = row_tot.cells[col_idx]
        c.width = col_widths[col_idx]
        TableBuilder.set_cell_margins(c, top_dxa=120, bottom_dxa=120, left_dxa=120, right_dxa=120)
        TableBuilder.set_cell_background(c, "EDF2F7")
        p = c.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.RIGHT if col_idx in (6, 7, 8) else (WD_ALIGN_PARAGRAPH.CENTER if col_idx == 0 else WD_ALIGN_PARAGRAPH.LEFT)
        run = p.add_run(val)
        run.font.name = StyleManager.FONT_NAME
        run.font.size = StyleManager.SIZE_TABLE_BODY
        run.font.bold = True
        run.font.color.rgb = StyleManager.hex_to_rgb("1A202C")

    # Estado global en la fila de totales
    c_tot_est = row_tot.cells[9]
    c_tot_est.width = col_widths[9]
    TableBuilder.set_cell_background(c_tot_est, "EDF2F7")
    TableBuilder.add_status_cell(c_tot_est, totales.estado_global, width=col_widths[9])
