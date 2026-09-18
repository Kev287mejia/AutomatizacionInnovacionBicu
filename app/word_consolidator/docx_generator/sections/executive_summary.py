"""
app.word_consolidator.docx_generator.sections.executive_summary

Renderizador de la Sección 2: Resumen Ejecutivo Institucional.
Presenta las métricas agregadas globales verificadas, recurrencia, distribución
por sexo y estado general de concordancia, sin narrativa inventada.
"""

import docx
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt

from app.word_consolidator.document.models import ResumenEjecutivoDocumental
from app.word_consolidator.docx_generator.style_manager import ColorPalette, StyleManager
from app.word_consolidator.docx_generator.table_builder import TableBuilder


def render_executive_summary(
    doc: docx.Document,
    resumen: ResumenEjecutivoDocumental,
) -> None:
    """
    Construye la sección de Resumen Ejecutivo.
    """
    StyleManager.add_heading_1(doc, "2. RESUMEN EJECUTIVO")

    StyleManager.add_paragraph(
        doc,
        f"El presente informe consolida las asistencias y participaciones del período {resumen.periodo.etiqueta}. "
        "A continuación se detallan los principales indicadores cuantitativos consolidados a partir de las fuentes oficiales.",
        space_after=8,
    )

    # 1. Tabla de Indicadores Clave de Participación
    StyleManager.add_heading_2(doc, "2.1. Métricas Globales de Asistencia y Participación")

    headers_kpi = ["Indicador Institucional", "Cifra Consolidada", "Descripción Técnica"]
    col_widths_kpi = [Inches(2.8), Inches(1.8), Inches(4.8)]

    table_kpi = TableBuilder.create_table(doc, headers_kpi, col_widths_kpi)

    kpis = [
        (
            "Actividades Consolidadas",
            str(resumen.total_actividades),
            "Total de eventos o actividades ejecutadas en la ventana del período.",
        ),
        (
            "Asistencia Bruta (Participaciones)",
            str(resumen.total_asistencia_bruta),
            "Total acumulado de registros de participación física o nominal.",
        ),
        (
            "Personas Únicas Netas",
            str(resumen.total_personas_unicas),
            "Individuos distintos identificados unívocamente tras resolución de identidad.",
        ),
        (
            "Asistencias Recurrentes",
            str(resumen.total_recurrencia),
            "Participaciones adicionales registradas por personas que asistieron a más de un evento.",
        ),
        (
            "Tasa de Recurrencia",
            f"{resumen.tasa_recurrencia:.2f}%",
            "Proporción porcentual de recurrencia sobre la asistencia bruta total.",
        ),
    ]

    for i, (indicador, valor, desc) in enumerate(kpis):
        TableBuilder.add_row(
            table_kpi,
            values=[indicador, valor, desc],
            col_widths=col_widths_kpi,
            alignments=[WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.RIGHT, WD_ALIGN_PARAGRAPH.LEFT],
            is_even=(i % 2 == 1),
            bold=(i in (1, 2)),  # Destacar asistencia bruta y personas únicas
        )

    p_sp1 = doc.add_paragraph()
    p_sp1.paragraph_format.space_before = Pt(4)
    p_sp1.paragraph_format.space_after = Pt(4)

    # 2. Resumen Demográfico por Sexo Canónico
    StyleManager.add_heading_2(doc, "2.2. Distribución General por Sexo (Semántica Canónica)")

    headers_sexo = ["Sexo Canónico", "Participantes Nominales", "Proporción (%)"]
    col_widths_sexo = [Inches(3.2), Inches(2.2), Inches(2.2)]

    table_sexo = TableBuilder.create_table(doc, headers_sexo, col_widths_sexo)

    ds = resumen.distribucion_sexo
    filas_sexo = [
        ("Femenino (Mujeres)", str(ds.femenino), f"{ds.porcentaje_femenino:.2f}%" if ds.porcentaje_femenino is not None else "—"),
        ("Masculino (Varones)", str(ds.masculino), f"{ds.porcentaje_masculino:.2f}%" if ds.porcentaje_masculino is not None else "—"),
    ]

    for i, (sexo, cant, pct) in enumerate(filas_sexo):
        TableBuilder.add_row(
            table_sexo,
            values=[sexo, cant, pct],
            col_widths=col_widths_sexo,
            alignments=[WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.RIGHT, WD_ALIGN_PARAGRAPH.RIGHT],
            is_even=(i % 2 == 1),
        )

    # Fila total
    TableBuilder.add_row(
        table_sexo,
        values=["Total General", str(ds.total), "100.00%"],
        col_widths=col_widths_sexo,
        alignments=[WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.RIGHT, WD_ALIGN_PARAGRAPH.RIGHT],
        is_total_row=True,
    )

    p_sp2 = doc.add_paragraph()
    p_sp2.paragraph_format.space_before = Pt(4)
    p_sp2.paragraph_format.space_after = Pt(4)

    # 3. Estado Institucional de Concordancia y Auditoría
    StyleManager.add_heading_2(doc, "2.3. Estado Institucional de Concordancia y Calidad de Datos")

    headers_aud = ["Dimensión de Auditoría", "Cantidad", "Estado General"]
    col_widths_aud = [Inches(4.0), Inches(2.0), Inches(3.4)]

    table_aud = TableBuilder.create_table(doc, headers_aud, col_widths_aud)

    filas_aud = [
        ("Actividades 100% Concordantes (0 diferencias)", str(resumen.actividades_concordantes), "CONCORDANTE"),
        ("Actividades con Observaciones / Discrepancias", str(resumen.actividades_con_observaciones), "REQUIERE_REVISION" if resumen.actividades_con_observaciones > 0 else "CONCORDANTE"),
        ("Total de Discrepancias Identificadas", str(resumen.total_discrepancias), "REQUIERE_REVISION" if resumen.total_discrepancias > 0 else "CONCORDANTE"),
        ("Fuentes Nominales Pendientes / Faltantes", str(resumen.total_fuentes_faltantes), "REQUIERE_REVISION" if resumen.total_fuentes_faltantes > 0 else "CONCORDANTE"),
    ]

    for i, (dim, cant, estado) in enumerate(filas_aud):
        row = table_aud.add_row()
        TableBuilder.prevent_row_split(row)
        
        # Celda 0
        c0 = row.cells[0]
        c0.width = col_widths_aud[0]
        TableBuilder.set_cell_margins(c0)
        p0 = c0.paragraphs[0]
        p0.add_run(dim).font.name = StyleManager.FONT_NAME
        p0.runs[0].font.size = StyleManager.SIZE_TABLE_BODY
        
        # Celda 1
        c1 = row.cells[1]
        c1.width = col_widths_aud[1]
        TableBuilder.set_cell_margins(c1)
        p1 = c1.paragraphs[0]
        p1.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        p1.add_run(cant).font.name = StyleManager.FONT_NAME
        p1.runs[0].font.size = StyleManager.SIZE_TABLE_BODY
        p1.runs[0].bold = True
        
        # Celda 2: Badge de auditoría
        c2 = row.cells[2]
        TableBuilder.add_status_cell(c2, estado, width=col_widths_aud[2])

    # 4. Alertas principales si existen
    if resumen.alertas_principales:
        StyleManager.add_heading_2(doc, "2.4. Alertas Principales de Auditoría", space_before=8)
        for alerta in resumen.alertas_principales:
            StyleManager.add_callout(
                doc,
                text=alerta,
                title="ALERTA DE AUDITORÍA",
                bg_hex=ColorPalette.REVISION_BG,
                border_hex=ColorPalette.REVISION_TEXT,
            )

    # 5. Sinopsis respaldada (sólo si existe)
    if resumen.sinopsis:
        StyleManager.add_heading_2(doc, "2.5. Sinopsis Institucional", space_before=8)
        StyleManager.add_paragraph(doc, resumen.sinopsis, italic=True)
