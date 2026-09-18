"""
app.word_consolidator.docx_generator.sections.demographics

Renderizador de la Sección 5: Análisis Demográfico y Sexo.
Aplica con rigor la semántica canónica de sexo (FEMENINO y MASCULINO),
presenta las métricas de recurrencia y expone el cotejo obligatorio entre
Matriz 1 y Nominales (ej. Femenino 12 vs 13, Masculino 6 vs 5) sin reconciliación forzada.
"""

import docx
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt

from app.word_consolidator.document.models import DemografiaDocumental
from app.word_consolidator.docx_generator.style_manager import ColorPalette, StyleManager
from app.word_consolidator.docx_generator.table_builder import TableBuilder
from app.word_consolidator.models import EstadoDiscrepancia


def render_demographics(
    doc: docx.Document,
    demografia: DemografiaDocumental,
) -> None:
    """
    Construye la sección de Análisis Demográfico y Sexo.
    """
    StyleManager.add_heading_1(doc, "5. ANÁLISIS DEMOGRÁFICO Y DISTRIBUCIÓN POR SEXO")

    StyleManager.add_paragraph(
        doc,
        "Esta sección expone la composición demográfica institucional calculada a partir de los "
        "registros nominales unívocos y contrasta estos resultados con los totales agregados "
        "declarados originalmente en Matriz 1. Se respeta estrictamente la semántica canónica de género "
        "(FEMENINO y MASCULINO), sin asumir abreviaturas ambiguas.",
        space_after=8,
    )

    # 1. Métricas de Asistencia Neta y Recurrencia
    StyleManager.add_heading_2(doc, "5.1. Métricas de Participación Neta vs Recurrencia")

    met = demografia.metricas_recurrencia
    headers_rec = ["Concepto de Medición", "Valor Numérico", "Proporción Relativa"]
    widths_rec = [Inches(3.8), Inches(2.2), Inches(3.4)]
    table_rec = TableBuilder.create_table(doc, headers_rec, widths_rec)

    filas_rec = [
        ("Total Asistencia Bruta (Participaciones)", str(met.total_asistencia_bruta), "100.00% de asistencias"),
        ("Total Personas Únicas Netas (Población Atendida)", str(met.total_personas_unicas), f"{((met.total_personas_unicas / met.total_asistencia_bruta * 100) if met.total_asistencia_bruta > 0 else 100):.2f}% personas distintas"),
        ("Asistencias Recurrentes Identificadas", str(met.total_recurrencia), f"{met.tasa_recurrencia:.2f}% tasa de recurrencia"),
    ]

    for r_idx, (c, v, p) in enumerate(filas_rec):
        TableBuilder.add_row(
            table_rec,
            values=[c, v, p],
            col_widths=widths_rec,
            alignments=[WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.RIGHT, WD_ALIGN_PARAGRAPH.LEFT],
            is_even=(r_idx % 2 == 1),
            bold=(r_idx in (0, 1)),
        )

    # 2. Distribución Nominal por Sexo Canónico
    StyleManager.add_heading_2(doc, "5.2. Distribución Nominal por Sexo (M2–M5)")

    nom = demografia.distribucion_sexo_nominal
    headers_nom_sex = ["Sexo Canónico", "Denominación Institucional", "Total Nominal", "Porcentaje (%)"]
    widths_nom_sex = [Inches(2.5), Inches(2.5), Inches(2.2), Inches(2.2)]
    table_nom_sex = TableBuilder.create_table(doc, headers_nom_sex, widths_nom_sex)

    f_pct = f"{nom.porcentaje_femenino:.2f}%" if nom.porcentaje_femenino is not None else "—"
    m_pct = f"{nom.porcentaje_masculino:.2f}%" if nom.porcentaje_masculino is not None else "—"

    TableBuilder.add_row(
        table_nom_sex,
        values=["FEMENINO", "Mujeres", str(nom.femenino), f_pct],
        col_widths=widths_nom_sex,
        alignments=[WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.RIGHT, WD_ALIGN_PARAGRAPH.RIGHT],
        is_even=False,
    )
    TableBuilder.add_row(
        table_nom_sex,
        values=["MASCULINO", "Varones", str(nom.masculino), m_pct],
        col_widths=widths_nom_sex,
        alignments=[WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.RIGHT, WD_ALIGN_PARAGRAPH.RIGHT],
        is_even=True,
    )
    TableBuilder.add_row(
        table_nom_sex,
        values=["TOTAL", "Total Consolidado", str(nom.total), "100.00%"],
        col_widths=widths_nom_sex,
        alignments=[WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.RIGHT, WD_ALIGN_PARAGRAPH.RIGHT],
        is_total_row=True,
    )

    # 3. Comparativa Obligatoria M1 vs Nominal por Sexo (Cotejo sin Reconciliación)
    StyleManager.add_heading_2(doc, "5.3. Cotejo Dimensional de Sexo: Declarado en M1 vs Nominal (M2–M5)")

    StyleManager.add_paragraph(
        doc,
        "Nota Técnica de Auditoría: En Matriz 1, las siglas corresponden a F = Femenino (Mujeres) "
        "y M = Masculino (Varones). Se presentan ambos valores de forma simultánea sin sustitución ni reconciliación forzada.",
        italic=True,
        space_after=6,
    )

    headers_comp = ["Dimensión / Sexo", "Declarado en M1", "Nominal Verificado", "Delta (Nom - M1)", "Estado de Concordancia"]
    widths_comp = [Inches(2.5), Inches(1.6), Inches(1.6), Inches(1.4), Inches(2.3)]
    table_comp = TableBuilder.create_table(doc, headers_comp, widths_comp)

    if demografia.comparativa_sexo:
        for idx_c, item_c in enumerate(demografia.comparativa_sexo):
            row_c = table_comp.add_row()
            TableBuilder.prevent_row_split(row_c)
            is_total = ("TOTAL" in item_c.dimension.upper())

            # Formato de valores
            val_a_str = str(item_c.valor_fuente_a) if item_c.valor_fuente_a is not None else "—"
            val_b_str = str(item_c.valor_fuente_b) if item_c.valor_fuente_b is not None else "—"
            delta_str = f"+{item_c.delta}" if (item_c.delta is not None and item_c.delta > 0) else str(item_c.delta or 0)

            vals = [item_c.dimension, val_a_str, val_b_str, delta_str]
            aligns = [WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.RIGHT, WD_ALIGN_PARAGRAPH.RIGHT, WD_ALIGN_PARAGRAPH.RIGHT]

            for c_idx, (v_txt, al_txt) in enumerate(zip(vals, aligns)):
                cell_c = row_c.cells[c_idx]
                cell_c.width = widths_comp[c_idx]
                TableBuilder.set_cell_margins(cell_c)
                if is_total:
                    TableBuilder.set_cell_background(cell_c, ColorPalette.BG_SUBHEADER)
                elif idx_c % 2 == 1:
                    TableBuilder.set_cell_background(cell_c, ColorPalette.BG_LIGHT_GRAY)
                
                pc = cell_c.paragraphs[0]
                pc.alignment = al_txt
                rc = pc.add_run(v_txt)
                rc.font.name = StyleManager.FONT_NAME
                rc.font.size = StyleManager.SIZE_TABLE_BODY
                rc.bold = is_total

            # Columna Estado
            cell_est = row_c.cells[4]
            TableBuilder.add_status_cell(cell_est, item_c.estado, width=widths_comp[4])
    else:
        # Si no vino la lista de comparativas, construir a partir de distribucion_sexo_m1 y nominal
        m1_dist = demografia.distribucion_sexo_m1
        filas_respaldo = [
            ("Femenino (Mujeres)", m1_dist.femenino if m1_dist else None, nom.femenino),
            ("Masculino (Varones)", m1_dist.masculino if m1_dist else None, nom.masculino),
            ("Total General", m1_dist.total if m1_dist else None, nom.total),
        ]
        for idx_r, (dim_lbl, v_m1, v_nom) in enumerate(filas_respaldo):
            row_r = table_comp.add_row()
            TableBuilder.prevent_row_split(row_r)
            is_tot = (idx_r == 2)
            d_val = (v_nom - v_m1) if v_m1 is not None else None
            d_txt = (f"+{d_val}" if d_val > 0 else str(d_val)) if d_val is not None else "—"
            v_m1_txt = str(v_m1) if v_m1 is not None else "—"
            
            vals_r = [dim_lbl, v_m1_txt, str(v_nom), d_txt]
            aligns_r = [WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.RIGHT, WD_ALIGN_PARAGRAPH.RIGHT, WD_ALIGN_PARAGRAPH.RIGHT]
            for cr_idx, (vr_txt, alr_txt) in enumerate(zip(vals_r, aligns_r)):
                cell_cr = row_r.cells[cr_idx]
                cell_cr.width = widths_comp[cr_idx]
                TableBuilder.set_cell_margins(cell_cr)
                if is_tot:
                    TableBuilder.set_cell_background(cell_cr, ColorPalette.BG_SUBHEADER)
                elif idx_r % 2 == 1:
                    TableBuilder.set_cell_background(cell_cr, ColorPalette.BG_LIGHT_GRAY)
                pcr = cell_cr.paragraphs[0]
                pcr.alignment = alr_txt
                rcr = pcr.add_run(vr_txt)
                rcr.font.name = StyleManager.FONT_NAME
                rcr.font.size = StyleManager.SIZE_TABLE_BODY
                rcr.bold = is_tot
            
            cell_est_r = row_r.cells[4]
            est_r = EstadoDiscrepancia.CONCORDANTE if (d_val == 0) else EstadoDiscrepancia.REQUIERE_REVISION
            TableBuilder.add_status_cell(cell_est_r, est_r, width=widths_comp[4])
