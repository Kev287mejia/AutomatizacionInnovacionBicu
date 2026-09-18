"""
app.word_consolidator.docx_generator.sections.estamentos

Renderizador de la Sección 6: Estamentos y Carreras Universitarias.
Aplica con rigor la separación institucional entre Docentes y Administrativos,
presentando cada estamento de forma independiente sin fusiones no autorizadas.
"""

import docx
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt

from app.word_consolidator.document.models import EstamentosDocumental
from app.word_consolidator.docx_generator.style_manager import ColorPalette, StyleManager
from app.word_consolidator.docx_generator.table_builder import TableBuilder


def render_estamentos(
    doc: docx.Document,
    estamentos: EstamentosDocumental,
) -> None:
    """
    Construye la sección de Estamentos y Carreras.
    """
    StyleManager.add_heading_1(doc, "6. ESTAMENTOS INSTITUCIONALES Y CARRERAS")

    StyleManager.add_paragraph(
        doc,
        "Esta sección presenta la composición de las asistencias desglosada por estamento institucional. "
        "En cumplimiento irrestricto de las normas universitarias, el personal académico (Docentes) "
        "y el personal de apoyo (Administrativos) se mantienen estrictamente desagregados sin fusión.",
        space_after=8,
    )

    # 1. Tabla de Estamentos Institucionales
    StyleManager.add_heading_2(doc, "6.1. Distribución Consolidada por Estamento")

    headers_est = ["Estamento Institucional", "Femenino", "Masculino", "Total Participantes", "Proporción (%)"]
    widths_est = [Inches(3.2), Inches(1.5), Inches(1.5), Inches(1.8), Inches(1.4)]
    table_est = TableBuilder.create_table(doc, headers_est, widths_est)

    tot_general = estamentos.total_general if estamentos.total_general > 0 else 1

    def _calc_pct(val: int) -> str:
        pct = (val / tot_general) * 100.0
        return f"{pct:.2f}%"

    filas_est = [
        ("Estudiantes Matriculados (M2)", estamentos.estudiantes.femenino, estamentos.estudiantes.masculino, estamentos.estudiantes.total),
        ("Docentes y Académicos (M3)", estamentos.docentes.femenino, estamentos.docentes.masculino, estamentos.docentes.total),
        ("Personal Administrativo (M3)", estamentos.administrativos.femenino, estamentos.administrativos.masculino, estamentos.administrativos.total),
        ("Colaboradores Externos (M4)", estamentos.colaboradores.femenino, estamentos.colaboradores.masculino, estamentos.colaboradores.total),
        ("Protagonistas y Beneficiarios (M5)", estamentos.beneficiarios.femenino, estamentos.beneficiarios.masculino, estamentos.beneficiarios.total),
    ]

    for i, (nombre, fem, masc, tot) in enumerate(filas_est):
        TableBuilder.add_row(
            table_est,
            values=[nombre, str(fem), str(masc), str(tot), _calc_pct(tot)],
            col_widths=widths_est,
            alignments=[
                WD_ALIGN_PARAGRAPH.LEFT,
                WD_ALIGN_PARAGRAPH.RIGHT,
                WD_ALIGN_PARAGRAPH.RIGHT,
                WD_ALIGN_PARAGRAPH.RIGHT,
                WD_ALIGN_PARAGRAPH.RIGHT,
            ],
            is_even=(i % 2 == 1),
        )

    # Fila de Total
    tot_fem = (
        estamentos.estudiantes.femenino
        + estamentos.docentes.femenino
        + estamentos.administrativos.femenino
        + estamentos.colaboradores.femenino
        + estamentos.beneficiarios.femenino
    )
    tot_masc = (
        estamentos.estudiantes.masculino
        + estamentos.docentes.masculino
        + estamentos.administrativos.masculino
        + estamentos.colaboradores.masculino
        + estamentos.beneficiarios.masculino
    )

    TableBuilder.add_row(
        table_est,
        values=["TOTAL GENERAL DE ESTAMENTOS", str(tot_fem), str(tot_masc), str(estamentos.total_general), "100.00%"],
        col_widths=widths_est,
        alignments=[
            WD_ALIGN_PARAGRAPH.LEFT,
            WD_ALIGN_PARAGRAPH.RIGHT,
            WD_ALIGN_PARAGRAPH.RIGHT,
            WD_ALIGN_PARAGRAPH.RIGHT,
            WD_ALIGN_PARAGRAPH.RIGHT,
        ],
        is_total_row=True,
    )

    p_s1 = doc.add_paragraph()
    p_s1.paragraph_format.space_before = Pt(4)
    p_s1.paragraph_format.space_after = Pt(4)

    # 2. Comparativa de Estamentos (si existe en el modelo)
    if estamentos.comparativa_estamentos:
        StyleManager.add_heading_2(doc, "6.2. Cotejo de Estamentos entre Fuentes")
        headers_comp_e = ["Estamento", "Fuente A", "Fuente B", "Delta", "Estado"]
        widths_comp_e = [Inches(2.5), Inches(1.8), Inches(1.8), Inches(1.0), Inches(2.3)]
        table_comp_e = TableBuilder.create_table(doc, headers_comp_e, widths_comp_e)

        for c_idx, ce in enumerate(estamentos.comparativa_estamentos):
            row_ce = table_comp_e.add_row()
            TableBuilder.prevent_row_split(row_ce)
            v_a_str = str(ce.valor_fuente_a) if ce.valor_fuente_a is not None else "—"
            v_b_str = str(ce.valor_fuente_b) if ce.valor_fuente_b is not None else "—"
            delta_str = f"+{ce.delta}" if (ce.delta is not None and ce.delta > 0) else str(ce.delta or 0)

            vals_ce = [ce.dimension, v_a_str, v_b_str, delta_str]
            aligns_ce = [WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.RIGHT, WD_ALIGN_PARAGRAPH.RIGHT, WD_ALIGN_PARAGRAPH.RIGHT]

            for col_i, (v_ce, al_ce) in enumerate(zip(vals_ce, aligns_ce)):
                cell_ce = row_ce.cells[col_i]
                cell_ce.width = widths_comp_e[col_i]
                TableBuilder.set_cell_margins(cell_ce)
                if c_idx % 2 == 1:
                    TableBuilder.set_cell_background(cell_ce, ColorPalette.BG_LIGHT_GRAY)
                pce = cell_ce.paragraphs[0]
                pce.alignment = al_ce
                rce = pce.add_run(v_ce)
                rce.font.name = StyleManager.FONT_NAME
                rce.font.size = StyleManager.SIZE_TABLE_BODY

            # Estado
            cell_est_ce = row_ce.cells[4]
            TableBuilder.add_status_cell(cell_est_ce, ce.estado, width=widths_comp_e[4])

        p_s2 = doc.add_paragraph()
        p_s2.paragraph_format.space_before = Pt(4)
        p_s2.paragraph_format.space_after = Pt(4)

    # 3. Desglose por Carreras Universitarias (sólo si existen datos respaldados)
    if estamentos.desglose_carreras:
        StyleManager.add_heading_2(doc, "6.3. Distribución de Estudiantes por Carrera Universitaria")
        headers_car = ["Carrera Universitaria", "Femenino", "Masculino", "Total Estudiantes"]
        widths_car = [Inches(4.4), Inches(1.6), Inches(1.6), Inches(1.8)]
        table_car = TableBuilder.create_table(doc, headers_car, widths_car)

        for k, car in enumerate(estamentos.desglose_carreras):
            TableBuilder.add_row(
                table_car,
                values=[car.carrera, str(car.femenino), str(car.masculino), str(car.total)],
                col_widths=widths_car,
                alignments=[
                    WD_ALIGN_PARAGRAPH.LEFT,
                    WD_ALIGN_PARAGRAPH.RIGHT,
                    WD_ALIGN_PARAGRAPH.RIGHT,
                    WD_ALIGN_PARAGRAPH.RIGHT,
                ],
                is_even=(k % 2 == 1),
            )
