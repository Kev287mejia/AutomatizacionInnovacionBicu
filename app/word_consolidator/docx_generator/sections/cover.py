"""
app.word_consolidator.docx_generator.sections.cover

Renderizador de la Sección 1: Portada Institucional.
Cumple estrictamente con la regla de NO INVENTAR información:
Si fecha, lugar o metadatos no existen en el modelo, se omiten limpiamente.
"""

from typing import Optional

import docx
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls
from docx.shared import Inches, Mm, Pt

from app.word_consolidator.document.models import PortadaDocumental
from app.word_consolidator.docx_generator.style_manager import ColorPalette, StyleManager
from app.word_consolidator.models import MetadatosInstitucionales, PeriodoConsolidacion


def render_cover(
    doc: docx.Document,
    portada: PortadaDocumental,
    periodo: PeriodoConsolidacion,
    metadatos: Optional[MetadatosInstitucionales] = None,
) -> None:
    """
    Construye la portada formal e institucional del informe.
    """
    # 1. Espaciado superior inicial para centrar visualmente
    p_top = doc.add_paragraph()
    p_top.paragraph_format.space_before = Pt(36)
    p_top.paragraph_format.space_after = Pt(12)
    p_top.alignment = WD_ALIGN_PARAGRAPH.CENTER

    # 2. Nombre de la Universidad (si existe)
    universidad = portada.universidad or (metadatos.universidad if metadatos else None)
    if universidad:
        p_univ = doc.add_paragraph()
        p_univ.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_univ.paragraph_format.space_before = Pt(0)
        p_univ.paragraph_format.space_after = Pt(4)
        run_univ = p_univ.add_run(universidad.upper())
        run_univ.font.name = StyleManager.FONT_NAME
        run_univ.font.size = Pt(14)
        run_univ.font.bold = True
        run_univ.font.color.rgb = StyleManager.hex_to_rgb(ColorPalette.PRIMARY_NAVY)

    # 3. Lema institucional (si existe)
    lema = portada.lema or (metadatos.lema if metadatos else None)
    if lema:
        p_lema = doc.add_paragraph()
        p_lema.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_lema.paragraph_format.space_before = Pt(0)
        p_lema.paragraph_format.space_after = Pt(20)
        run_lema = p_lema.add_run(f'"{lema}"')
        run_lema.font.name = StyleManager.FONT_NAME
        run_lema.font.size = Pt(9.5)
        run_lema.font.italic = True
        run_lema.font.color.rgb = StyleManager.hex_to_rgb(ColorPalette.ACCENT_SLATE)

    # 4. Línea decorativa central sobria
    table_line = doc.add_table(rows=1, cols=1)
    table_line.alignment = docx.enum.table.WD_TABLE_ALIGNMENT.CENTER
    table_line.autofit = False
    c_line = table_line.cell(0, 0)
    c_line.width = Inches(4.0)
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{ColorPalette.PRIMARY_NAVY}"/>')
    c_line._tc.get_or_add_tcPr().append(shd)
    p_l = c_line.paragraphs[0]
    p_l.paragraph_format.space_before = Pt(1)
    p_l.paragraph_format.space_after = Pt(1)

    # 5. Espaciador intermedio
    p_mid = doc.add_paragraph()
    p_mid.paragraph_format.space_before = Pt(24)
    p_mid.paragraph_format.space_after = Pt(0)

    # 6. Título formal del informe
    titulo = portada.titulo or "INFORME CONSOLIDADO INSTITUCIONAL DE ACTIVIDADES"
    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_title.paragraph_format.space_before = Pt(0)
    p_title.paragraph_format.space_after = Pt(8)
    run_title = p_title.add_run(titulo.upper())
    run_title.font.name = StyleManager.FONT_NAME
    run_title.font.size = Pt(18)
    run_title.font.bold = True
    run_title.font.color.rgb = StyleManager.hex_to_rgb(ColorPalette.PRIMARY_NAVY)

    # 7. Subtítulo (si existe)
    if portada.subtitulo:
        p_sub = doc.add_paragraph()
        p_sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_sub.paragraph_format.space_before = Pt(0)
        p_sub.paragraph_format.space_after = Pt(16)
        run_sub = p_sub.add_run(portada.subtitulo)
        run_sub.font.name = StyleManager.FONT_NAME
        run_sub.font.size = Pt(11)
        run_sub.font.color.rgb = StyleManager.hex_to_rgb(ColorPalette.ACCENT_SLATE)

    # 8. Período evaluado
    periodo_str = portada.periodo_texto or periodo.etiqueta
    p_per = doc.add_paragraph()
    p_per.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_per.paragraph_format.space_before = Pt(8)
    p_per.paragraph_format.space_after = Pt(4)
    r_per_lbl = p_per.add_run("Período Consolidado: ")
    r_per_lbl.font.name = StyleManager.FONT_NAME
    r_per_lbl.font.size = Pt(11)
    r_per_lbl.font.bold = True
    r_per_lbl.font.color.rgb = StyleManager.hex_to_rgb(ColorPalette.PRIMARY_NAVY)
    
    r_per_val = p_per.add_run(periodo_str)
    r_per_val.font.name = StyleManager.FONT_NAME
    r_per_val.font.size = Pt(11)
    r_per_val.font.color.rgb = StyleManager.hex_to_rgb(ColorPalette.TEXT_DARK)

    # 9. Área responsable (si existe)
    area = portada.area_responsable or (metadatos.area_responsable if metadatos else None)
    if area:
        p_area = doc.add_paragraph()
        p_area.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_area.paragraph_format.space_before = Pt(4)
        p_area.paragraph_format.space_after = Pt(4)
        r_area = p_area.add_run(area)
        r_area.font.name = StyleManager.FONT_NAME
        r_area.font.size = Pt(10)
        r_area.font.color.rgb = StyleManager.hex_to_rgb(ColorPalette.ACCENT_SLATE)

    # 10. Bloque inferior de emisión (sólo si fecha o lugar están explícitamente presentes)
    fecha = portada.fecha_emision
    lugar = portada.lugar_emision or (metadatos.lugar_emision if metadatos else None)
    
    p_bot_space = doc.add_paragraph()
    p_bot_space.paragraph_format.space_before = Pt(48)
    p_bot_space.paragraph_format.space_after = Pt(0)

    if lugar or fecha:
        p_emis = doc.add_paragraph()
        p_emis.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_emis.paragraph_format.space_before = Pt(0)
        p_emis.paragraph_format.space_after = Pt(2)
        texto_emision = " — ".join(filter(None, [lugar, fecha]))
        r_emis = p_emis.add_run(texto_emision)
        r_emis.font.name = StyleManager.FONT_NAME
        r_emis.font.size = Pt(9)
        r_emis.font.color.rgb = StyleManager.hex_to_rgb(ColorPalette.TEXT_MUTED)

    # Versión del sistema si está en metadatos
    version = portada.version_sistema or (metadatos.version_sistema if metadatos else None)
    if version:
        p_ver = doc.add_paragraph()
        p_ver.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_ver.paragraph_format.space_before = Pt(0)
        p_ver.paragraph_format.space_after = Pt(0)
        r_ver = p_ver.add_run(f"Versión del Sistema: {version}")
        r_ver.font.name = StyleManager.FONT_NAME
        r_ver.font.size = Pt(8)
        r_ver.font.color.rgb = StyleManager.hex_to_rgb(ColorPalette.TEXT_MUTED)
