"""
app.word_consolidator.docx_generator.sections.activity_cards

Renderizador de la Sección 4: Fichas de Detalle por Actividad.
Genera una ficha técnica institucional e independiente para cada actividad,
con comparativa M1 vs Nominal, desglose de estamentos, discrepancias y trazabilidad.
"""

from typing import List

import docx
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt

from app.word_consolidator.document.models import FichaActividadDocumental
from app.word_consolidator.docx_generator.style_manager import ColorPalette, StyleManager
from app.word_consolidator.docx_generator.table_builder import TableBuilder
from app.word_consolidator.models import EstadoDiscrepancia


def render_activity_cards(
    doc: docx.Document,
    fichas: List[FichaActividadDocumental],
) -> None:
    """
    Construye las fichas individuales de detalle por actividad.
    """
    StyleManager.add_heading_1(doc, "4. FICHAS DE DETALLE POR ACTIVIDAD")

    StyleManager.add_paragraph(
        doc,
        "A continuación se presenta la ficha técnica pormenorizada de cada actividad consolidada, "
        "incluyendo su localización, eje estratégico, cotejo numérico M1 frente a Nominal, "
        "desglose por estamentos, observaciones de auditoría y trazabilidad de celdas físicas.",
        space_after=12,
    )

    for idx, ficha in enumerate(fichas):
        num_act = idx + 1
        
        # Título de la Ficha
        StyleManager.add_heading_2(
            doc,
            f"Ficha {num_act}: {ficha.nombre_actividad}",
            space_before=10 if idx > 0 else 4,
        )

        # 1. Tabla de Metadatos y Clasificación
        headers_meta = ["Campo Institucional", "Detalle Registrado"]
        widths_meta = [Inches(2.5), Inches(6.9)]
        table_meta = TableBuilder.create_table(doc, headers_meta, widths_meta)

        # Territorio
        terr_parts = []
        if ficha.departamento:
            terr_parts.append(f"Depto: {ficha.departamento}")
        if ficha.municipio:
            terr_parts.append(f"Municipio: {ficha.municipio}")
        if ficha.comunidad:
            terr_parts.append(f"Comunidad: {ficha.comunidad}")
        terr_info = " — ".join(terr_parts) if terr_parts else "—"

        meta_rows = [
            ("Sede Universitaria", ficha.sede),
            ("Fecha de Ejecución", str(ficha.fecha) if ficha.fecha else "—"),
            ("Ubicación Geográfica", terr_info),
            ("Eje / Ámbito Estratégico", ficha.eje or "—"),
            ("Área Responsable", ficha.area_responsable or "—"),
            ("Tipo de Evento / Modalidad", ficha.tipo_evento or "—"),
        ]

        if ficha.programa:
            meta_rows.append(("Programa Institucional", ficha.programa))
        if ficha.proyecto:
            meta_rows.append(("Proyecto Institucional", ficha.proyecto))
        if ficha.descripcion_resultados:
            meta_rows.append(("Resultados Reportados", ficha.descripcion_resultados))

        for j, (campo, val) in enumerate(meta_rows):
            TableBuilder.add_row(
                table_meta,
                values=[campo, val],
                col_widths=widths_meta,
                alignments=[WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.LEFT],
                is_even=(j % 2 == 1),
                bold=(j == 0),
            )

        p_s1 = doc.add_paragraph()
        p_s1.paragraph_format.space_before = Pt(3)
        p_s1.paragraph_format.space_after = Pt(3)

        # 2. Tabla Comparativa M1 vs Nominal por Sexo y Total
        StyleManager.add_heading_3(doc, "Cotejo Numérico: Declarado en M1 vs Nominal Verificado (M2–M5)")

        headers_comp = ["Dimensión de Participación", "Valor M1", "Valor Nominal", "Delta (Nom - M1)", "Estado de Concordancia"]
        widths_comp = [Inches(2.5), Inches(1.5), Inches(1.5), Inches(1.5), Inches(2.4)]
        table_comp = TableBuilder.create_table(doc, headers_comp, widths_comp)

        comp = ficha.comparativa_m1_vs_nominal

        def _fmt_val(v):
            return str(v) if v is not None else "—"

        def _fmt_delta(d):
            if d is None:
                return "—"
            return f"+{d}" if d > 0 else str(d)

        comp_items = [
            ("Femenino (Mujeres)", comp.femenino_m1, comp.femenino_nominal, comp.delta_femenino),
            ("Masculino (Varones)", comp.masculino_m1, comp.masculino_nominal, comp.delta_masculino),
            ("Total General", comp.total_m1, comp.total_nominal, comp.delta_total),
        ]

        for k, (dim_label, v_m1, v_nom, delta_val) in enumerate(comp_items):
            row_c = table_comp.add_row()
            TableBuilder.prevent_row_split(row_c)
            is_tot = (k == 2)
            bg = ColorPalette.BG_SUBHEADER if is_tot else (ColorPalette.BG_LIGHT_GRAY if k % 2 == 1 else "FFFFFF")

            # Col 0: Dimensión
            c0 = row_c.cells[0]
            c0.width = widths_comp[0]
            if bg != "FFFFFF":
                TableBuilder.set_cell_background(c0, bg)
            TableBuilder.set_cell_margins(c0)
            p0 = c0.paragraphs[0]
            r0 = p0.add_run(dim_label)
            r0.font.name = StyleManager.FONT_NAME
            r0.font.size = StyleManager.SIZE_TABLE_BODY
            r0.bold = is_tot

            # Col 1: M1
            c1 = row_c.cells[1]
            c1.width = widths_comp[1]
            if bg != "FFFFFF":
                TableBuilder.set_cell_background(c1, bg)
            TableBuilder.set_cell_margins(c1)
            p1 = c1.paragraphs[0]
            p1.alignment = WD_ALIGN_PARAGRAPH.RIGHT
            r1 = p1.add_run(_fmt_val(v_m1))
            r1.font.name = StyleManager.FONT_NAME
            r1.font.size = StyleManager.SIZE_TABLE_BODY
            r1.bold = is_tot

            # Col 2: Nominal
            c2 = row_c.cells[2]
            c2.width = widths_comp[2]
            if bg != "FFFFFF":
                TableBuilder.set_cell_background(c2, bg)
            TableBuilder.set_cell_margins(c2)
            p2 = c2.paragraphs[0]
            p2.alignment = WD_ALIGN_PARAGRAPH.RIGHT
            r2 = p2.add_run(_fmt_val(v_nom))
            r2.font.name = StyleManager.FONT_NAME
            r2.font.size = StyleManager.SIZE_TABLE_BODY
            r2.bold = is_tot

            # Col 3: Delta
            c3 = row_c.cells[3]
            c3.width = widths_comp[3]
            if bg != "FFFFFF":
                TableBuilder.set_cell_background(c3, bg)
            TableBuilder.set_cell_margins(c3)
            p3 = c3.paragraphs[0]
            p3.alignment = WD_ALIGN_PARAGRAPH.RIGHT
            r3 = p3.add_run(_fmt_delta(delta_val))
            r3.font.name = StyleManager.FONT_NAME
            r3.font.size = StyleManager.SIZE_TABLE_BODY
            r3.bold = is_tot

            # Col 4: Estado
            c4 = row_c.cells[4]
            # Si delta es 0 y v_m1 existe, es CONCORDANTE; si delta != 0, REQUIERE_REVISION
            est_fila = EstadoDiscrepancia.CONCORDANTE if (delta_val == 0) else EstadoDiscrepancia.REQUIERE_REVISION
            TableBuilder.add_status_cell(c4, est_fila, width=widths_comp[4])

        p_s2 = doc.add_paragraph()
        p_s2.paragraph_format.space_before = Pt(3)
        p_s2.paragraph_format.space_after = Pt(3)

        # 3. Desglose de Estamentos
        if ficha.desglose_estamentos:
            StyleManager.add_heading_3(doc, "Desglose por Estamentos Institucionales")
            headers_est = ["Estamento", "Femenino", "Masculino", "Total", "Condición"]
            widths_est = [Inches(3.4), Inches(1.5), Inches(1.5), Inches(1.5), Inches(1.5)]
            table_est = TableBuilder.create_table(doc, headers_est, widths_est)

            for m, fila_est in enumerate(ficha.desglose_estamentos):
                row_e = table_est.add_row()
                TableBuilder.prevent_row_split(row_e)
                vals_e = [fila_est.tipo_estamento, str(fila_est.femenino), str(fila_est.masculino), str(fila_est.total)]
                aligns_e = [WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.RIGHT, WD_ALIGN_PARAGRAPH.RIGHT, WD_ALIGN_PARAGRAPH.RIGHT]
                
                for col_e_idx, (v_e, al_e) in enumerate(zip(vals_e, aligns_e)):
                    ce = row_e.cells[col_e_idx]
                    ce.width = widths_est[col_e_idx]
                    TableBuilder.set_cell_margins(ce)
                    if m % 2 == 1:
                        TableBuilder.set_cell_background(ce, ColorPalette.BG_LIGHT_GRAY)
                    pe = ce.paragraphs[0]
                    pe.alignment = al_e
                    re = pe.add_run(v_e)
                    re.font.name = StyleManager.FONT_NAME
                    re.font.size = StyleManager.SIZE_TABLE_BODY

                # Columna Condición
                ce_cond = row_e.cells[4]
                TableBuilder.add_status_cell(ce_cond, fila_est.estado, width=widths_est[4])

            p_s3 = doc.add_paragraph()
            p_s3.paragraph_format.space_before = Pt(3)
            p_s3.paragraph_format.space_after = Pt(3)

        # 4. Datos Narrativos (si existen en el modelo)
        if ficha.datos_narrativos:
            StyleManager.add_heading_3(doc, "Datos del Informe Narrativo Asociado")
            headers_narr = ["Dimensión Narrativa", "Valor Reportado"]
            widths_narr = [Inches(4.0), Inches(5.4)]
            table_narr = TableBuilder.create_table(doc, headers_narr, widths_narr)
            for n_idx, (k_dim, v_dim) in enumerate(ficha.datos_narrativos.items()):
                TableBuilder.add_row(
                    table_narr,
                    values=[k_dim.capitalize(), str(v_dim)],
                    col_widths=widths_narr,
                    alignments=[WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.LEFT],
                    is_even=(n_idx % 2 == 1),
                )
            p_s4 = doc.add_paragraph()
            p_s4.paragraph_format.space_before = Pt(3)
            p_s4.paragraph_format.space_after = Pt(3)

        # 5. Discrepancias Específicas de la Actividad
        if ficha.discrepancias:
            StyleManager.add_heading_3(doc, "Discrepancias y Hallazgos de Auditoría para esta Actividad")
            headers_d = ["Dimensión", "Fuente A", "Fuente B", "Delta", "Estado", "Descripción"]
            widths_d = [Inches(1.5), Inches(1.8), Inches(1.8), Inches(0.8), Inches(1.5), Inches(2.0)]
            table_d = TableBuilder.create_table(doc, headers_d, widths_d)

            for d_idx, item_d in enumerate(ficha.discrepancias):
                row_d = table_d.add_row()
                TableBuilder.prevent_row_split(row_d)
                delta_txt = f"+{item_d.delta}" if (item_d.delta is not None and item_d.delta > 0) else str(item_d.delta or 0)
                
                fa_txt = f"{item_d.fuente_a} ({item_d.valor_fuente_a})" if item_d.valor_fuente_a is not None else item_d.fuente_a
                fb_txt = f"{item_d.fuente_b} ({item_d.valor_fuente_b})" if item_d.valor_fuente_b is not None else item_d.fuente_b

                vals_d = [item_d.dimension, fa_txt, fb_txt, delta_txt]
                aligns_d = [WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.RIGHT]

                for col_d_idx, (v_d, al_d) in enumerate(zip(vals_d, aligns_d)):
                    cd = row_d.cells[col_d_idx]
                    cd.width = widths_d[col_d_idx]
                    TableBuilder.set_cell_margins(cd)
                    if d_idx % 2 == 1:
                        TableBuilder.set_cell_background(cd, ColorPalette.BG_LIGHT_GRAY)
                    pd = cd.paragraphs[0]
                    pd.alignment = al_d
                    rd = pd.add_run(v_d)
                    rd.font.name = StyleManager.FONT_NAME
                    rd.font.size = StyleManager.SIZE_TABLE_BODY

                # Estado
                cd_est = row_d.cells[4]
                TableBuilder.add_status_cell(cd_est, item_d.estado, width=widths_d[4])

                # Descripción
                cd_desc = row_d.cells[5]
                cd_desc.width = widths_d[5]
                TableBuilder.set_cell_margins(cd_desc)
                if d_idx % 2 == 1:
                    TableBuilder.set_cell_background(cd_desc, ColorPalette.BG_LIGHT_GRAY)
                pd_desc = cd_desc.paragraphs[0]
                rd_desc = pd_desc.add_run(item_d.descripcion or "—")
                rd_desc.font.name = StyleManager.FONT_NAME
                rd_desc.font.size = StyleManager.SIZE_TABLE_BODY

            p_s5 = doc.add_paragraph()
            p_s5.paragraph_format.space_before = Pt(3)
            p_s5.paragraph_format.space_after = Pt(3)

        # 6. Trazabilidad de Fuentes Físicas
        if ficha.trazabilidad_fuentes:
            StyleManager.add_heading_3(doc, "Trazabilidad de Fuentes de Origen")
            headers_t = ["Matriz", "Archivo Origen", "Hoja", "Fila Física", "Período"]
            widths_t = [Inches(1.0), Inches(3.4), Inches(2.2), Inches(1.2), Inches(1.6)]
            table_t = TableBuilder.create_table(doc, headers_t, widths_t)

            for t_idx, tr in enumerate(ficha.trazabilidad_fuentes):
                TableBuilder.add_row(
                    table_t,
                    values=[
                        tr.matriz,
                        tr.archivo or "Archivo Matriz Oficial",
                        tr.hoja or "—",
                        str(tr.fila) if tr.fila else "—",
                        tr.periodo or "—",
                    ],
                    col_widths=widths_t,
                    alignments=[
                        WD_ALIGN_PARAGRAPH.CENTER,
                        WD_ALIGN_PARAGRAPH.LEFT,
                        WD_ALIGN_PARAGRAPH.LEFT,
                        WD_ALIGN_PARAGRAPH.CENTER,
                        WD_ALIGN_PARAGRAPH.LEFT,
                    ],
                    is_even=(t_idx % 2 == 1),
                )

        # Separador visual entre fichas
        if idx < len(fichas) - 1:
            p_div = doc.add_paragraph()
            p_div.paragraph_format.space_before = Pt(12)
            p_div.paragraph_format.space_after = Pt(12)
            r_div = p_div.add_run("―" * 40)
            r_div.font.color.rgb = StyleManager.hex_to_rgb(ColorPalette.BORDER_STRONG)
