"""
app.word_consolidator.institutional_docx.generator

Orquestador principal del Generador DOCX Institucional del Producto B.
Convierte un 'InformeSemanalInstitucional' en un documento Word institucional
idéntico al formato de 'Documento de Kevds3.docx'.

Implementa la cadena completa:
    InformeSemanalInstitucional
        → InstitutionalDocxSpec (build_document_layout_spec)
        → page_setup  (configure_page + add_institutional_headers)
        → table_builder (build_institutional_table)
        → evidences_builder (build_evidences_section)
        → Document.save()

Política de fuente predeterminada: ModoFuenteDatosEnum.DECLARADO_M1
(conforme a la decisión funcional oficial de la Fase 16.5).
"""

import hashlib
import io
import os
from pathlib import Path
from typing import Union

from docx import Document

from app.word_consolidator.document.institutional_models import (
    InformeSemanalInstitucional,
)
from app.word_consolidator.institutional_docx.spec import (
    ModoFuenteDatosEnum,
    build_document_layout_spec,
)
from app.word_consolidator.institutional_docx.page_setup import (
    configure_page,
    add_institutional_headers,
)
from app.word_consolidator.institutional_docx.table_builder import (
    build_institutional_table,
)
from app.word_consolidator.institutional_docx.evidences_builder import (
    build_evidences_section,
)


class InstitutionalDocxGenerator:
    """
    Generador del Informe Semanal Ejecutivo Institucional BICU (Producto B).

    Usage:
        generator = InstitutionalDocxGenerator()
        buffer = generator.generate(informe)
        # o bien:
        generator.generate(informe, output_path="ruta/al/archivo.docx")
    """

    def __init__(
        self,
        modo_fuente: ModoFuenteDatosEnum = ModoFuenteDatosEnum.DECLARADO_M1,
    ) -> None:
        """
        Args:
            modo_fuente: Política de fuente de datos para la tabla ejecutiva.
                Predeterminado: DECLARADO_M1 (datos declarados en Matriz 1).
        """
        self.modo_fuente = modo_fuente

    def generate(
        self,
        informe: InformeSemanalInstitucional,
        output_path: Union[str, Path, None] = None,
    ) -> io.BytesIO:
        """
        Genera el documento Word institucional a partir del informe.

        Args:
            informe: Modelo inmutable con todos los datos del informe semanal.
            output_path: Ruta opcional donde guardar el DOCX.
                Si se provee, el archivo se escribe en disco.
                Siempre se retorna el buffer en memoria.

        Returns:
            BytesIO con el contenido completo del archivo DOCX generado.
        """
        # 1. Calcular especificación completa (pura, sin efectos)
        doc_spec = build_document_layout_spec(informe, modo_fuente=self.modo_fuente)

        # 2. Inicializar documento Word en blanco
        doc = Document()

        # 3. Configurar página (Landscape, Letter, márgenes institucionales)
        configure_page(doc)

        # 4. Eliminar el párrafo vacío inicial que python-docx inserta por defecto
        for para in list(doc.paragraphs):
            p = para._p
            p.getparent().remove(p)

        # 5. Agregar encabezados institucionales
        add_institutional_headers(doc)

        # 6. Agregar un párrafo separador entre los títulos y la tabla
        sep = doc.add_paragraph()
        sep.paragraph_format.space_before = __import__("docx.shared", fromlist=["Pt"]).Pt(4)
        sep.paragraph_format.space_after = __import__("docx.shared", fromlist=["Pt"]).Pt(4)

        # 7. Construir la tabla ejecutiva de 13 columnas
        build_institutional_table(doc, doc_spec.tabla_ejecutiva)

        # 8. Construir la sección EVIDENCIAS ANEXAS
        build_evidences_section(
            doc=doc,
            evidencias_spec=doc_spec.seccion_evidencias,
            tiene_discrepancias=doc_spec.tiene_discrepancias,
            nombre_depto=informe.departamento_responsable,
        )

        # 9. Serializar a BytesIO
        buffer = io.BytesIO()
        doc.save(buffer)
        buffer.seek(0)

        # 10. Guardar en disco si se especificó ruta
        if output_path is not None:
            path = Path(output_path)
            path.parent.mkdir(parents=True, exist_ok=True)
            with open(path, "wb") as f:
                f.write(buffer.read())
            buffer.seek(0)

        return buffer

    def generate_and_compute_sha256(
        self,
        informe: InformeSemanalInstitucional,
        output_path: Union[str, Path, None] = None,
    ) -> tuple[io.BytesIO, str]:
        """
        Genera el documento y computa su hash SHA-256.

        Returns:
            Tuple[BytesIO, str]: buffer del DOCX y hash SHA-256 hexadecimal.
        """
        buffer = self.generate(informe, output_path=output_path)
        content = buffer.read()
        sha256 = hashlib.sha256(content).hexdigest()
        buffer.seek(0)
        return buffer, sha256


def generate_institutional_docx(
    informe: InformeSemanalInstitucional,
    output_path: Union[str, Path, None] = None,
    modo_fuente: ModoFuenteDatosEnum = ModoFuenteDatosEnum.DECLARADO_M1,
) -> io.BytesIO:
    """
    Función de conveniencia para generar el DOCX institucional.

    Args:
        informe: Modelo institucional completo.
        output_path: Ruta de escritura opcional.
        modo_fuente: Política de fuente de datos.

    Returns:
        BytesIO con el DOCX generado.
    """
    generator = InstitutionalDocxGenerator(modo_fuente=modo_fuente)
    return generator.generate(informe, output_path=output_path)
