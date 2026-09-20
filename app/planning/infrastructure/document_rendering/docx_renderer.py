"""
Adaptador de Infraestructura — Generador DOCX Institucional para Diseño Metodológico.

Fuente arquitectónica:
  FASE_29_5_1_ANALISIS_INTEGRAL_DOCX_Y_M1_M5.md §4, §5, §10
  FASE_29_5_2_AUDITORIA_CIERRE_DISENO_GENERADOR_DOCX.md §11, §14, §16
  Fase 29.5.3 — Implementación Controlada del Renderer DOCX

Clasificación: ADAPTADOR DE INFRAESTRUCTURA (Inversión de Dependencias).
Implementa: MethodologicalDocumentRendererPort

REGLAS ABSOLUTAS:
  - Cero dependencias de app.word_consolidator (núcleo protegido).
  - Cero consultas a SQLite, UnitOfWork o repositorios (Ceguera a Reglas).
  - Cero llamadas a IA (OpenRouter, Gemini, Ollama).
  - Cero invención de datos o reglas no provistas en el DTO.
  - Representación fidedigna de los cinco bloques institucionales.
"""
from __future__ import annotations

import os
import pathlib
from typing import Optional, Sequence

import docx
from docx.enum.section import WD_ORIENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls
from docx.shared import Inches, Pt, RGBColor

from app.planning.domain.dtos import (
    FAQTableDTO,
    MethodologicalDesignDTO,
    OperationalActivityDTO,
    TimeBlockDTO,
)
from app.planning.domain.ports import MethodologicalDocumentRendererPort


class DocumentRenderingError(Exception):
    """Excepción técnica levantada ante fallos en la generación o precondiciones del DOCX.

    Pertenece exclusivamente a la capa de infraestructura/adaptador y no contamina el dominio.
    """
    pass


# ---------------------------------------------------------------------------
# CONSTANTES DE ESTILO Y MAQUETACIÓN INSTITUCIONAL
# Fuente: FASE_29_5_1 §4.B, FASE_29_5_2 §16.3
# ---------------------------------------------------------------------------
COLOR_NAVY_HEX = "002060"
COLOR_NAVY_RGB = RGBColor(0, 32, 96)
COLOR_LIGHT_BLUE_HEX = "B4C6E7"
COLOR_ICE_BLUE_HEX = "DEEAF6"
COLOR_BORDER_HEX = "B4C6E7"

FONT_BASE = "Times New Roman"
FONT_HEADING = "Tahoma"

PT_TITLE = Pt(14)
PT_HEADING = Pt(11)
PT_BASE = Pt(12)
PT_TABLE = Pt(10)


# ---------------------------------------------------------------------------
# HELPERS DE XML PARA FORMATO Y ESTILOS
# ---------------------------------------------------------------------------
def _set_cell_shading(cell, color_hex: str) -> None:
    """Aplica sombreado de fondo a una celda de tabla."""
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{color_hex}"/>')
    tcPr.append(shd)


def _set_cell_margins(cell, top: int = 100, bottom: int = 100, left: int = 150, right: int = 150) -> None:
    """Establece márgenes internos (padding) en dxa para legibilidad."""
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(
        f'<w:tcMar {nsdecls("w")}>'
        f'  <w:top w:w="{top}" w:type="dxa"/>'
        f'  <w:bottom w:w="{bottom}" w:type="dxa"/>'
        f'  <w:left w:w="{left}" w:type="dxa"/>'
        f'  <w:right w:w="{right}" w:type="dxa"/>'
        f'</w:tcMar>'
    )
    tcPr.append(tcMar)


def _set_table_borders(table, color_hex: str = COLOR_BORDER_HEX, sz: str = "4", val: str = "single") -> None:
    """Aplica bordes institucionales homogéneos a la tabla."""
    tblPr = table._tbl.tblPr
    borders = parse_xml(
        f'<w:tblBorders {nsdecls("w")}>'
        f'  <w:top w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color_hex}"/>'
        f'  <w:left w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color_hex}"/>'
        f'  <w:bottom w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color_hex}"/>'
        f'  <w:right w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color_hex}"/>'
        f'  <w:insideH w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color_hex}"/>'
        f'  <w:insideV w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color_hex}"/>'
        f'</w:tblBorders>'
    )
    tblPr.append(borders)


def _set_row_cant_split(row) -> None:
    """Evita que una fila se divida entre saltos de página."""
    trPr = row._tr.get_or_add_trPr()
    trPr.append(parse_xml(f'<w:cantSplit {nsdecls("w")}/>'))


def _set_row_header(row) -> None:
    """Marca la fila como encabezado repetible en saltos de página y evita partición."""
    trPr = row._tr.get_or_add_trPr()
    trPr.append(parse_xml(f'<w:tblHeader {nsdecls("w")}/>'))
    trPr.append(parse_xml(f'<w:cantSplit {nsdecls("w")}/>'))


def _set_column_widths(table, widths: Sequence[float]) -> None:
    """Establece el ancho de cada columna en pulgadas en todas las filas."""
    for row in table.rows:
        for idx, width in enumerate(widths):
            if idx < len(row.cells):
                row.cells[idx].width = Inches(width)


# ---------------------------------------------------------------------------
# ADAPTADOR CONCRETO: DocxMethodologicalDocumentRenderer
# ---------------------------------------------------------------------------
class DocxMethodologicalDocumentRenderer(MethodologicalDocumentRendererPort):
    """Adaptador de infraestructura para la renderización de Diseños Metodológicos a DOCX.

    Implementa MethodologicalDocumentRendererPort conforme al diseño de la Fase 29.5.2.
    Construye programáticamente el documento Word sin requerir plantillas binarias externas.
    """

    def render(self, design_dto: MethodologicalDesignDTO, output_path: str) -> str:
        """Renderiza el diseño metodológico a un archivo DOCX y retorna el path canónico.

        Args:
            design_dto: Contrato inmutable de transferencia previamente validado.
            output_path: Ruta de destino en el sistema de archivos.

        Returns:
            Ruta canónica del archivo .docx generado exitosamente.

        Raises:
            DocumentRenderingError: Si design_dto es inválido, incompleto, inconsistente
                                    estructuralmente, o si falla la escritura del archivo.
        """
        # 1. Validaciones estructurales y de contrato
        self._validate_preconditions(design_dto, output_path)

        # 2. Inicializar documento Word
        doc = docx.Document()

        # 3. Configuración de página (Letter Landscape, márgenes institucionales)
        self._setup_page(doc)

        # 4. Renderizar Bloque 1: Título General e Introducción
        self._render_block_1_introduction(doc, design_dto)

        # 5. Renderizar Bloque 3: Tabla 1 — Preguntas Frecuentes (FAQ)
        # Nota: En el corpus institucional (DOC-01..05), la Tabla 1 se ubica inmediatamente
        # después de la Introducción y antes de los Objetivos (FASE_29_5_1 §4.C).
        self._render_block_3_faq(doc, design_dto)

        # 6. Renderizar Bloque 2: Objetivos
        self._render_block_2_objectives(doc, design_dto)

        # 7. Renderizar Bloque 4: Sección III — Tabla 2 — Programa / Agenda
        self._render_block_4_program(doc, design_dto)

        # 8. Renderizar Bloque 5: Sección VI — Tabla 3 — Matriz de Planificación
        self._render_block_5_matrix(doc, design_dto)

        # 9. Guardar y retornar ruta
        return self._save_document(doc, output_path)

    # -----------------------------------------------------------------------
    # MÉTODOS PRIVADOS DE MAQUETACIÓN Y VALIDACIÓN
    # -----------------------------------------------------------------------
    def _validate_preconditions(self, design_dto: MethodologicalDesignDTO, output_path: str) -> None:
        """Verifica que el DTO cumpla las precondiciones estructurales indispensables.

        El renderer no duplica reglas de negocio institucionales, pero debe evitar
        generar documentos físicamente imposibles o corruptos.
        """
        if design_dto is None or not isinstance(design_dto, MethodologicalDesignDTO):
            raise DocumentRenderingError(
                "DTO inválido: se requiere una instancia válida de MethodologicalDesignDTO."
            )

        if not output_path or not isinstance(output_path, str):
            raise DocumentRenderingError(
                "output_path inválido: debe ser una cadena con la ruta de destino no vacía."
            )

        if not design_dto.agenda:
            raise DocumentRenderingError(
                "Inconsistencia estructural: el bloque de agenda no puede estar vacío."
            )

        if not design_dto.operational_matrix:
            raise DocumentRenderingError(
                "Inconsistencia estructural: la matriz operacional no puede estar vacía."
            )

        # Comprobación de coherencia estructural básica (sin inventar reglas de negocio)
        if len(design_dto.agenda) != len(design_dto.operational_matrix):
            raise DocumentRenderingError(
                f"Inconsistencia estructural: la cantidad de bloques en agenda ({len(design_dto.agenda)}) "
                f"no coincide con la matriz operativa ({len(design_dto.operational_matrix)})."
            )

        calculated_minutes = sum(block.minutes for block in design_dto.agenda)
        if calculated_minutes != design_dto.total_minutes:
            raise DocumentRenderingError(
                f"Inconsistencia estructural en tiempos: la suma de bloques de agenda ({calculated_minutes} min.) "
                f"no coincide con total_minutes ({design_dto.total_minutes} min.)."
            )

        for i, (block, act) in enumerate(zip(design_dto.agenda, design_dto.operational_matrix)):
            if block.minutes != act.minutes:
                raise DocumentRenderingError(
                    f"Inconsistencia estructural en tiempos: el bloque {i+1} en agenda ({block.minutes} min.) "
                    f"no coincide con la actividad operativa ({act.minutes} min.)."
                )

    def _setup_page(self, doc: docx.Document) -> None:
        """Configura la orientación Landscape, tamaño Letter y márgenes auditados."""
        section = doc.sections[0]
        section.orientation = WD_ORIENT.LANDSCAPE
        section.page_width = Inches(11.0)
        section.page_height = Inches(8.5)

        # Márgenes auditados (FASE_29_5_1 §4.A): 2.5 cm superior/inferior, 3.0 cm laterales
        section.top_margin = Inches(0.984)     # ~2.5 cm
        section.bottom_margin = Inches(0.984)  # ~2.5 cm
        section.left_margin = Inches(1.181)    # ~3.0 cm
        section.right_margin = Inches(1.181)   # ~3.0 cm

    def _render_block_1_introduction(self, doc: docx.Document, dto: MethodologicalDesignDTO) -> None:
        """Renderiza el título del documento y el Bloque 1 (Introducción)."""
        # Título General Centrado
        title_p = doc.add_paragraph()
        title_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        title_p.paragraph_format.space_after = Pt(12)
        title_run = title_p.add_run(dto.document_title.upper())
        title_run.font.name = FONT_HEADING
        title_run.font.size = PT_TITLE
        title_run.font.bold = True
        title_run.font.color.rgb = COLOR_NAVY_RGB

        # Encabezado "INTRODUCCIÓN"
        intro_h = doc.add_paragraph()
        intro_h.paragraph_format.space_after = Pt(6)
        intro_h_run = intro_h.add_run("INTRODUCCIÓN")
        intro_h_run.font.name = FONT_HEADING
        intro_h_run.font.size = PT_HEADING
        intro_h_run.font.bold = True
        intro_h_run.font.color.rgb = COLOR_NAVY_RGB

        # Párrafo 1: Contexto / Vinculación
        intro_p1 = doc.add_paragraph()
        intro_p1.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        intro_p1.paragraph_format.space_after = Pt(6)
        intro_p1_run = intro_p1.add_run(dto.introduction)
        intro_p1_run.font.name = FONT_BASE
        intro_p1_run.font.size = PT_BASE

        # Párrafo 2: Enfoque Metodológico (si está presente)
        if dto.methodological_approach:
            intro_p2 = doc.add_paragraph()
            intro_p2.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            intro_p2.paragraph_format.space_after = Pt(12)
            intro_p2_run = intro_p2.add_run(dto.methodological_approach)
            intro_p2_run.font.name = FONT_BASE
            intro_p2_run.font.size = PT_BASE

    def _render_block_3_faq(self, doc: docx.Document, dto: MethodologicalDesignDTO) -> None:
        """Renderiza el Bloque 3: Tabla 1 de Preguntas Frecuentes (8 filas × 2 columnas)."""
        table = doc.add_table(rows=8, cols=2)
        _set_table_borders(table)
        _set_column_widths(table, (2.70, 5.94))
        _set_row_header(table.rows[0])

        # Fila 0: Encabezado institucional
        for cell in table.rows[0].cells:
            cell.text = "Preguntas frecuentes"
            _set_cell_shading(cell, COLOR_LIGHT_BLUE_HEX)
            _set_cell_margins(cell)
            for p in cell.paragraphs:
                p.alignment = WD_ALIGN_PARAGRAPH.LEFT
                for r in p.runs:
                    r.font.name = FONT_HEADING
                    r.font.size = PT_TABLE
                    r.font.bold = True
                    r.font.color.rgb = COLOR_NAVY_RGB

        # Filas 1..7: Las 7 preguntas institucionales auditadas
        faq_data = [
            ("¿En qué es esta actividad?", dto.faq.q1_que_es),
            ("¿Para qué se realiza?", dto.faq.q2_para_que or "—"),
            ("¿Cuántas sesiones se realizarán?", dto.faq.q3_sesiones),
            ("¿Cuánto son los protagonistas?", dto.faq.q4_protagonistas),
            ("¿Quién va facilitar??", dto.faq.q5_facilitador),
            ("¿Qué materiales se van a utilizar?", dto.faq.q6_materiales),
            ("¿Cuánto tiempo dura el taller?", dto.faq.q7_duracion),
        ]

        for idx, (question, answer) in enumerate(faq_data, start=1):
            row = table.rows[idx]
            _set_row_cant_split(row)
            q_cell = row.cells[0]
            a_cell = row.cells[1]

            q_cell.text = question
            a_cell.text = answer

            # Sombreado alterno azul hielo en filas impares (FASE_29_5_1 §4.B)
            shd_color = COLOR_ICE_BLUE_HEX if idx % 2 == 1 else None

            for c_i, cell in enumerate((q_cell, a_cell)):
                if shd_color:
                    _set_cell_shading(cell, shd_color)
                _set_cell_margins(cell)
                for p in cell.paragraphs:
                    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
                    for r in p.runs:
                        r.font.size = PT_TABLE
                        if c_i == 0:
                            r.font.name = FONT_HEADING
                            r.font.bold = True
                        else:
                            r.font.name = FONT_BASE

        # Separador posterior
        sep = doc.add_paragraph()
        sep.paragraph_format.space_after = Pt(6)

    def _render_block_2_objectives(self, doc: docx.Document, dto: MethodologicalDesignDTO) -> None:
        """Renderiza el Bloque 2: Objetivos institucionales (viñetas)."""
        obj_h = doc.add_paragraph()
        obj_h.paragraph_format.space_after = Pt(6)
        obj_h_run = obj_h.add_run(dto.objectives_heading)
        obj_h_run.font.name = FONT_HEADING
        obj_h_run.font.size = PT_HEADING
        obj_h_run.font.bold = True
        obj_h_run.font.color.rgb = COLOR_NAVY_RGB

        # Viñeta 1
        obj_p1 = doc.add_paragraph(style="List Bullet")
        obj_p1.paragraph_format.space_after = Pt(3)
        obj_p1_run = obj_p1.add_run(dto.objective_1)
        obj_p1_run.font.name = FONT_BASE
        obj_p1_run.font.size = PT_BASE

        # Viñeta 2
        obj_p2 = doc.add_paragraph(style="List Bullet")
        obj_p2.paragraph_format.space_after = Pt(12)
        obj_p2_run = obj_p2.add_run(dto.objective_2)
        obj_p2_run.font.name = FONT_BASE
        obj_p2_run.font.size = PT_BASE

    def _render_block_4_program(self, doc: docx.Document, dto: MethodologicalDesignDTO) -> None:
        """Renderiza el Bloque 4: Sección III — Tabla 2 — Programa / Agenda."""
        prog_h = doc.add_paragraph()
        prog_h.paragraph_format.space_after = Pt(6)
        prog_h_run = prog_h.add_run(dto.program_section_label)
        prog_h_run.font.name = FONT_HEADING
        prog_h_run.font.size = PT_HEADING
        prog_h_run.font.bold = True
        prog_h_run.font.color.rgb = COLOR_NAVY_RGB

        n_blocks = len(dto.agenda)
        table = doc.add_table(rows=n_blocks + 2, cols=2)
        _set_table_borders(table)
        _set_column_widths(table, (6.64, 2.00))
        _set_row_header(table.rows[0])

        # Fila 0: Encabezado
        table.rows[0].cells[0].text = "Actividad"
        table.rows[0].cells[1].text = "Tiempo"
        for cell in table.rows[0].cells:
            _set_cell_shading(cell, COLOR_LIGHT_BLUE_HEX)
            _set_cell_margins(cell)
            for p in cell.paragraphs:
                for r in p.runs:
                    r.font.name = FONT_HEADING
                    r.font.size = PT_TABLE
                    r.font.bold = True
                    r.font.color.rgb = COLOR_NAVY_RGB

        # Filas 1..N: Bloques de agenda
        for idx, block in enumerate(dto.agenda, start=1):
            row = table.rows[idx]
            _set_row_cant_split(row)
            row.cells[0].text = block.label
            row.cells[1].text = f"{block.minutes} min."

            for c_i, cell in enumerate(row.cells):
                _set_cell_margins(cell)
                for p in cell.paragraphs:
                    if c_i == 1:
                        p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
                    for r in p.runs:
                        r.font.name = FONT_BASE
                        r.font.size = PT_TABLE

        # Fila N+1: Fila Total
        total_row = table.rows[n_blocks + 1]
        _set_row_cant_split(total_row)
        total_row.cells[0].text = "Total"
        total_row.cells[1].text = f"{dto.total_minutes} min."

        for c_i, cell in enumerate(total_row.cells):
            _set_cell_shading(cell, COLOR_ICE_BLUE_HEX)
            _set_cell_margins(cell)
            for p in cell.paragraphs:
                if c_i == 1:
                    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
                for r in p.runs:
                    r.font.name = FONT_HEADING
                    r.font.size = PT_TABLE
                    r.font.bold = True

        # Separador posterior
        sep = doc.add_paragraph()
        sep.paragraph_format.space_after = Pt(6)

    def _render_block_5_matrix(self, doc: docx.Document, dto: MethodologicalDesignDTO) -> None:
        """Renderiza el Bloque 5: Sección VI — Tabla 3 — Matriz de Planificación (6 columnas)."""
        mat_h = doc.add_paragraph()
        mat_h.paragraph_format.space_after = Pt(6)
        mat_h_run = mat_h.add_run(dto.matrix_section_label)
        mat_h_run.font.name = FONT_HEADING
        mat_h_run.font.size = PT_HEADING
        mat_h_run.font.bold = True
        mat_h_run.font.color.rgb = COLOR_NAVY_RGB

        n_ops = len(dto.operational_matrix)
        table = doc.add_table(rows=n_ops + 1, cols=6)
        _set_table_borders(table)
        _set_column_widths(table, (0.50, 1.40, 1.74, 2.80, 1.30, 0.90))
        _set_row_header(table.rows[0])

        # Fila 0: Encabezados de las 6 columnas oficiales (FASE_29_5_1 §4.C)
        col_headers = [
            "No.",
            "Actividad",
            "Objetivos Operativos",
            "Procedimiento",
            "Materiales",
            "Tiempo",
        ]
        for c_i, h_text in enumerate(col_headers):
            cell = table.rows[0].cells[c_i]
            cell.text = h_text
            _set_cell_shading(cell, COLOR_LIGHT_BLUE_HEX)
            _set_cell_margins(cell)
            for p in cell.paragraphs:
                if c_i in (0, 5):
                    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                for r in p.runs:
                    r.font.name = FONT_HEADING
                    r.font.size = PT_TABLE
                    r.font.bold = True
                    r.font.color.rgb = COLOR_NAVY_RGB

        # Filas 1..M: Pasos operacionales
        for idx, act in enumerate(dto.operational_matrix, start=1):
            row = table.rows[idx]
            _set_row_cant_split(row)
            row.cells[0].text = str(act.step_number)
            row.cells[1].text = act.phase_label
            row.cells[2].text = act.operative_goal
            row.cells[3].text = act.procedure
            row.cells[4].text = act.materials
            row.cells[5].text = f"{act.minutes} min."

            for c_i, cell in enumerate(row.cells):
                _set_cell_margins(cell)
                for p in cell.paragraphs:
                    if c_i in (0, 5):
                        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    for r in p.runs:
                        r.font.name = FONT_BASE
                        r.font.size = PT_TABLE

    def _save_document(self, doc: docx.Document, output_path: str) -> str:
        """Guarda el documento Word en el disco de manera segura y controlada."""
        try:
            target = pathlib.Path(output_path).resolve()
            target.parent.mkdir(parents=True, exist_ok=True)
            doc.save(str(target))
            return str(target)
        except Exception as e:
            raise DocumentRenderingError(
                f"Fallo al guardar el documento DOCX en '{output_path}': {e}"
            ) from e
