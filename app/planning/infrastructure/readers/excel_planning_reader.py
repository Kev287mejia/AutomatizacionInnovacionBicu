"""Adaptador de Infraestructura — Lector de Matrices Institucionales POA en Excel.

Fase 29.7 — Implementación Controlada del Planning Reader Institucional.
Implementa PlanningSourceReaderPort conforme a los principios de:
  - Ceguera a Persistencia (cero imports de sqlite3 o SQL).
  - Cero Dependencias de app.word_consolidator (núcleo protegido).
  - Cero Dependencias de IA (OpenRouter, Gemini, Ollama).
  - Cero Fabricación de Datos (los valores desconocidos se preservan como None).
"""
from __future__ import annotations

import datetime
import os
import re
from typing import Any, Optional

import openpyxl

from app.planning.domain.dtos import (
    PlanningSourceActivityDTO,
    PlanningSourceReadResultDTO,
)
from app.planning.domain.ports import PlanningSourceReaderPort


def _normalize_text(val: Any) -> str:
    """Elimina tildes, espacios extras y convierte a minúsculas para matching de encabezados."""
    if val is None:
        return ""
    text = str(val).strip().lower()
    trans = str.maketrans("áéíóúÁÉÍÓÚñÑ", "aeiouAEIOUnN")
    text = text.translate(trans)
    return re.sub(r"\s+", " ", text)


class ExcelPlanningReader(PlanningSourceReaderPort):
    """Lector programático de matrices POA institucionales en formato Excel (.xlsx).

    Soporta los formatos institucionales observados:
      - Formato A: Matriz agregada con metas de participantes (ej. 'Programas, proyectos y act').
      - Formato B: Matriz detallada de actividades (ej. 'Pregrado_Grado en Actividades').
    """

    # Diccionario de sinónimos auditados para mapeo de encabezados institucionales
    _HEADER_SYNONYMS: dict[str, set[str]] = {
        "raw_no": {"no", "n°", "numero", "num"},
        "planning_id": {"planning_id", "id_actividad", "codigo_actividad", "id"},
        "activity_name": {"actividad", "nombre_actividad", "nombre de la actividad"},
        "sede": {"sede", "nombre de la sede", "recinto"},
        "dep_sede": {"dep_sede", "departamento_sede", "departamento de ubicacion sede"},
        "departamento": {"departamento"},
        "mun_sede": {"mun_sede", "municipio_sede", "municipio de ubicacion sede"},
        "municipio": {"municipio"},
        "programa": {"programa", "nombre del programa"},
        "otro_programa": {"otro_programa", "otro programa"},
        "proyecto": {"proyecto", "nombre del proyecto"},
        "tipo_proyecto": {"tipo_proyecto", "tipo de proyecto"},
        "ambito": {"ambito", "ámbito"},
        "eje_estrategia": {"eje_estrategia", "eje", "linea estrategica", "línea estratégica"},
        "codigo_presupuestario": {"area", "área", "codigo_presupuestario", "código presupuestario"},
        "area_responsable": {"area_responsable", "área_responsable", "nombre_area", "nombre_área", "instancia_responsable"},
        "departamento_responsable": {"departamento_responsable"},
        "tipo_evento": {"tipo_evento", "evento"},
        "proposito": {"proposito", "propósito", "objetivo"},
        "fecha_evento": {"fecha_evento", "fecha"},
        "convenio": {"convenio", "nombre_convenio"},
        "entidades_cooperantes": {"entidad_vinculada_1", "nombre_entidad_1", "entidades_cooperantes"},
        # Metas cuantitativas
        "est_grado_m": {"total, estud_m_grado", "estud_m_grado", "est_grado_m"},
        "est_grado_f": {"total, estud_f_grado", "estud_f_grado", "est_grado_f"},
        "est_postgrado_m": {"total, estud_m_posgrado", "estud_m_posgrado", "est_postgrado_m"},
        "est_postgrado_f": {"total, estud_f_posgrado", "estud_f_posgrado", "est_postgrado_f"},
        "docentes_m": {"total, docente_m", "docente_m", "docentes_m"},
        "docentes_f": {"total, docente_f", "docente_f", "docentes_f"},
        "administrativos_m": {"total, administrativos_m", "administrativos_m"},
        "administrativos_f": {"total, administrativos_f", "administrativos_f"},
        # Componentes externos M
        "ext_pub_m": {"total, partic_m_inst_publica", "total, partic_m_inst_pública"},
        "ext_priv_m": {"total, partic_m_inst_privada"},
        "ext_ong_m": {"total, partic_m_ong"},
        "ext_prot_m": {"total, protagonistas_m"},
        # Componentes externos F
        "ext_pub_f": {"total, partic_f_inst_publica", "total, partic_f_inst_pública"},
        "ext_priv_f": {"total, partic_f_inst_privada"},
        "ext_ong_f": {"total, partic_f_ong"},
        "ext_prot_f": {"total, protagonistas_f"},
    }

    def read_planning_source(
        self,
        file_path: str,
        sheet_name: Optional[str] = None,
    ) -> PlanningSourceReadResultDTO:
        """Lee un archivo institucional de planificación y retorna los datos normalizados técnicamente.

        Args:
            file_path: Ruta del archivo Excel.
            sheet_name: Nombre de la hoja opcional a procesar.

        Returns:
            PlanningSourceReadResultDTO con las actividades leídas y cualquier advertencia/error.
        """
        if not os.path.isfile(file_path):
            return PlanningSourceReadResultDTO(
                source_file=file_path,
                sheet_name=sheet_name or "",
                activities=(),
                structural_errors=(f"El archivo '{file_path}' no existe o no es un archivo válido.",),
            )

        try:
            wb = openpyxl.load_workbook(file_path, data_only=True, read_only=True)
        except Exception as e:
            return PlanningSourceReadResultDTO(
                source_file=file_path,
                sheet_name=sheet_name or "",
                activities=(),
                structural_errors=(f"Error al abrir el archivo Excel '{file_path}': {e}",),
            )

        try:
            # 1. Identificar la hoja relevante
            target_sheet_name, sheet_err = self._resolve_sheet_name(wb, sheet_name)
            if sheet_err:
                return PlanningSourceReadResultDTO(
                    source_file=file_path,
                    sheet_name=sheet_name or "",
                    activities=(),
                    structural_errors=(sheet_err,),
                )

            ws = wb[target_sheet_name]

            # 2. Localizar la fila de encabezados
            header_row_idx, col_mapping, header_err = self._locate_headers(ws)
            if header_err:
                return PlanningSourceReadResultDTO(
                    source_file=file_path,
                    sheet_name=target_sheet_name,
                    activities=(),
                    structural_errors=(header_err,),
                )

            # 3. Leer filas de datos
            activities, warnings = self._read_data_rows(
                ws=ws,
                file_path=file_path,
                sheet_name=target_sheet_name,
                header_row_idx=header_row_idx,
                col_mapping=col_mapping,
            )

            return PlanningSourceReadResultDTO(
                source_file=file_path,
                sheet_name=target_sheet_name,
                activities=tuple(activities),
                structural_errors=(),
                reading_warnings=tuple(warnings),
            )
        finally:
            wb.close()

    def _resolve_sheet_name(
        self,
        wb: openpyxl.Workbook,
        requested_sheet: Optional[str],
    ) -> tuple[str, Optional[str]]:
        """Determina la hoja a procesar."""
        if requested_sheet:
            if requested_sheet in wb.sheetnames:
                return requested_sheet, None
            return "", f"La hoja solicitada '{requested_sheet}' no existe en el archivo. Hojas: {wb.sheetnames}"

        # Candidatos conocidos en orden de prioridad
        priority_candidates = [
            "Programas, proyectos y act",
            "Pregrado_Grado en Actividades",
        ]
        for candidate in priority_candidates:
            if candidate in wb.sheetnames:
                return candidate, None

        # Si ninguna coincide exactamente, inspeccionar si alguna hoja tiene encabezados de planificación
        for sname in wb.sheetnames:
            ws = wb[sname]
            for r in range(1, min(6, ws.max_row + 1)):
                row_vals = [_normalize_text(ws.cell(r, c).value) for c in range(1, min(30, ws.max_column + 1))]
                if any("actividad" in v for v in row_vals) and any("sede" in v for v in row_vals):
                    return sname, None

        return "", f"No se encontró una hoja válida de planificación en el archivo Excel. Hojas disponibles: {wb.sheetnames}"

    def _locate_headers(
        self,
        ws: Any,
    ) -> tuple[int, dict[str, int], Optional[str]]:
        """Localiza la fila de encabezados y construye el mapeo columna -> índice."""
        for r in range(1, min(10, ws.max_row + 1)):
            row_vals = [_normalize_text(ws.cell(r, c).value) for c in range(1, ws.max_column + 1)]
            has_act = any("actividad" in v for v in row_vals)
            has_sede = any("sede" in v for v in row_vals)
            if has_act and has_sede:
                # Fila de encabezado localizada
                mapping: dict[str, int] = {}
                for col_idx, raw_val in enumerate(row_vals, start=1):
                    if not raw_val:
                        continue
                    for field_name, synonyms in self._HEADER_SYNONYMS.items():
                        if field_name not in mapping and raw_val in synonyms:
                            mapping[field_name] = col_idx

                # Verificar obligatoriedad mínima de columnas para poder interpretar
                if "activity_name" not in mapping:
                    return r, {}, "No se localizó la columna obligatoria 'Actividad' en la fila de encabezados."
                if "sede" not in mapping:
                    return r, {}, "No se localizó la columna obligatoria 'Sede' en la fila de encabezados."

                return r, mapping, None

        return 0, {}, "No se encontraron los encabezados institucionales mínimos ('Sede', 'Actividad') en las primeras filas."

    def _read_data_rows(
        self,
        ws: Any,
        file_path: str,
        sheet_name: str,
        header_row_idx: int,
        col_mapping: dict[str, int],
    ) -> tuple[list[PlanningSourceActivityDTO], list[str]]:
        """Lee y extrae las actividades de las filas subsiguientes a la fila de encabezados."""
        activities: list[PlanningSourceActivityDTO] = []
        warnings: list[str] = []

        def get_val(r: int, field_key: str) -> Any:
            c = col_mapping.get(field_key)
            if c is None:
                return None
            return ws.cell(r, c).value

        def get_clean_str(r: int, field_key: str) -> Optional[str]:
            v = get_val(r, field_key)
            if v is None:
                return None
            s = str(v).strip()
            return s if s else None

        def get_int(r: int, field_key: str) -> Optional[int]:
            v = get_val(r, field_key)
            if v is None or v == "":
                return None
            try:
                return int(v)
            except (ValueError, TypeError):
                try:
                    return int(float(v))
                except (ValueError, TypeError):
                    warnings.append(
                        f"Fila {r}: El valor '{v}' en '{field_key}' no es un entero válido."
                    )
                    return None

        def get_date_str(r: int, field_key: str) -> Optional[str]:
            v = get_val(r, field_key)
            if v is None:
                return None
            if isinstance(v, (datetime.date, datetime.datetime)):
                return v.strftime("%Y-%m-%d")
            s = str(v).strip()
            return s if s else None

        for r in range(header_row_idx + 1, ws.max_row + 1):
            act_name = get_clean_str(r, "activity_name")
            sede_val = get_clean_str(r, "sede")

            # Si ambos están vacíos, omitir fila (fila en blanco o de formato)
            if not act_name and not sede_val:
                continue

            row_warnings: list[str] = []

            # Extraer metas numéricas
            g_est_grado_m = get_int(r, "est_grado_m")
            g_est_grado_f = get_int(r, "est_grado_f")
            g_est_pos_m = get_int(r, "est_postgrado_m")
            g_est_pos_f = get_int(r, "est_postgrado_f")
            g_doc_m = get_int(r, "docentes_m")
            g_doc_f = get_int(r, "docentes_f")
            g_adm_m = get_int(r, "administrativos_m")
            g_adm_f = get_int(r, "administrativos_f")

            # Sumar componentes de externos si existen en la hoja
            ext_m_parts = [get_int(r, k) for k in ("ext_pub_m", "ext_priv_m", "ext_ong_m", "ext_prot_m")]
            ext_m_parts_valid = [p for p in ext_m_parts if p is not None]
            externos_m = sum(ext_m_parts_valid) if ext_m_parts_valid else None

            ext_f_parts = [get_int(r, k) for k in ("ext_pub_f", "ext_priv_f", "ext_ong_f", "ext_prot_f")]
            ext_f_parts_valid = [p for p in ext_f_parts if p is not None]
            externos_f = sum(ext_f_parts_valid) if ext_f_parts_valid else None

            # Validar enteros no negativos en metas
            for g_name, g_val in [
                ("est_grado_m", g_est_grado_m),
                ("est_grado_f", g_est_grado_f),
                ("est_postgrado_m", g_est_pos_m),
                ("est_postgrado_f", g_est_pos_f),
                ("docentes_m", g_doc_m),
                ("docentes_f", g_doc_f),
                ("administrativos_m", g_adm_m),
                ("administrativos_f", g_adm_f),
                ("externos_m", externos_m),
                ("externos_f", externos_f),
            ]:
                if g_val is not None and g_val < 0:
                    row_warnings.append(f"La meta '{g_name}' no puede ser negativa ({g_val}).")

            dep_sede_val = get_clean_str(r, "dep_sede") or get_clean_str(r, "departamento")
            mun_sede_val = get_clean_str(r, "mun_sede") or get_clean_str(r, "municipio")

            dto = PlanningSourceActivityDTO(
                source_file=file_path,
                sheet_name=sheet_name,
                row_number=r,
                raw_no=get_clean_str(r, "raw_no"),
                planning_id=get_clean_str(r, "planning_id"),
                activity_name=act_name,
                sede=sede_val,
                dep_sede=dep_sede_val,
                mun_sede=mun_sede_val,
                programa=get_clean_str(r, "programa"),
                otro_programa=get_clean_str(r, "otro_programa"),
                proyecto=get_clean_str(r, "proyecto"),
                tipo_proyecto=get_clean_str(r, "tipo_proyecto"),
                ambito=get_clean_str(r, "ambito"),
                eje_estrategia=get_clean_str(r, "eje_estrategia"),
                codigo_presupuestario=get_clean_str(r, "codigo_presupuestario"),
                area_responsable=get_clean_str(r, "area_responsable"),
                departamento_responsable=get_clean_str(r, "departamento_responsable"),
                tipo_evento=get_clean_str(r, "tipo_evento"),
                proposito=get_clean_str(r, "proposito"),
                fecha_evento=get_date_str(r, "fecha_evento"),
                convenio=get_clean_str(r, "convenio"),
                entidades_cooperantes=get_clean_str(r, "entidades_cooperantes"),
                est_grado_m=g_est_grado_m,
                est_grado_f=g_est_grado_f,
                est_postgrado_m=g_est_pos_m,
                est_postgrado_f=g_est_pos_f,
                docentes_m=g_doc_m,
                docentes_f=g_doc_f,
                administrativos_m=g_adm_m,
                administrativos_f=g_adm_f,
                externos_m=externos_m,
                externos_f=externos_f,
                reading_warnings=tuple(row_warnings),
            )
            activities.append(dto)

        return activities, warnings
