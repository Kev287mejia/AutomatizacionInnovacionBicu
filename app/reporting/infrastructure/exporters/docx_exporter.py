"""app.reporting.infrastructure.exporters.docx_exporter

Exportador de Reportes Institucionales BICU en formato Word (.docx).
Fase 29.20.3 — Integración Controlada de Reporting + Dashboard.

PRINCIPIOS:
  - Generación de archivos derivados limpios utilizando python-docx.
  - Cero modificación de plantillas maestras o documentos patrimoniales.
  - Jerarquía visual institucional con tablas formateadas.
  - Preservación íntegra de metadatos periciales y notas de auditoría.
"""

from pathlib import Path
from typing import Union

import docx
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

from app.reporting.domain.dto import ReportDocumentDTO


class DOCXReportExporter:
    """Genera documentos Word institucionales derivados a partir de ReportDocumentDTO."""

    @staticmethod
    def export(document: ReportDocumentDTO, output_path: Union[str, Path]) -> Path:
        """Exporta el informe a la ruta especificada en formato .docx."""
        out = Path(output_path)
        out.parent.mkdir(parents=True, exist_ok=True)

        doc = docx.Document()

        # Configuración de márgenes estándar (1 pulgada)
        for s in doc.sections:
            s.top_margin = Inches(1.0)
            s.bottom_margin = Inches(1.0)
            s.left_margin = Inches(1.0)
            s.right_margin = Inches(1.0)

        # -----------------------------------------------------------------
        # 1. ENCABEZADO INSTITUCIONAL
        # -----------------------------------------------------------------
        p_inst = doc.add_paragraph()
        p_inst.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r_inst = p_inst.add_run(document.institution_name.upper())
        r_inst.font.name = "Segoe UI"
        r_inst.font.size = Pt(14)
        r_inst.font.bold = True
        r_inst.font.color.rgb = RGBColor(11, 60, 93)  # #0B3C5D

        p_title = doc.add_paragraph()
        p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r_title = p_title.add_run(f"{document.report_type.value}: {document.title}")
        r_title.font.name = "Segoe UI"
        r_title.font.size = Pt(12)
        r_title.font.bold = True
        r_title.font.color.rgb = RGBColor(26, 38, 52)  # #1A2634

        if document.subtitle:
            p_sub = doc.add_paragraph()
            p_sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
            r_sub = p_sub.add_run(document.subtitle)
            r_sub.font.name = "Segoe UI"
            r_sub.font.size = Pt(10)
            r_sub.font.italic = True
            r_sub.font.color.rgb = RGBColor(100, 100, 100)

        doc.add_paragraph()  # Espacio

        # -----------------------------------------------------------------
        # 2. CUADRO DE METADATOS Y FILTROS APLICADOS
        # -----------------------------------------------------------------
        meta_table = doc.add_table(rows=0, cols=2)
        meta_table.alignment = WD_TABLE_ALIGNMENT.CENTER

        meta_rows = [
            ("ID del Reporte:", document.report_id),
            ("Tipo Institucional:", document.report_type.value),
            ("Fecha de Emisión:", document.generated_at),
            ("Emitido Por:", document.generated_by),
            ("Año del Período:", str(document.filters_applied.period_year or "Todos los años")),
            ("Rango de Fechas:", f"{document.filters_applied.start_date or 'Inicio'} a {document.filters_applied.end_date or 'Fin'}"),
            ("Sede Institucional:", document.filters_applied.sede or "Todas las Sedes"),
            ("Modo Multisesión:", document.filters_applied.multisession_goal_mode),
            ("Sello Criptográfico:", document.cryptographic_seal or "No computado"),
        ]

        for k, v in meta_rows:
            row_cells = meta_table.add_row().cells
            row_cells[0].width = Inches(2.2)
            row_cells[1].width = Inches(4.3)
            p0 = row_cells[0].paragraphs[0]
            r0 = p0.add_run(k)
            r0.font.name = "Segoe UI"
            r0.font.size = Pt(9.5)
            r0.font.bold = True

            p1 = row_cells[1].paragraphs[0]
            r1 = p1.add_run(str(v))
            r1.font.name = "Segoe UI"
            r1.font.size = Pt(9.5)

        doc.add_paragraph()  # Espacio

        # -----------------------------------------------------------------
        # 3. SECCIONES Y TABLAS DEL DOCUMENTO
        # -----------------------------------------------------------------
        for sec in document.sections:
            # Título de la sección
            h_sec = doc.add_heading(level=2)
            r_h = h_sec.add_run(f"{sec.section_id}: {sec.title}")
            r_h.font.name = "Segoe UI"
            r_h.font.size = Pt(11.5)
            r_h.font.bold = True
            r_h.font.color.rgb = RGBColor(11, 60, 93)

            if sec.description:
                p_desc = doc.add_paragraph()
                r_d = p_desc.add_run(sec.description)
                r_d.font.name = "Segoe UI"
                r_d.font.size = Pt(9.5)
                r_d.font.italic = True

            # Métricas si existen
            if sec.metrics:
                m_table = doc.add_table(rows=1, cols=5)
                m_table.alignment = WD_TABLE_ALIGNMENT.CENTER
                m_headers = ["Indicador", "Concepto", "Valor", "Unidad", "Estado"]
                hdr_cells = m_table.rows[0].cells
                for idx, h_text in enumerate(m_headers):
                    hdr_cells[idx].paragraphs[0].text = h_text
                    hdr_cells[idx].paragraphs[0].runs[0].font.bold = True
                    hdr_cells[idx].paragraphs[0].runs[0].font.name = "Segoe UI"
                    hdr_cells[idx].paragraphs[0].runs[0].font.size = Pt(9)
                    hdr_cells[idx].paragraphs[0].runs[0].font.color.rgb = RGBColor(255, 255, 255)
                    DOCXReportExporter._set_cell_background(hdr_cells[idx], "0B3C5D")

                for m in sec.metrics:
                    row_cells = m_table.add_row().cells
                    row_cells[0].paragraphs[0].text = m.indicator_id
                    row_cells[1].paragraphs[0].text = m.label
                    row_cells[2].paragraphs[0].text = m.value_display
                    row_cells[3].paragraphs[0].text = m.unit
                    row_cells[4].paragraphs[0].text = m.alert_level
                    for c in row_cells:
                        if c.paragraphs[0].runs:
                            c.paragraphs[0].runs[0].font.name = "Segoe UI"
                            c.paragraphs[0].runs[0].font.size = Pt(8.5)

                doc.add_paragraph()

            # Tablas de datos si existen
            if sec.table_headers and sec.table_rows:
                t_table = doc.add_table(rows=1, cols=len(sec.table_headers))
                t_table.alignment = WD_TABLE_ALIGNMENT.CENTER
                hdr_cells = t_table.rows[0].cells
                for idx, h_text in enumerate(sec.table_headers):
                    hdr_cells[idx].paragraphs[0].text = str(h_text)
                    hdr_cells[idx].paragraphs[0].runs[0].font.bold = True
                    hdr_cells[idx].paragraphs[0].runs[0].font.name = "Segoe UI"
                    hdr_cells[idx].paragraphs[0].runs[0].font.size = Pt(9)
                    hdr_cells[idx].paragraphs[0].runs[0].font.color.rgb = RGBColor(255, 255, 255)
                    DOCXReportExporter._set_cell_background(hdr_cells[idx], "1A2634")

                for r_data in sec.table_rows:
                    row_cells = t_table.add_row().cells
                    for idx, val in enumerate(r_data):
                        val_str = str(val) if val is not None else "N/D"
                        row_cells[idx].paragraphs[0].text = val_str
                        if row_cells[idx].paragraphs[0].runs:
                            row_cells[idx].paragraphs[0].runs[0].font.name = "Segoe UI"
                            row_cells[idx].paragraphs[0].runs[0].font.size = Pt(8.5)

                doc.add_paragraph()

            # Notas de auditoría
            if sec.audit_notes:
                for note in sec.audit_notes:
                    p_note = doc.add_paragraph()
                    r_note = p_note.add_run(f"• Nota Pericial: {note}")
                    r_note.font.name = "Segoe UI"
                    r_note.font.size = Pt(8.5)
                    r_note.font.italic = True
                    r_note.font.color.rgb = RGBColor(80, 80, 80)

        doc.save(str(out))
        return out

    @staticmethod
    def _set_cell_background(cell: Any, hex_color: str) -> None:
        """Aplica color de fondo hexadecimal al XML de una celda de tabla."""
        tc_pr = cell._tc.get_or_add_tcPr()
        shd = OxmlElement("w:shd")
        shd.set(qn("w:val"), "clear")
        shd.set(qn("w:color"), "auto")
        shd.set(qn("w:fill"), hex_color)
        tc_pr.append(shd)
