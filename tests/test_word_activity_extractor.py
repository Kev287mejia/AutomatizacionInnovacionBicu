"""tests.test_word_activity_extractor

Suite de pruebas unitarias y de integración para Fase 26.9 — Bloque 1:
Extractor de actividades desde documentos Word institucionales (.docx).

Cobertura requerida según especificación:
    - T-01: Documento válido con Tabla 1 (Ficha Técnica detectada).
    - T-02: Extracción correcta de campos de Tabla 1 (Actividad, Línea, Indicador, Lugar, Fecha, Resultados, etc.).
    - T-03: Documento válido con Tabla 2 (Matriz Cuantitativa detectada).
    - T-04: Extracción correcta de matriz cuantitativa (F, M, Total, Estamentos, Etnias).
    - T-05: Tabla ausente (manejo controlado sin excepción no capturada).
    - T-06: Campo vacío (manejo controlado, valores None, no inventa datos).
    - T-07: Documento no compatible (detecta NO_COMPATIBLE, compatible_flujo_principal=False).
    - T-08: DOCX corrupto o ilegible (captura error, reporta en diagnostico.errores).
    - T-09: Lista de asistencia como imagen no convertida artificialmente en participantes (TOTAL != NOMINAL).
    - T-10: Conservación de advertencias de extracción (discrepancia aritmética F + M != Total).
    - T-11: No dependencia de SQLite (extractor opera sin base de datos).
    - T-12: No dependencia de Quality / Routing / Export (extractor desacoplado de fases posteriores).
    - T-13: Aislamiento arquitectónico (AST: Application no importa docx; extractor no importa sqlite3 ni openpyxl).
    - Adicionalmente: Validación contra archivos reales de la institución si están disponibles en el entorno.
"""

import ast
import glob
import inspect
import io
import os
from pathlib import Path
from typing import List, Optional
import docx
import pytest

from app.application.dto.word_extraction_dtos import (
    DiagnosticoExtraccionDTO,
    FichaTecnicaDTO,
    MatrizCuantitativaDTO,
    TipoDocumentoWord,
    TipoEvidenciaWord,
    WordActivityExtractionResultDTO,
)
from app.application.ports.word_activity_extractor import (
    IWordActivityExtractor,
)
from app.infrastructure.word_reader.word_activity_extractor import (
    WordActivityExtractor,
)


# ---------------------------------------------------------------------------
# Helpers para crear documentos DOCX sintéticos en memoria / archivos temporales
# ---------------------------------------------------------------------------

def _crear_docx_actividad_valido(
    tmp_path: Path,
    nombre_archivo: str = "actividad_valida.docx",
    actividad: str = "Taller de Capacitación en Formularios Digitales",
    linea: str = "Consolidar el modelo educativo comunitario",
    indicador: str = "16. Porcentaje de estudiantes capacitados",
    lugar: str = "Laboratorio BICU Bilwi",
    fecha: str = "06/06/2026",
    total: int = 15,
    femenino: int = 7,
    masculino: int = 8,
    estudiantes: int = 15,
    docentes: int = 0,
    admo: int = 0,
    otros: int = 0,
    incluir_tabla1: bool = True,
    incluir_tabla2: bool = True,
    corromper_suma_genero: bool = False,
    campo_vacio: bool = False,
) -> Path:
    """Crea un documento .docx sintético con la estructura exacta de las plantillas BICU."""
    doc = docx.Document()

    doc.add_heading("Bluefields Indian & Caribbean University", level=1)
    doc.add_heading("Resumen Técnico", level=2)

    doc.add_heading("Objetivo de la Actividad", level=3)
    doc.add_paragraph("2.1. Fortalecer competencias digitales en recolección de datos.")
    doc.add_paragraph("2.2. Capacitar en el diseño visual de reportes.")

    doc.add_heading("Diseño Metodológico", level=3)
    doc.add_paragraph("Metodología participativa de laboratorio práctico.")

    doc.add_heading("Desarrollo de la Actividad", level=3)
    doc.add_paragraph("La actividad se desarrolló con asistencia completa en fecha programada.")

    # Tabla 1: Ficha Técnica (19 filas x 2 cols)
    if incluir_tabla1:
        t1 = doc.add_table(rows=19, cols=2)
        filas_t1 = [
            ("Actividad General", "" if campo_vacio else actividad),
            ("Línea Estratégica de la Actividad", linea),
            ("3. Indicador al que aporta", indicador),
            ("Lugar", lugar),
            ("Fecha de Realización", fecha),
            ("Participantes", f"Total: [{total}]"),
            ("Participantes", "Distribución por Género"),
            ("Participantes", f"Masculino: [{masculino}]"),
            ("Participantes", f"Femenino: [{femenino}]"),
            ("Participantes", "Distribución por Etnia"),
            ("Participantes", "Creole [1]   Mestizos [4] Miskito [10] Garífuna [ 0] Rama [0]    Mayangna [0] Ulwa[ 0]"),
            ("Principales Resultados", "Resultado 1: Estudiantes capacitados."),
            ("Principales Resultados", "Resultado 2: Formatos estandarizados."),
            ("Principales Resultados", "Resultado 3: Eficiencia en recolección."),
            ("8. Principales Dificultades", "Dificultad 1: Conectividad intermitente."),
            ("8. Principales Dificultades", "Dificultad 2: Sin impresora local."),
            ("9. Acuerdos de la acción", "Acuerdo 1: Aplicación en investigaciones."),
            ("9. Acuerdos de la acción", "Acuerdo 2: Estandarización de plantillas."),
            ("9. Acuerdos de la acción", "7.3 Seguimiento mensual."),
        ]
        for idx, (lbl, val) in enumerate(filas_t1):
            t1.rows[idx].cells[0].text = lbl
            t1.rows[idx].cells[1].text = val

    # Tabla 2: Matriz Cuantitativa (3 filas x 15 cols)
    if incluir_tabla2:
        t2 = doc.add_table(rows=3, cols=15)
        # Fila 1: Supercabeceras
        for c in t2.rows[0].cells:
            c.text = "Encabezado"
        # Fila 2: Subcabeceras
        headers = [
            "Actividad realizada", "F", "M", "Total",
            "Estudiante", "Docente", "Trab. Admo", "Otros",
            "Creole", "Mzt", "Miskito", "Rama", "Ulwa", "Mayagna", "Garífuna"
        ]
        for c_idx, h in enumerate(headers):
            t2.rows[1].cells[c_idx].text = h

        # Fila 3: Valores
        val_fem = femenino
        val_masc = masculino
        val_tot = (femenino + masculino + 10) if corromper_suma_genero else total

        vals = [
            actividad, str(val_fem), str(val_masc), str(val_tot),
            str(estudiantes), str(docentes), str(admo), str(otros),
            "1", "4", "10", "0", "0", "0", "0"
        ]
        for c_idx, v in enumerate(vals):
            t2.rows[2].cells[c_idx].text = v

    doc.add_heading("Anexos", level=2)
    p_anexo2 = doc.add_paragraph("Anexo 2. Lista de Asistencia.")
    # Simular una imagen en Anexo 2 agregando un run
    p_anexo2.add_run(" (Imagen adjunta de lista firmada)")

    p_anexo3 = doc.add_paragraph("Anexo 3. Fotografía")
    p_anexo3.add_run(" (Foto 1 del evento)")

    fpath = tmp_path / nombre_archivo
    doc.save(str(fpath))
    return fpath


# ===========================================================================
# SUITE DE PRUEBAS DEL EXTRACTOR
# ===========================================================================

class TestWordActivityExtractor:
    """Suite de pruebas para el extractor físico de actividades Word."""

    @pytest.fixture
    def extractor(self) -> WordActivityExtractor:
        return WordActivityExtractor()

    # -----------------------------------------------------------------------
    # T-01: Documento válido con Tabla 1
    # -----------------------------------------------------------------------
    def test_t01_documento_valido_con_tabla_1_detecta_ficha_tecnica(self, extractor, tmp_path):
        """T-01: Documento válido con Tabla 1 es reconocido como INFORME_ACTIVIDAD."""
        doc_path = _crear_docx_actividad_valido(tmp_path, "t01_doc.docx", incluir_tabla1=True, incluir_tabla2=False)
        res = extractor.extract(doc_path)

        assert res.diagnostico.es_valido is True
        assert res.diagnostico.tipo_documento == TipoDocumentoWord.INFORME_ACTIVIDAD
        assert res.diagnostico.compatible_flujo_principal is True
        assert res.diagnostico.tabla_1_encontrada is True
        assert res.ficha_tecnica is not None
        assert res.nombre_archivo == "t01_doc.docx"
        assert res.hash_sha256 is not None
        assert len(res.hash_sha256) == 64

    # -----------------------------------------------------------------------
    # T-02: Extracción correcta de campos de Tabla 1
    # -----------------------------------------------------------------------
    def test_t02_extraccion_correcta_campos_tabla_1(self, extractor, tmp_path):
        """T-02: Se extraen fielmente Actividad, Línea, Indicador, Lugar, Fecha y listas."""
        doc_path = _crear_docx_actividad_valido(
            tmp_path,
            "t02_doc.docx",
            actividad="Taller de Capacitación en Google Forms",
            linea="Eje 11: Innovación",
            indicador="16. Estudiantes capacitados",
            lugar="Laboratorio Bilwi",
            fecha="06/06/2026",
            total=15,
            femenino=7,
            masculino=8,
        )
        res = extractor.extract(doc_path)
        ft = res.ficha_tecnica

        assert ft is not None
        assert ft.actividad_general == "Taller de Capacitación en Google Forms"
        assert ft.linea_estrategica == "Eje 11: Innovación"
        assert ft.indicador == "16. Estudiantes capacitados"
        assert ft.lugar == "Laboratorio Bilwi"
        assert ft.fecha_realizacion == "06/06/2026"
        assert ft.fecha_iso == "2026-06-06"
        assert ft.total_participantes_declarado == 15
        assert ft.masculino_declarado == 8
        assert ft.femenino_declarado == 7
        assert len(ft.resultados) == 3
        assert len(ft.dificultades) == 2
        assert len(ft.acuerdos) == 3
        assert ft.distribucion_etnica_declarada.get("Miskito") == 10
        assert ft.distribucion_etnica_declarada.get("Mestizo") == 4
        assert ft.distribucion_etnica_declarada.get("Creole") == 1

    # -----------------------------------------------------------------------
    # T-03: Documento válido con Tabla 2
    # -----------------------------------------------------------------------
    def test_t03_documento_valido_con_tabla_2_detecta_matriz_cuantitativa(self, extractor, tmp_path):
        """T-03: Se reconoce la presencia de la Matriz Cuantitativa (Tabla 2)."""
        doc_path = _crear_docx_actividad_valido(tmp_path, "t03_doc.docx", incluir_tabla1=False, incluir_tabla2=True)
        res = extractor.extract(doc_path)

        assert res.diagnostico.es_valido is True
        assert res.diagnostico.tipo_documento == TipoDocumentoWord.INFORME_ACTIVIDAD
        assert res.diagnostico.tabla_2_encontrada is True
        assert res.matriz_cuantitativa is not None

    # -----------------------------------------------------------------------
    # T-04: Extracción correcta de la matriz cuantitativa
    # -----------------------------------------------------------------------
    def test_t04_extraccion_correcta_matriz_cuantitativa(self, extractor, tmp_path):
        """T-04: Cifras de género, estamentos y etnias se extraen como enteros exactos."""
        doc_path = _crear_docx_actividad_valido(
            tmp_path,
            "t04_doc.docx",
            total=24,
            femenino=12,
            masculino=12,
            estudiantes=9,
            docentes=10,
            admo=5,
            otros=0,
        )
        res = extractor.extract(doc_path)
        mc = res.matriz_cuantitativa

        assert mc is not None
        assert mc.total == 24
        assert mc.femenino == 12
        assert mc.masculino == 12
        assert mc.estudiantes == 9
        assert mc.docentes == 10
        assert mc.trabajadores_administrativos == 5
        assert mc.otros == 0
        assert mc.distribucion_etnica.get("Miskito") == 10
        assert mc.discrepancia_suma_genero is False
        assert mc.discrepancia_suma_estamento is False

    # -----------------------------------------------------------------------
    # T-05: Tabla ausente
    # -----------------------------------------------------------------------
    def test_t05_tabla_ausente_manejo_controlado(self, extractor, tmp_path):
        """T-05: Si no hay tablas o falta alguna, no se lanzan excepciones no capturadas."""
        doc = docx.Document()
        doc.add_paragraph("Texto sin tablas institucionales.")
        p = tmp_path / "sin_tablas.docx"
        doc.save(str(p))

        res = extractor.extract(p)

        assert res.diagnostico.es_valido is False
        assert res.diagnostico.tipo_documento == TipoDocumentoWord.NO_COMPATIBLE
        assert res.diagnostico.compatible_flujo_principal is False
        assert res.diagnostico.tabla_1_encontrada is False
        assert res.diagnostico.tabla_2_encontrada is False
        assert len(res.diagnostico.errores) > 0
        assert res.ficha_tecnica is None
        assert res.matriz_cuantitativa is None

    # -----------------------------------------------------------------------
    # T-06: Campo vacío
    # -----------------------------------------------------------------------
    def test_t06_campo_vacio_no_inventa_datos(self, extractor, tmp_path):
        """T-06: Celda vacía en la Ficha Técnica se mapea a None, sin inventar datos."""
        doc_path = _crear_docx_actividad_valido(tmp_path, "t06_doc.docx", campo_vacio=True)
        res = extractor.extract(doc_path)

        assert res.ficha_tecnica is not None
        assert res.ficha_tecnica.actividad_general is None

    # -----------------------------------------------------------------------
    # T-07: Documento no compatible
    # -----------------------------------------------------------------------
    def test_t07_documento_no_compatible_clasificacion_clara(self, extractor, tmp_path):
        """T-07: Un archivo .docx administrativo o externo se clasifica como NO_COMPATIBLE."""
        doc = docx.Document()
        doc.add_heading("Carta de Solicitud Administrativa", level=1)
        doc.add_paragraph("Por medio de la presente solicito autorización...")
        t = doc.add_table(rows=2, cols=2)
        t.rows[0].cells[0].text = "Item"
        t.rows[0].cells[1].text = "Monto"
        t.rows[1].cells[0].text = "Papelería"
        t.rows[1].cells[1].text = "100"
        p = tmp_path / "carta.docx"
        doc.save(str(p))

        res = extractor.extract(p)

        assert res.diagnostico.tipo_documento == TipoDocumentoWord.NO_COMPATIBLE
        assert res.diagnostico.compatible_flujo_principal is False
        assert len(res.diagnostico.errores) > 0

    # -----------------------------------------------------------------------
    # T-08: DOCX corrupto o ilegible
    # -----------------------------------------------------------------------
    def test_t08_docx_corrupto_o_ilegible_reporta_error(self, extractor, tmp_path):
        """T-08: Archivo binario dañado o no zip se captura controladamente."""
        corrupt_path = tmp_path / "archivo_corrupto.docx"
        with open(corrupt_path, "wb") as f:
            f.write(b"Este no es un archivo ZIP ni DOCX valido.")

        res = extractor.extract(corrupt_path)

        assert res.diagnostico.es_valido is False
        assert res.diagnostico.compatible_flujo_principal is False
        assert len(res.diagnostico.errores) > 0
        assert "corrupto" in res.diagnostico.errores[0].lower() or "formato" in res.diagnostico.errores[0].lower()

    def test_t08b_archivo_inexistente_reporta_error(self, extractor, tmp_path):
        """T-08b: Ruta que no existe devuelve DTO con error sin crashear."""
        inexistente = tmp_path / "no_existe.docx"
        res = extractor.extract(inexistente)

        assert res.diagnostico.es_valido is False
        assert len(res.diagnostico.errores) > 0
        assert "no encontrado" in res.diagnostico.errores[0].lower()

    # -----------------------------------------------------------------------
    # T-09: Lista de asistencia como imagen no genera participantes nominales
    # -----------------------------------------------------------------------
    def test_t09_lista_asistencia_como_imagen_no_genera_personas_nominales(self, extractor, tmp_path):
        """T-09: La imagen del Anexo 2 se detecta como evidencia pero NO genera personas nominales."""
        doc_path = _crear_docx_actividad_valido(tmp_path, "t09_doc.docx")
        res = extractor.extract(doc_path)

        # Se detecta el anexo de asistencia
        assert res.evidencias is not None
        assert res.evidencias.tiene_anexo_asistencia is True

        # El DTO de extracción NO contiene listas de personas ni objetos Person
        assert not hasattr(res, "personas")
        assert not hasattr(res, "participantes_nominales")

    # -----------------------------------------------------------------------
    # T-10: Conservación de advertencias de extracción
    # -----------------------------------------------------------------------
    def test_t10_conservacion_advertencias_discrepancia_suma_genero(self, extractor, tmp_path):
        """T-10: Discrepancia en Tabla 2 (Total != F + M) genera advertencia sin autocorregir."""
        doc_path = _crear_docx_actividad_valido(
            tmp_path,
            "t10_doc.docx",
            femenino=7,
            masculino=8,
            total=15,
            corromper_suma_genero=True,  # Total se grabará como 25
        )
        res = extractor.extract(doc_path)

        assert res.matriz_cuantitativa is not None
        assert res.matriz_cuantitativa.discrepancia_suma_genero is True
        assert res.matriz_cuantitativa.total == 25
        assert res.matriz_cuantitativa.femenino == 7
        assert res.matriz_cuantitativa.masculino == 8
        # Se asentó advertencia en diagnóstico
        assert any("Discrepancia aritmética" in a for a in res.diagnostico.advertencias)

    # -----------------------------------------------------------------------
    # T-11: No dependencia de SQLite
    # -----------------------------------------------------------------------
    def test_t11_no_dependencia_de_sqlite(self):
        """T-11: El módulo extractor no importa sqlite3 ni módulos de base de datos."""
        import app.infrastructure.word_reader.word_activity_extractor as mod
        source = inspect.getsource(mod)
        tree = ast.parse(source)

        imports = []
        for node in ast.walk(tree):
            if isinstance(node, ast.Import):
                for a in node.names:
                    imports.append(a.name)
            elif isinstance(node, ast.ImportFrom):
                if node.module:
                    imports.append(node.module)

        assert not any(i == "sqlite3" or i.startswith("sqlite3.") for i in imports), (
            f"El extractor importa sqlite3 directamente: {imports}"
        )

    # -----------------------------------------------------------------------
    # T-12: No dependencia de Quality / Routing / Export
    # -----------------------------------------------------------------------
    def test_t12_no_dependencia_de_quality_routing_export(self):
        """T-12: El extractor no depende de los motores de fases posteriores."""
        import app.infrastructure.word_reader.word_activity_extractor as mod
        source = inspect.getsource(mod)
        tree = ast.parse(source)

        imports = []
        for node in ast.walk(tree):
            if isinstance(node, ast.Import):
                for a in node.names:
                    imports.append(a.name)
            elif isinstance(node, ast.ImportFrom):
                if node.module:
                    imports.append(node.module)

        for prohibited in ["app.quality", "app.routing", "app.exporters"]:
            assert not any(i.startswith(prohibited) for i in imports), (
                f"El extractor importa {prohibited} (prohibido en este bloque): {imports}"
            )

    # -----------------------------------------------------------------------
    # T-13: Aislamiento arquitectónico (Application no importa docx)
    # -----------------------------------------------------------------------
    def test_t13_aislamiento_arquitectonico_application_sin_docx(self):
        """T-13: Los DTOs y puertos de Application NO importan python-docx."""
        import app.application.ports.word_activity_extractor as mod_port
        import app.application.dto.word_extraction_dtos as mod_dto

        for m in [mod_port, mod_dto]:
            source = inspect.getsource(m)
            tree = ast.parse(source)
            imports = []
            for node in ast.walk(tree):
                if isinstance(node, ast.Import):
                    for a in node.names:
                        imports.append(a.name)
                elif isinstance(node, ast.ImportFrom):
                    if node.module:
                        imports.append(node.module)

            assert not any(i == "docx" or i.startswith("docx.") for i in imports), (
                f"Módulo de Application {m.__name__} importa docx: {imports}"
            )

    # -----------------------------------------------------------------------
    # -----------------------------------------------------------------------
    # PRUEBA CON DOCUMENTO REAL DE LA INSTITUCIÓN
    # -----------------------------------------------------------------------
    def test_extraccion_documento_real_institucional(self, extractor):
        """Prueba sobre un archivo real de BICU en Downloads si existe en el entorno."""
        sample_path = Path(r"C:\Users\LENOVO X1 YOGA\Downloads\Indicador 16.  Capacitación sobre uso y manejo de la herramienta Google Form y Canva, para seguimiento y desarrollo de investigaciones.docx")
        if not sample_path.exists():
            pytest.skip("Archivo real no disponible en esta máquina.")

        res = extractor.extract(sample_path)

        assert res.diagnostico.es_valido is True
        assert res.diagnostico.tipo_documento == TipoDocumentoWord.INFORME_ACTIVIDAD
        assert res.diagnostico.compatible_flujo_principal is True
        assert res.ficha_tecnica is not None
        assert "Google Form" in res.ficha_tecnica.actividad_general
        assert res.ficha_tecnica.fecha_iso == "2026-06-06"
        assert res.matriz_cuantitativa is not None
        assert res.matriz_cuantitativa.total == 15
        assert res.matriz_cuantitativa.femenino == 7
        assert res.matriz_cuantitativa.masculino == 8
        assert res.evidencias is not None
        assert res.evidencias.tiene_anexo_asistencia is True
        assert res.evidencias.tiene_anexo_fotografias is True

    def test_extraccion_muestra_documentos_reales_institucionales(self, extractor):
        """Valida que todos los informes 'Indicador *.docx' reales se extraigan con estructura válida."""
        downloads = Path(r"C:\Users\LENOVO X1 YOGA\Downloads")
        if not downloads.exists():
            pytest.skip("Directorio de descargas no disponible.")

        archivos = [downloads / f for f in os.listdir(downloads) if f.lower().startswith("indicador") and f.lower().endswith(".docx")]
        if not archivos:
            pytest.skip("No se encontraron archivos de indicador en Downloads.")

        for fpath in archivos:
            res = extractor.extract(fpath)
            assert res.diagnostico.es_valido is True, f"Fallo en archivo {fpath.name}: {res.diagnostico.errores}"
            assert res.diagnostico.tipo_documento == TipoDocumentoWord.INFORME_ACTIVIDAD
            assert res.diagnostico.compatible_flujo_principal is True
            assert res.diagnostico.tabla_1_encontrada is True
            assert res.diagnostico.tabla_2_encontrada is True
            assert res.ficha_tecnica is not None
            assert res.ficha_tecnica.actividad_general is not None
            assert res.matriz_cuantitativa is not None
            assert res.matriz_cuantitativa.total is not None
            assert res.matriz_cuantitativa.total > 0

    def test_extraccion_informe_semanal_real_clasificado_adecuadamente(self, extractor):
        """Valida que un informe semanal real sea clasificado como INFORME_SEMANAL y no compatible."""
        fpath = Path(r"C:\Users\LENOVO X1 YOGA\Downloads\Informe semanal primera semana de Junio y activiadedes del 29 de mayo 2026.docx")
        if not fpath.exists():
            pytest.skip("Informe semanal no disponible.")

        res = extractor.extract(fpath)
        assert res.diagnostico.tipo_documento == TipoDocumentoWord.INFORME_SEMANAL
        assert res.diagnostico.compatible_flujo_principal is False
        assert any("Informe Semanal" in a for a in res.diagnostico.advertencias)

    def test_extraccion_documento_incompatible_real_clasificado_adecuadamente(self, extractor):
        """Valida que un documento no institucional o administrativo sea clasificado como NO_COMPATIBLE."""
        fpath = Path(r"C:\Users\LENOVO X1 YOGA\Downloads\Carta de Solicitud a Roosman.docx")
        if not fpath.exists():
            pytest.skip("Documento administrativo no disponible.")

        res = extractor.extract(fpath)
        assert res.diagnostico.tipo_documento == TipoDocumentoWord.NO_COMPATIBLE
        assert res.diagnostico.compatible_flujo_principal is False
        assert len(res.diagnostico.errores) > 0
