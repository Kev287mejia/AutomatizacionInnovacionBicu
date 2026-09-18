"""
tests.test_word_consolidator.test_institutional_docx_generator

Suite de pruebas del generador DOCX institucional (Producto B - Fase 16.5).

Valida los 19 escenarios requeridos:
 1. Documento básico: se genera sin excepción y contiene bytes válidos.
 2. N=0 actividades: tabla con fila informativa, sección evidencias vacía.
 3. Una actividad, un estamento: tabla de 5 filas.
 4. Múltiples actividades: bloques independientes, spans correctos.
 5. Solo Estudiantes: 1 fila de datos, vMerge=none.
 6. Estudiantes + Administrativos: 2 filas, vMerge=restart/continue.
 7. Todos los estamentos: 5 filas de datos.
 8. Discrepancias presentes: observación textual en documento, datos M1 sin alteración.
 9. Fuente DECLARADO_M1: tabla muestra cifras M1.
10. Fuente NOMINAL: tabla muestra cifras nominales.
11. Evidencias vacías: sección EVIDENCIAS ANEXAS sin inventar contenido.
12. Fotografías: intento de insertar imagen (ruta inexistente → omisión limpia).
13. Registro de asistencia: idem a fotografías.
14. Enlaces: URLs reales en sección de evidencias.
15. Merges horizontales: gridSpan exacto por fila.
16. cantSplit: todas las filas tienen cantSplit activo.
17. tblHeader: filas 2 y 3 tienen tblHeader; filas de datos no.
18. Geometría de 13 columnas: total_span==13 en todas las filas.
19. Caso real Septiembre 2026: M1 declarado, datos nominales conservados, discrepancia.
"""

import hashlib
import io
import pytest

from docx import Document
from docx.oxml.ns import qn

from app.word_consolidator.document.institutional_models import (
    ActividadInstitucional,
    ConteoSexoInstitucional,
    DiscrepanciaActividadInstitucional,
    DiscrepanciaEstamento,
    EstadoEvidenciaEnum,
    InformeSemanalInstitucional,
    ItemEvidenciaInstitucional,
    ProtagonistaEstamentoInstitucional,
    SeccionEvidenciasActividad,
    TipoEstamentoInstitucional,
    TipoEvidenciaEnum,
    TrazabilidadActividad,
)
from app.word_consolidator.institutional_docx.generator import (
    InstitutionalDocxGenerator,
    generate_institutional_docx,
)
from app.word_consolidator.institutional_docx.spec import ModoFuenteDatosEnum


# ============================================================================
# HELPERS AUXILIARES DE FIXTURES
# ============================================================================

def _conteo(m: int, v: int) -> ConteoSexoInstitucional:
    return ConteoSexoInstitucional(mujeres=m, varones=v, total=m + v)


def _evidencias_vacias(id_act: str, num: int, nombre: str) -> SeccionEvidenciasActividad:
    return SeccionEvidenciasActividad(
        id_actividad=id_act, numero_actividad=num, nombre_actividad=nombre
    )


def _trazabilidad(id_act: str) -> TrazabilidadActividad:
    return TrazabilidadActividad(id_actividad=id_act, matriz_m1_origen="M1.xlsx")


def _estamento(tipo: TipoEstamentoInstitucional, label: str, m: int, v: int) -> ProtagonistaEstamentoInstitucional:
    conteo = _conteo(m, v)
    return ProtagonistaEstamentoInstitucional(
        estamento_tipo=tipo,
        denominacion_visible=label,
        conteo_presentacion=conteo,
        dato_declarado=conteo,
    )


def _actividad(
    id_act: str,
    num: int,
    nombre: str,
    estamentos: list,
    total_m: int,
    total_v: int,
    discrepancias=None,
    evidencias=None,
) -> ActividadInstitucional:
    return ActividadInstitucional(
        id_actividad=id_act,
        numero_orden=num,
        nombre_actividad=nombre,
        eje_vinculado="11.41.67",
        descripcion="Descripcion de prueba",
        sede="Bilwi",
        departamento="RACCN",
        municipio="Puerto Cabezas",
        tipo_actividad="7. Creatividad, Ciencias e Innovacion",
        protagonistas=estamentos,
        totales_actividad=_conteo(total_m, total_v),
        discrepancias=discrepancias or [],
        evidencias=evidencias or _evidencias_vacias(id_act, num, nombre),
        trazabilidad=_trazabilidad(id_act),
    )


def _informe(depto: str, actividades: list) -> InformeSemanalInstitucional:
    return InformeSemanalInstitucional(
        departamento_responsable=depto,
        mes_planificado="Septiembre",
        semana="1",
        actividades=actividades,
    )


def _generar(informe: InformeSemanalInstitucional, modo=ModoFuenteDatosEnum.DECLARADO_M1) -> Document:
    """Genera y abre el DOCX institucional como objeto Document."""
    buf = generate_institutional_docx(informe, modo_fuente=modo)
    return Document(buf)


def _get_table_rows(doc: Document):
    """Retorna los elementos w:tr de la tabla institucional."""
    return doc.tables[0].rows


def _total_span_row(row) -> int:
    """Calcula el span total de una fila (suma de gridSpan de todos los w:tc)."""
    tcs = row._tr.xpath(".//w:tc")
    return sum(
        int(tc.xpath("./w:tcPr/w:gridSpan/@w:val")[0])
        if tc.xpath("./w:tcPr/w:gridSpan/@w:val")
        else 1
        for tc in tcs
    )


def _cell_texts(row) -> list:
    tcs = row._tr.xpath(".//w:tc")
    return ["".join(tc.xpath(".//w:t/text()")).strip() for tc in tcs]


def _row_has_cant_split(row) -> bool:
    return bool(row._tr.xpath("./w:trPr/w:cantSplit"))


def _row_has_tbl_header(row) -> bool:
    return bool(row._tr.xpath("./w:trPr/w:tblHeader"))


def _cell_v_merge(tc_elem) -> str:
    v_merge = tc_elem.xpath("./w:tcPr/w:vMerge")
    if not v_merge:
        return "none"
    val = v_merge[0].get(qn("w:val"), "")
    return val if val else "continue"


def _all_paragraphs_text(doc: Document) -> list:
    return [p.text for p in doc.paragraphs]


# ============================================================================
# CASO REAL SEPTIEMBRE 2026 (fixture compartido)
# ============================================================================

@pytest.fixture
def informe_sept_2026_real():
    """Informe institucional real de Septiembre 2026 con 3 estamentos y 3 discrepancias."""
    est_est = ProtagonistaEstamentoInstitucional(
        estamento_tipo=TipoEstamentoInstitucional.ESTUDIANTES,
        denominacion_visible="2. Estudiantes",
        conteo_presentacion=_conteo(10, 5),
        dato_declarado=_conteo(10, 5),
        dato_nominal=_conteo(12, 5),
        tiene_discrepancia=True,
        discrepancia_detalle=DiscrepanciaEstamento(delta_mujeres=2, delta_varones=0, delta_total=2),
    )
    est_adm = ProtagonistaEstamentoInstitucional(
        estamento_tipo=TipoEstamentoInstitucional.PERSONAL_ADMINISTRATIVO,
        denominacion_visible="3. Personal administrativo",
        conteo_presentacion=_conteo(2, 0),
        dato_declarado=_conteo(2, 0),
        dato_nominal=_conteo(0, 1),
        tiene_discrepancia=True,
        discrepancia_detalle=DiscrepanciaEstamento(delta_mujeres=-2, delta_varones=1, delta_total=-1),
    )
    est_doc = ProtagonistaEstamentoInstitucional(
        estamento_tipo=TipoEstamentoInstitucional.MAESTRAS_Y_MAESTROS,
        denominacion_visible="1. Maestras y maestros",
        conteo_presentacion=_conteo(0, 1),
        dato_declarado=_conteo(0, 1),
        dato_nominal=_conteo(0, 0),
        tiene_discrepancia=True,
        discrepancia_detalle=DiscrepanciaEstamento(delta_mujeres=0, delta_varones=-1, delta_total=-1),
    )
    discrepancias = [
        DiscrepanciaActividadInstitucional(
            estamento="ESTUDIANTES",
            fuente_declarada="Matriz 1",
            fuente_nominal="Matriz 2",
            declarado_m=10, declarado_v=5, declarado_total=15,
            nominal_m=12, nominal_v=5, nominal_total=17,
            delta_total=2, delta_m=2, delta_v=0,
            descripcion="Diferencia en estudiantes (+2 nominal)",
        ),
        DiscrepanciaActividadInstitucional(
            estamento="PERSONAL_ADMINISTRATIVO",
            fuente_declarada="Matriz 1",
            fuente_nominal="Matriz 3",
            declarado_m=2, declarado_v=0, declarado_total=2,
            nominal_m=0, nominal_v=1, nominal_total=1,
            delta_total=-1, delta_m=-2, delta_v=1,
            descripcion="Diferencia en administrativos (-1 nominal)",
        ),
        DiscrepanciaActividadInstitucional(
            estamento="MAESTRAS_Y_MAESTROS",
            fuente_declarada="Matriz 1",
            fuente_nominal="Matriz 4",
            declarado_m=0, declarado_v=1, declarado_total=1,
            nominal_m=0, nominal_v=0, nominal_total=0,
            delta_total=-1, delta_m=0, delta_v=-1,
            descripcion="Diferencia en docentes (-1 nominal)",
        ),
    ]
    evidencias = SeccionEvidenciasActividad(
        id_actividad="act-sept-2026-001",
        numero_actividad=1,
        nombre_actividad="BICU CUR Bilwi fortalece conocimientos estudiantiles",
        enlaces_publicaciones=[
            ItemEvidenciaInstitucional(
                tipo=TipoEvidenciaEnum.ENLACE_MEDIOS,
                titulo="Publicacion Facebook",
                referencia="https://www.facebook.com/share/1J19aqgr3U/",
                estado=EstadoEvidenciaEnum.VINCULADO,
                orden=1,
            )
        ],
    )
    actividad = _actividad(
        "act-sept-2026-001",
        1,
        "BICU CUR Bilwi fortalece conocimientos estudiantiles en diseno de logotipos",
        [est_est, est_adm, est_doc],
        12, 6,
        discrepancias=discrepancias,
        evidencias=evidencias,
    )
    return _informe("Innovacion y Emprendimiento", [actividad])


# ============================================================================
# TEST SUITE PRINCIPAL
# ============================================================================

class TestInstitutionalDocxGenerator:

    # ─────────────────────────────────────────────────
    # TEST 1: Documento básico
    # ─────────────────────────────────────────────────
    def test_01_documento_basico_genera_bytes_validos(self):
        """El generador produce un DOCX con bytes válidos sin excepción."""
        informe = _informe("Dpto Test", [])
        buf = generate_institutional_docx(informe)
        assert isinstance(buf, io.BytesIO)
        content = buf.read()
        assert len(content) > 5000  # Un DOCX mínimo siempre pesa > 5 KB
        # Verificar que es un ZIP válido (DOCX = ZIP)
        assert content[:2] == b"PK"

    # ─────────────────────────────────────────────────
    # TEST 2: N=0 actividades
    # ─────────────────────────────────────────────────
    def test_02_n0_actividades_tabla_con_fila_informativa(self):
        """Informe sin actividades → tabla con 5 filas (4 encabezados + 1 estado vacío)."""
        informe = _informe("Dpto Test", [])
        doc = _generar(informe)
        assert len(doc.tables) == 1
        rows = _get_table_rows(doc)
        assert len(rows) == 5  # 4 encabezados + 1 estado vacío
        # Fila de estado vacío: span completo de 13
        empty_row = rows[4]
        tcs = empty_row._tr.xpath(".//w:tc")
        assert len(tcs) == 1
        span = int(tcs[0].xpath("./w:tcPr/w:gridSpan/@w:val")[0])
        assert span == 13

    # ─────────────────────────────────────────────────
    # TEST 3: Una actividad con un estamento
    # ─────────────────────────────────────────────────
    def test_03_una_actividad_un_estamento_5_filas(self):
        """Una actividad con 1 estamento → tabla de 5 filas."""
        est = _estamento(TipoEstamentoInstitucional.ESTUDIANTES, "Estudiantes", 10, 5)
        act = _actividad("a1", 1, "Actividad Uno", [est], 10, 5)
        doc = _generar(_informe("Depto Test", [act]))
        assert len(doc.tables[0].rows) == 5

    # ─────────────────────────────────────────────────
    # TEST 4: Múltiples actividades (N=3)
    # ─────────────────────────────────────────────────
    def test_04_multiples_actividades_bloques_independientes(self):
        """N=3 actividades → filas correctas y datos no mezclados entre actividades."""
        est_a1_1 = _estamento(TipoEstamentoInstitucional.ESTUDIANTES, "Est A1", 5, 5)
        est_a1_2 = _estamento(TipoEstamentoInstitucional.PERSONAL_ADMINISTRATIVO, "Adm A1", 1, 0)
        act1 = _actividad("a1", 1, "Actividad 1", [est_a1_1, est_a1_2], 6, 5)

        est_a2 = _estamento(TipoEstamentoInstitucional.ESTUDIANTES, "Est A2", 20, 10)
        act2 = _actividad("a2", 2, "Actividad 2", [est_a2], 20, 10)

        est_a3_1 = _estamento(TipoEstamentoInstitucional.MAESTRAS_Y_MAESTROS, "Doc A3", 3, 2)
        est_a3_2 = _estamento(TipoEstamentoInstitucional.ESTUDIANTES, "Est A3", 7, 3)
        est_a3_3 = _estamento(TipoEstamentoInstitucional.COLABORADORES, "Col A3", 1, 1)
        act3 = _actividad("a3", 3, "Actividad 3", [est_a3_1, est_a3_2, est_a3_3], 11, 6)

        doc = _generar(_informe("Dpto Test", [act1, act2, act3]))
        # 4 encabezados + 2 + 1 + 3 = 10 filas
        assert len(doc.tables[0].rows) == 10

        rows = doc.tables[0].rows
        # Fila 4 → actividad 1 primer estamento
        texts_4 = _cell_texts(rows[4])
        assert "Actividad 1" in texts_4[0]
        # Fila 6 → actividad 2 (span = none, no vMerge)
        texts_6 = _cell_texts(rows[6])
        assert "Actividad 2" in texts_6[0]

    # ─────────────────────────────────────────────────
    # TEST 5: Solo Estudiantes
    # ─────────────────────────────────────────────────
    def test_05_solo_estudiantes_1_fila_datos(self):
        """1 actividad con solo Estudiantes → 1 fila de datos, vMerge=none."""
        est = _estamento(TipoEstamentoInstitucional.ESTUDIANTES, "2. Estudiantes", 10, 5)
        act = _actividad("a1", 1, "Taller IA", [est], 10, 5)
        doc = _generar(_informe("Dpto", [act]))
        row_4 = doc.tables[0].rows[4]
        tcs = row_4._tr.xpath(".//w:tc")
        # Col 0: vMerge debe ser none (solo 1 estamento)
        assert _cell_v_merge(tcs[0]) == "none"
        texts = _cell_texts(row_4)
        assert "2. Estudiantes" in texts[7]
        assert texts[8] == "10"
        assert texts[9] == "5"
        assert texts[10] == "15"

    # ─────────────────────────────────────────────────
    # TEST 6: Estudiantes + Administrativos (vMerge restart/continue)
    # ─────────────────────────────────────────────────
    def test_06_estudiantes_y_administrativos_vmerge_correcto(self):
        """2 estamentos → fila 0: vMerge=restart, fila 1: vMerge=continue."""
        est1 = _estamento(TipoEstamentoInstitucional.ESTUDIANTES, "Estudiantes", 10, 5)
        est2 = _estamento(TipoEstamentoInstitucional.PERSONAL_ADMINISTRATIVO, "Administrativos", 2, 0)
        act = _actividad("a1", 1, "Capacitacion", [est1, est2], 12, 5)
        doc = _generar(_informe("Dpto", [act]))

        rows = doc.tables[0].rows
        row_4 = rows[4]
        row_5 = rows[5]

        tcs_4 = row_4._tr.xpath(".//w:tc")
        tcs_5 = row_5._tr.xpath(".//w:tc")

        # Fila 4: col 0 vMerge=restart
        assert _cell_v_merge(tcs_4[0]) == "restart"
        # Fila 5: col 0 vMerge=continue
        assert _cell_v_merge(tcs_5[0]) == "continue"
        # Fila 5: el texto del nombre debe estar vacío (fusión)
        assert _cell_texts(row_5)[0] == ""
        assert "Administrativos" in _cell_texts(row_5)[7]

    # ─────────────────────────────────────────────────
    # TEST 7: Todos los estamentos (5 filas de datos)
    # ─────────────────────────────────────────────────
    def test_07_todos_los_estamentos_5_filas_datos(self):
        """5 estamentos → 9 filas totales (4 enc + 5 datos)."""
        estamentos = [
            _estamento(TipoEstamentoInstitucional.ESTUDIANTES, "Estudiantes", 10, 5),
            _estamento(TipoEstamentoInstitucional.PERSONAL_ADMINISTRATIVO, "Administrativos", 2, 1),
            _estamento(TipoEstamentoInstitucional.MAESTRAS_Y_MAESTROS, "Docentes", 1, 1),
            _estamento(TipoEstamentoInstitucional.COLABORADORES, "Colaboradores", 1, 0),
            _estamento(TipoEstamentoInstitucional.BENEFICIARIOS_POBLADORES, "Poblacion general", 5, 5),
        ]
        act = _actividad("a1", 1, "Evento Grande", estamentos, 19, 12)
        doc = _generar(_informe("Dpto", [act]))
        assert len(doc.tables[0].rows) == 9

    # ─────────────────────────────────────────────────
    # TEST 8: Discrepancias → observación en documento, datos M1 sin alterar
    # ─────────────────────────────────────────────────
    def test_08_discrepancias_generan_observacion_discreta(self, informe_sept_2026_real):
        """Discrepancias presentes → párrafo de observación en EVIDENCIAS, datos M1 intactos."""
        doc = _generar(informe_sept_2026_real, modo=ModoFuenteDatosEnum.DECLARADO_M1)
        paragraphs_text = _all_paragraphs_text(doc)

        # Debe haber un párrafo de observación de discrepancias
        obs_found = any(
            "Observaci" in t and ("diferencia" in t.lower() or "declarados" in t.lower())
            for t in paragraphs_text
        )
        assert obs_found, "Debe existir párrafo de observación de discrepancias"

        # Los datos en la tabla son M1 (10, 5, 15) no Nominal (12, 5, 17)
        rows = doc.tables[0].rows
        texts_row4 = _cell_texts(rows[4])
        assert texts_row4[8] == "10"   # M declarado
        assert texts_row4[9] == "5"    # V declarado
        assert texts_row4[10] == "15"  # Total declarado

    # ─────────────────────────────────────────────────
    # TEST 9: Fuente DECLARADO_M1
    # ─────────────────────────────────────────────────
    def test_09_fuente_declarado_m1_muestra_cifras_m1(self, informe_sept_2026_real):
        """Modo DECLARADO_M1: tabla muestra cifras declaradas, no nominales."""
        doc = _generar(informe_sept_2026_real, modo=ModoFuenteDatosEnum.DECLARADO_M1)
        rows = doc.tables[0].rows
        # Estudiantes: M1=(10,5,15), Nominal=(12,5,17)
        texts_est = _cell_texts(rows[4])
        assert texts_est[8] == "10"   # M1
        assert texts_est[9] == "5"    # M1
        assert texts_est[10] == "15"  # M1

        # Admin: M1=(2,0,2)
        texts_adm = _cell_texts(rows[5])
        assert texts_adm[8] == "2"
        assert texts_adm[9] == "0"
        assert texts_adm[10] == "2"

    # ─────────────────────────────────────────────────
    # TEST 10: Fuente NOMINAL
    # ─────────────────────────────────────────────────
    def test_10_fuente_nominal_muestra_cifras_nominales(self, informe_sept_2026_real):
        """Modo NOMINAL: tabla muestra cifras nominales verificadas (M2–M5)."""
        doc = _generar(informe_sept_2026_real, modo=ModoFuenteDatosEnum.NOMINAL)
        rows = doc.tables[0].rows
        # Estudiantes Nominal: (12, 5, 17)
        texts_est = _cell_texts(rows[4])
        assert texts_est[8] == "12"
        assert texts_est[9] == "5"
        assert texts_est[10] == "17"

        # Admin Nominal: (0, 1, 1)
        texts_adm = _cell_texts(rows[5])
        assert texts_adm[8] == "0"
        assert texts_adm[9] == "1"
        assert texts_adm[10] == "1"

        # Docentes Nominal: (0, 0, 0)
        texts_doc = _cell_texts(rows[6])
        assert texts_doc[8] == "0"
        assert texts_doc[9] == "0"
        assert texts_doc[10] == "0"

    # ─────────────────────────────────────────────────
    # TEST 11: Evidencias vacías → no se inventa contenido
    # ─────────────────────────────────────────────────
    def test_11_evidencias_vacias_no_inventa_contenido(self):
        """Evidencias vacías → EVIDENCIAS ANEXAS presente pero sin fotos o URLs ficticias."""
        act = _actividad("a1", 1, "Evento Sin Evidencias", [], 0, 0)
        doc = _generar(_informe("Dpto", [act]))
        paragraphs_text = _all_paragraphs_text(doc)

        # El título de evidencias siempre aparece
        ev_title = any("EVIDENCIAS ANEXAS" in t for t in paragraphs_text)
        assert ev_title

        # NO debe haber párrafos de "Fotografías" o "Registro de asistencias"
        fotos_label = any("Fotograf" in t for t in paragraphs_text)
        assert not fotos_label, "No se debe insertar rótulo 'Fotografías' si no hay fotos"
        asistencia_label = any("asistencia" in t.lower() for t in paragraphs_text)
        assert not asistencia_label, "No se debe insertar rótulo de asistencias si no hay registros"

    # ─────────────────────────────────────────────────
    # TEST 12: Fotografías (ruta inexistente → omisión limpia)
    # ─────────────────────────────────────────────────
    def test_12_fotografia_ruta_inexistente_omision_limpia(self):
        """Si la imagen referenciada no existe, se omite sin error y sin placeholder."""
        foto = ItemEvidenciaInstitucional(
            tipo=TipoEvidenciaEnum.FOTOGRAFIA,
            titulo="Foto prueba inexistente",
            referencia="C:/no_existe/foto_falsa.jpg",
            estado=EstadoEvidenciaEnum.PENDIENTE,
            orden=1,
        )
        sec = SeccionEvidenciasActividad(
            id_actividad="a1", numero_actividad=1, nombre_actividad="Act1",
            fotografias=[foto],
        )
        act = _actividad("a1", 1, "Taller Fotos", [], 0, 0, evidencias=sec)

        import warnings
        with warnings.catch_warnings(record=True) as w_list:
            warnings.simplefilter("always")
            buf = generate_institutional_docx(_informe("Dpto", [act]))
            # Debe haber una advertencia sobre imagen no encontrada
            foto_warnings = [w for w in w_list if "Imagen no encontrada" in str(w.message)]
            assert len(foto_warnings) >= 1

        doc = Document(buf)
        # El documento se generó sin excepción y tiene bytes válidos
        assert len(doc.tables) == 1

    # ─────────────────────────────────────────────────
    # TEST 13: Registro de asistencia (ruta inexistente → omisión)
    # ─────────────────────────────────────────────────
    def test_13_registro_asistencia_ruta_inexistente_omision_limpia(self):
        """Registro de asistencia con ruta inexistente → omisión sin error."""
        asistencia = ItemEvidenciaInstitucional(
            tipo=TipoEvidenciaEnum.REGISTRO_ASISTENCIA,
            titulo="Lista asistencia falsa",
            referencia="C:/no_existe/asistencia.pdf",
            estado=EstadoEvidenciaEnum.PENDIENTE,
            orden=1,
        )
        sec = SeccionEvidenciasActividad(
            id_actividad="a1", numero_actividad=1, nombre_actividad="Act1",
            registros_asistencia=[asistencia],
        )
        act = _actividad("a1", 1, "Taller", [], 0, 0, evidencias=sec)

        import warnings
        with warnings.catch_warnings(record=True):
            warnings.simplefilter("always")
            buf = generate_institutional_docx(_informe("Dpto", [act]))

        doc = Document(buf)
        assert len(doc.tables) == 1  # Se generó sin excepciones

    # ─────────────────────────────────────────────────
    # TEST 14: URLs reales en sección de evidencias
    # ─────────────────────────────────────────────────
    def test_14_enlaces_urls_reales_en_evidencias(self):
        """URLs reales del modelo aparecen en sección de evidencias."""
        url_real = "https://www.facebook.com/share/1J19aqgr3U/"
        enlace = ItemEvidenciaInstitucional(
            tipo=TipoEvidenciaEnum.ENLACE_MEDIOS,
            titulo="Publicacion Facebook",
            referencia=url_real,
            estado=EstadoEvidenciaEnum.VINCULADO,
            orden=1,
        )
        sec = SeccionEvidenciasActividad(
            id_actividad="a1", numero_actividad=1, nombre_actividad="Act1",
            enlaces_publicaciones=[enlace],
        )
        act = _actividad("a1", 1, "Evento Con Enlace", [], 0, 0, evidencias=sec)
        doc = _generar(_informe("Dpto", [act]))
        paragraphs_text = _all_paragraphs_text(doc)
        url_found = any(url_real in t for t in paragraphs_text)
        assert url_found, "La URL real debe aparecer en la sección de evidencias"

    # ─────────────────────────────────────────────────
    # TEST 15: Merges horizontales (gridSpan exacto)
    # ─────────────────────────────────────────────────
    def test_15_grid_span_exacto_por_fila(self):
        """Verifica gridSpan exacto en filas de metadatos y encabezados."""
        informe = _informe("Innovacion", [])
        doc = _generar(informe)
        rows = doc.tables[0].rows
        tbl = doc.tables[0]

        # Fila 0: 2 celdas, spans = [1, 12]
        row0_tcs = rows[0]._tr.xpath(".//w:tc")
        spans_0 = [
            int(tc.xpath("./w:tcPr/w:gridSpan/@w:val")[0])
            if tc.xpath("./w:tcPr/w:gridSpan/@w:val") else 1
            for tc in row0_tcs
        ]
        assert spans_0 == [1, 12]

        # Fila 1: 4 celdas, spans = [1, 6, 2, 4]
        row1_tcs = rows[1]._tr.xpath(".//w:tc")
        spans_1 = [
            int(tc.xpath("./w:tcPr/w:gridSpan/@w:val")[0])
            if tc.xpath("./w:tcPr/w:gridSpan/@w:val") else 1
            for tc in row1_tcs
        ]
        assert spans_1 == [1, 6, 2, 4]

        # Fila 2: 3 celdas, spans = [1, 6, 6]
        row2_tcs = rows[2]._tr.xpath(".//w:tc")
        spans_2 = [
            int(tc.xpath("./w:tcPr/w:gridSpan/@w:val")[0])
            if tc.xpath("./w:tcPr/w:gridSpan/@w:val") else 1
            for tc in row2_tcs
        ]
        assert spans_2 == [1, 6, 6]

    # ─────────────────────────────────────────────────
    # TEST 16: cantSplit en todas las filas
    # ─────────────────────────────────────────────────
    def test_16_cant_split_en_todas_las_filas(self):
        """Todas las filas de la tabla tienen w:cantSplit activo."""
        est = _estamento(TipoEstamentoInstitucional.ESTUDIANTES, "Estudiantes", 5, 5)
        act = _actividad("a1", 1, "Taller", [est], 5, 5)
        doc = _generar(_informe("Dpto", [act]))
        for i, row in enumerate(doc.tables[0].rows):
            assert _row_has_cant_split(row), f"Fila {i} no tiene cantSplit"

    # ─────────────────────────────────────────────────
    # TEST 17: tblHeader en filas de encabezado
    # ─────────────────────────────────────────────────
    def test_17_tbl_header_en_filas_de_encabezado_no_en_datos(self):
        """Filas 2 y 3 tienen tblHeader; metadatos y datos no lo tienen."""
        est = _estamento(TipoEstamentoInstitucional.ESTUDIANTES, "Estudiantes", 5, 5)
        act = _actividad("a1", 1, "Taller", [est], 5, 5)
        doc = _generar(_informe("Dpto", [act]))
        rows = doc.tables[0].rows

        # Filas 0 y 1 (metadatos): no son tblHeader
        assert not _row_has_tbl_header(rows[0]), "Fila 0 (metadato depto) no debe ser tblHeader"
        assert not _row_has_tbl_header(rows[1]), "Fila 1 (metadato mes/semana) no debe ser tblHeader"
        # Filas 2 y 3 (encabezados): sí son tblHeader
        assert _row_has_tbl_header(rows[2]), "Fila 2 (superheader) debe ser tblHeader"
        assert _row_has_tbl_header(rows[3]), "Fila 3 (subheader) debe ser tblHeader"
        # Fila 4 (datos): no es tblHeader
        assert not _row_has_tbl_header(rows[4]), "Fila 4 (datos) no debe ser tblHeader"

    # ─────────────────────────────────────────────────
    # TEST 18: Geometría de 13 columnas invariante en TODAS las filas
    # ─────────────────────────────────────────────────
    def test_18_geometria_13_columnas_invariante_todas_las_filas(self):
        """En TODAS las filas, la suma de gridSpan es exactamente 13."""
        estamentos = [
            _estamento(TipoEstamentoInstitucional.ESTUDIANTES, "Estudiantes", 10, 5),
            _estamento(TipoEstamentoInstitucional.PERSONAL_ADMINISTRATIVO, "Admin", 2, 0),
            _estamento(TipoEstamentoInstitucional.MAESTRAS_Y_MAESTROS, "Docentes", 0, 1),
        ]
        act1 = _actividad("a1", 1, "Actividad 1", estamentos, 12, 6)
        act2 = _actividad("a2", 2, "Actividad 2", [estamentos[0]], 10, 5)
        informe = _informe("Dpto", [act1, act2])

        doc = _generar(informe)
        for i, row in enumerate(doc.tables[0].rows):
            span = _total_span_row(row)
            assert span == 13, f"Fila {i} tiene span={span}, debe ser 13"

    # ─────────────────────────────────────────────────
    # TEST 19: Caso Real Septiembre 2026 completo
    # ─────────────────────────────────────────────────
    def test_19_caso_real_septiembre_2026_m1_declarado_con_discrepancias(self, informe_sept_2026_real):
        """
        Caso real Septiembre 2026:
        - 3 estamentos: Estudiantes 10/5/15, Administrativos 2/0/2, Docentes 0/1/1.
        - Datos nominales conservados en el modelo (no alterados).
        - 3 discrepancias presentes.
        - Observación de discrepancias en documento.
        - URL de Facebook en sección de evidencias.
        """
        gen = InstitutionalDocxGenerator(modo_fuente=ModoFuenteDatosEnum.DECLARADO_M1)
        buf, sha256 = gen.generate_and_compute_sha256(informe_sept_2026_real)

        # Verificar hash SHA-256
        assert len(sha256) == 64
        assert all(c in "0123456789abcdef" for c in sha256)

        doc = Document(buf)
        rows = doc.tables[0].rows

        # 4 encabezados + 3 estamentos = 7 filas
        assert len(rows) == 7

        # Verificar metadatos
        texts_0 = _cell_texts(rows[0])
        assert "Nombre del Depto:" in texts_0[0]
        assert "Innovacion" in texts_0[1]

        texts_1 = _cell_texts(rows[1])
        assert "Septiembre" in texts_1[1]
        assert "1" in texts_1[3]

        # Datos M1 en fila de estudiantes
        texts_est = _cell_texts(rows[4])
        assert "2. Estudiantes" in texts_est[7]
        assert texts_est[8] == "10"
        assert texts_est[9] == "5"
        assert texts_est[10] == "15"

        # Datos M1 en fila de administrativos
        texts_adm = _cell_texts(rows[5])
        assert "Personal administrativo" in texts_adm[7]
        assert texts_adm[8] == "2"
        assert texts_adm[9] == "0"
        assert texts_adm[10] == "2"

        # Datos M1 en fila de docentes
        texts_doc = _cell_texts(rows[6])
        assert "Maestras y maestros" in texts_doc[7]
        assert texts_doc[8] == "0"
        assert texts_doc[9] == "1"
        assert texts_doc[10] == "1"

        # URL de Facebook en evidencias
        paragraphs_text = _all_paragraphs_text(doc)
        url_found = any("facebook.com" in t for t in paragraphs_text)
        assert url_found

        # Observación de discrepancias presente
        obs_found = any("diferencia" in t.lower() or "declarados" in t.lower() for t in paragraphs_text)
        assert obs_found

        # Invariante de 13 columnas en todas las filas
        for i, row in enumerate(rows):
            assert _total_span_row(row) == 13, f"Fila {i} viola invariante de 13 columnas"

        # Los datos nominales continúan en el modelo sin modificación
        act = informe_sept_2026_real.actividades[0]
        assert act.protagonistas[0].dato_nominal.total == 17  # Estudiantes nominal
        assert act.protagonistas[1].dato_nominal.total == 1   # Admin nominal
        assert act.protagonistas[2].dato_nominal.total == 0   # Docentes nominal
        assert len(act.discrepancias) == 3


# ============================================================================
# PRUEBAS ADICIONALES: GEOMETRÍA DE PÁGINA Y ESTILO
# ============================================================================

class TestInstitutionalDocxPageGeometry:

    def test_page_geometry_landscape_letter(self):
        """Página Landscape Letter: 11in × 8.5in con márgenes correctos."""
        doc = _generar(_informe("Dpto", []))
        section = doc.sections[0]
        # Landscape Letter
        assert abs(section.page_width.inches - 11.0) < 0.05
        assert abs(section.page_height.inches - 8.5) < 0.05
        # Márgenes
        assert abs(section.top_margin.cm - 3.0) < 0.1
        assert abs(section.bottom_margin.cm - 3.0) < 0.1
        assert abs(section.left_margin.cm - 2.5) < 0.1
        assert abs(section.right_margin.cm - 2.5) < 0.1

    def test_titulos_institucionales_presentes(self):
        """Los dos títulos institucionales están presentes en el documento."""
        doc = _generar(_informe("Dpto", []))
        paragraphs_text = _all_paragraphs_text(doc)
        assert "MECANISMO INSTITUCIONAL" in paragraphs_text
        assert "INFORME SEMANAL" in paragraphs_text

    def test_sha256_determinista(self):
        """El mismo informe produce el mismo SHA-256."""
        informe = _informe("Dpto Reproducible", [])
        gen = InstitutionalDocxGenerator()
        _, sha1 = gen.generate_and_compute_sha256(informe)
        _, sha2 = gen.generate_and_compute_sha256(informe)
        assert sha1 == sha2

    def test_modo_fuente_predeterminado_es_declarado_m1(self):
        """El modo fuente predeterminado es DECLARADO_M1."""
        gen = InstitutionalDocxGenerator()
        assert gen.modo_fuente == ModoFuenteDatosEnum.DECLARADO_M1

    def test_output_path_escribe_archivo_en_disco(self, tmp_path):
        """El parámetro output_path guarda el DOCX en disco correctamente."""
        out = tmp_path / "test_output.docx"
        informe = _informe("Dpto", [])
        buf = generate_institutional_docx(informe, output_path=out)
        assert out.exists()
        assert out.stat().st_size > 5000

    def test_tabla_tiene_bordes_table_grid(self):
        """La tabla tiene bordes definidos explícitamente."""
        doc = _generar(_informe("Dpto", []))
        tbl = doc.tables[0]
        tblPr = tbl._tbl.tblPr
        tblBorders = tblPr.findall(qn("w:tblBorders"))
        assert len(tblBorders) == 1, "Debe existir exactamente un elemento w:tblBorders"
