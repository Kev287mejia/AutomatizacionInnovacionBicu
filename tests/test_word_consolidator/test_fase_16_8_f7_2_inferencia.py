"""
tests.test_word_consolidator.test_fase_16_8_f7_2_inferencia

Suite de pruebas automatizadas para la Fase 16.8-F.7.2:
Corrección quirúrgica del defecto de integración e inferencia de cabecera institucional en Product B.

Cubre rigurosamente los 12 tests obligatorios más los 5 tests adicionales requeridos en la autorización:
- Test 1: Simulación de ruta real GUI (MainWindow -> ConsolidationWorker -> pipeline) sin metadata manual.
- Test 2: Metadatos explícitos ganan sobre M1/M2/M3.
- Test 3: M1 informado gana sobre M2/M3 según jerarquía.
- Test 4: M1 vacío + M2/M3 coincidentes resuelve valor unívoco.
- Test 5: M1 vacío + solo M2 informado resuelve valor unívoco.
- Test 6: M1 vacío + solo M3 informado resuelve valor unívoco.
- Test 7: M2/M3 con valores diferentes detecta ambigüedad, no elige arbitrariamente y usa fallback.
- Test 8: Todo vacío mantiene fallback 'Área No Especificada'.
- Test 9: Integridad de Producto A sin regresiones funcionales.
- Test 10: Preservación intacta de los 32 registros históricos de M5.
- Test 11: Inmutabilidad criptográfica byte por byte de los 5 Excel de entrada (SHA-256).
- Test 12: Determinismo estricto (dos ejecuciones con mismos inputs producen idéntico resultado).
- Test Adicional A: Dos actividades diferentes no se mezclan.
- Test Adicional B: Misma actividad en dos períodos diferentes con programas diferentes (aislamiento temporal).
- Test Adicional C: M2 tiene un programa y M3 otro para la misma actividad (detección de ambigüedad).
- Test Adicional D: M2 con múltiples filas del mismo programa sigue siendo un único valor inequívoco.
- Test Adicional E: Una actividad con nominales y otra sin nominales no reutiliza el programa entre sí.
"""

import hashlib
from pathlib import Path
import queue
import shutil
import pytest
from docx import Document

from app.core.constants.participant_types import CategoriaParticipacion
from app.word_consolidator.models import (
    ConjuntoMatricesLeidas,
    FilaLeidaM1,
    FilaLeidaNominal,
    MetadatosInstitucionales,
    PeriodoConsolidacion,
    TipoPeriodo,
)
from app.word_consolidator.engine.consolidation_engine import (
    ActividadConsolidada,
    ConsolidationEngine,
    ParticipacionConsolidada,
    ResultadoConsolidacion,
    TotalesM1,
    TotalesNominales,
)
from app.word_consolidator.document.institutional_models import (
    InformeSemanalInstitucional,
    TipoEstamentoInstitucional,
)
from app.word_consolidator.document.institutional_transformer import InstitutionalReportTransformer
from app.word_consolidator.pipeline import WordConsolidationPipeline, calcular_sha256
from app.word_consolidator.readers.matrix_reader import MatrixReader
from app.word_consolidator.readers.period_filter import PeriodFilter
from app.word_consolidator.ui.services.application_service import ConsolidationAppService
from app.word_consolidator.ui.workers.consolidation_worker import (
    ConsolidationWorker,
    ConsolidationEventType,
)


@pytest.fixture
def periodo_septiembre_2026() -> PeriodoConsolidacion:
    return PeriodoConsolidacion(
        tipo_periodo=TipoPeriodo.SEMANA,
        anio=2026,
        mes=9,
        semana=1,
        etiqueta="Septiembre 2026 — Semana 1",
    )


@pytest.fixture
def dir_output_real() -> Path:
    p = Path("output")
    if not (p / "Matriz_1_Consolidado_Actividades.xlsx").exists():
        pytest.skip("Matrices reales en output/ no disponibles en este entorno")
    return p


# ============================================================================
# TEST 1: RUTA REAL GUI SIN METADATA MANUAL INYECTADA
# ============================================================================

def test_01_gui_sin_metadata_externa_caso_real(tmp_path, dir_output_real, periodo_septiembre_2026):
    """
    Test 1: Simula la ruta real de la GUI:
    MainWindow -> ConsolidationWorker -> pipeline
    sin inyectar manualmente metadatos con departamento institucional.
    Debe producir 'Innovación y Emprendimiento' en la cabecera de Product B y
    mantener los conteos declarados por M1.
    """
    matrices = {
        "M1": dir_output_real / "Matriz_1_Consolidado_Actividades.xlsx",
        "M2": dir_output_real / "Matriz_2_Estudiantes.xlsx",
        "M3": dir_output_real / "Matriz_3_Academicos_Administrativos.xlsx",
        "M4": dir_output_real / "Matriz_4_Colaboradores.xlsx",
        "M5": dir_output_real / "Matriz_5_Protagonistas_Beneficiados.xlsx",
    }
    event_q: queue.Queue = queue.Queue()

    # Simular MainWindow instanciando el worker sin metadata manual (usando servicio o default None)
    worker = ConsolidationWorker(
        fuentes=matrices,
        periodo=periodo_septiembre_2026,
        salida_dir=tmp_path,
        event_queue=event_q,
        metadatos=None,  # Simulación exacta del flujo GUI que causaba el fallo
        generar_tecnico=False,
        generar_institucional=True,
    )
    worker.run()

    # Recolectar resultado final de los eventos del worker
    resultado_final = None
    while not event_q.empty():
        ev = event_q.get_nowait()
        if ev.event_type == ConsolidationEventType.COMPLETED:
            resultado_final = ev.result

    assert resultado_final is not None, "El worker no emitió resultado de finalización COMPLETED"
    assert resultado_final.status_institucional == "SUCCESS"
    assert resultado_final.ruta_docx_institucional is not None
    docx_path = Path(resultado_final.ruta_docx_institucional)
    assert docx_path.exists()

    # Inspeccionar el Word generado físicamente
    doc = Document(docx_path)
    tabla = doc.tables[0]

    # Verificar cabecera institucional exacta
    cabecera_depto = tabla.rows[0].cells[1].text
    assert "Innovación y Emprendimiento" in cabecera_depto
    assert "Área No Especificada" not in cabecera_depto

    # Verificar preservación estricta de conteos DECLARADO_M1
    # Fila Estudiantes: 11 M, 6 V, 17 Total
    fila_est = [c.text.strip() for c in tabla.rows[4].cells]
    assert fila_est[8] == "Estudiantes"
    assert fila_est[10] == "11"
    assert fila_est[11] == "6"
    assert fila_est[12] == "17"

    # Fila Personal administrativo: 1 M, 0 V, 1 Total
    fila_adm = [c.text.strip() for c in tabla.rows[5].cells]
    assert fila_adm[8] == "Personal administrativo"
    assert fila_adm[10] == "1"
    assert fila_adm[11] == "0"
    assert fila_adm[12] == "1"


# ============================================================================
# TEST 2: METADATA EXPLÍCITA GANA SOBRE M1/M2/M3 (NIVEL 1)
# ============================================================================

def test_02_metadata_explicita_prioridad_maxima(periodo_septiembre_2026):
    """
    Test 2: Si departamento_institucional='Valor Explícito' en metadatos,
    debe ganar sobre M1 y sobre M2/M3 según la prioridad del Nivel 1.
    """
    fila_m1 = FilaLeidaM1(
        hoja="Consolidado", fila=2, sede="Bilwi",
        actividad="Actividad Prueba",
        programa="Programa M1", area_responsable="Area M1",
    )
    p_nom = ParticipacionConsolidada(
        id_actividad="act-1", id_persona="per-1", nombre_apellidos="Juan Perez",
        categoria=CategoriaParticipacion.ESTUDIANTE, matriz_origen="M2",
        hoja="Estudiantes", fila=2, actividad_asociada="Actividad Prueba",
        sede="Bilwi", programa="Programa M2",
    )
    act = ActividadConsolidada(
        clave_negocio=("actividad prueba", "bilwi", None),
        nombre_actividad="Actividad Prueba", sede="Bilwi",
        fuente_m1=fila_m1, programa="Programa M1", area_responsable="Area M1",
        participaciones=[p_nom],
    )
    cons = ResultadoConsolidacion(
        actividades=[act], fuentes={}, total_actividades=1,
        total_asistencia_bruta=1, total_personas_unicas=1,
        total_recurrencia=0, tasa_recurrencia=1.0,
        periodo=periodo_septiembre_2026,
    )

    metadatos = MetadatosInstitucionales(departamento_institucional="Valor Explícito")
    informe = InstitutionalReportTransformer.transformar(
        consolidacion=cons,
        metadatos=metadatos,
        politica_presentacion="DECLARADO_M1",
    )
    assert informe.departamento_responsable == "Valor Explícito"
    assert informe.trazabilidad_global.get("departamento_responsable_origen") == "NIVEL_1_METADATOS_EXPLICITOS"


# ============================================================================
# TEST 3: M1 INFORMADO GANA SOBRE M2/M3 SEGÚN JERARQUÍA (NIVEL 2)
# ============================================================================

def test_03_m1_informado_gana_sobre_nominales(periodo_septiembre_2026):
    """
    Test 3: Si M1 tiene Programa='Programa M1' y M2/M3 tienen otro programa,
    debe respetarse M1 según la jerarquía definida (Nivel 2 antes que Nivel 3).
    """
    fila_m1 = FilaLeidaM1(
        hoja="Consolidado", fila=2, sede="Bilwi",
        actividad="Actividad Prueba",
        programa="Programa M1 Oficial",
    )
    p_nom = ParticipacionConsolidada(
        id_actividad="act-1", id_persona="per-1", nombre_apellidos="Ana Gomez",
        categoria=CategoriaParticipacion.ESTUDIANTE, matriz_origen="M2",
        hoja="Estudiantes", fila=2, actividad_asociada="Actividad Prueba",
        sede="Bilwi", programa="Programa M2 Diferente",
    )
    act = ActividadConsolidada(
        clave_negocio=("actividad prueba", "bilwi", None),
        nombre_actividad="Actividad Prueba", sede="Bilwi",
        fuente_m1=fila_m1, programa="Programa M1 Oficial",
        participaciones=[p_nom],
    )
    cons = ResultadoConsolidacion(
        actividades=[act], fuentes={}, total_actividades=1,
        total_asistencia_bruta=1, total_personas_unicas=1,
        total_recurrencia=0, tasa_recurrencia=1.0,
        periodo=periodo_septiembre_2026,
    )

    informe = InstitutionalReportTransformer.transformar(
        consolidacion=cons,
        departamento_responsable=None,
        metadatos=None,
        politica_presentacion="DECLARADO_M1",
    )
    assert informe.departamento_responsable == "Programa M1 Oficial"
    assert informe.trazabilidad_global.get("departamento_responsable_origen") == "NIVEL_2_M1"


# ============================================================================
# TEST 4: M1 VACÍO + M2/M3 COINCIDENTES (NIVEL 3)
# ============================================================================

def test_04_m1_vacio_m2_m3_coincidentes(periodo_septiembre_2026):
    """
    Test 4: M1 vacío. M2: 'Innovación y Emprendimiento', M3: 'Innovación y Emprendimiento'.
    Resultado: 'Innovación y Emprendimiento' inferido de contexto nominal unívoco.
    """
    fila_m1 = FilaLeidaM1(
        hoja="Consolidado", fila=2, sede="Bilwi",
        actividad="Taller Robótica",
        programa=None, area_responsable=None, nombre_area=None,
    )
    p2 = ParticipacionConsolidada(
        id_actividad="act-1", id_persona="per-1", nombre_apellidos="Estudiante A",
        categoria=CategoriaParticipacion.ESTUDIANTE, matriz_origen="M2",
        hoja="Estudiantes", fila=2, actividad_asociada="Taller Robótica",
        sede="Bilwi", programa="Innovación y Emprendimiento",
    )
    p3 = ParticipacionConsolidada(
        id_actividad="act-1", id_persona="per-2", nombre_apellidos="Docente B",
        categoria=CategoriaParticipacion.ADMINISTRATIVO, matriz_origen="M3",
        hoja="Admin", fila=2, actividad_asociada="Taller Robótica",
        sede="Bilwi", programa="Innovación y Emprendimiento",
    )
    act = ActividadConsolidada(
        clave_negocio=("taller robotica", "bilwi", None),
        nombre_actividad="Taller Robótica", sede="Bilwi",
        fuente_m1=fila_m1, programa=None, area_responsable=None,
        participaciones=[p2, p3],
    )
    cons = ResultadoConsolidacion(
        actividades=[act], fuentes={}, total_actividades=1,
        total_asistencia_bruta=2, total_personas_unicas=2,
        total_recurrencia=0, tasa_recurrencia=1.0,
        periodo=periodo_septiembre_2026,
    )

    informe = InstitutionalReportTransformer.transformar(
        consolidacion=cons,
        departamento_responsable=None,
        metadatos=None,
        politica_presentacion="DECLARADO_M1",
    )
    assert informe.departamento_responsable == "Innovación y Emprendimiento"
    assert informe.trazabilidad_global.get("departamento_responsable_origen") == "NIVEL_3_CONTEXTO_NOMINAL"
    assert informe.trazabilidad_global.get("procedencia_departamento") == "M2/M3_NOMINAL_UNIVOCO"


# ============================================================================
# TEST 5: M1 VACÍO + SOLO M2 INFORMADO (NIVEL 3)
# ============================================================================

def test_05_m1_vacio_solo_m2_informado(periodo_septiembre_2026):
    """
    Test 5: M1 vacío + solo M2 informado. Debe resolver el valor unívoco asociado.
    """
    fila_m1 = FilaLeidaM1(
        hoja="Consolidado", fila=2, sede="Bilwi",
        actividad="Charla IA",
        programa=None, area_responsable=None, nombre_area=None,
    )
    p2 = ParticipacionConsolidada(
        id_actividad="act-1", id_persona="per-1", nombre_apellidos="Estudiante A",
        categoria=CategoriaParticipacion.ESTUDIANTE, matriz_origen="M2",
        hoja="Estudiantes", fila=2, actividad_asociada="Charla IA",
        sede="Bilwi", programa="Innovación y Emprendimiento",
    )
    act = ActividadConsolidada(
        clave_negocio=("charla ia", "bilwi", None),
        nombre_actividad="Charla IA", sede="Bilwi",
        fuente_m1=fila_m1, programa=None, area_responsable=None,
        participaciones=[p2],
    )
    cons = ResultadoConsolidacion(
        actividades=[act], fuentes={}, total_actividades=1,
        total_asistencia_bruta=1, total_personas_unicas=1,
        total_recurrencia=0, tasa_recurrencia=1.0,
        periodo=periodo_septiembre_2026,
    )

    informe = InstitutionalReportTransformer.transformar(
        consolidacion=cons,
        politica_presentacion="DECLARADO_M1",
    )
    assert informe.departamento_responsable == "Innovación y Emprendimiento"
    assert informe.trazabilidad_global.get("departamento_responsable_origen") == "NIVEL_3_CONTEXTO_NOMINAL"


# ============================================================================
# TEST 6: M1 VACÍO + SOLO M3 INFORMADO (NIVEL 3)
# ============================================================================

def test_06_m1_vacio_solo_m3_informado(periodo_septiembre_2026):
    """
    Test 6: M1 vacío + solo M3 informado. Debe resolver el valor unívoco asociado.
    """
    fila_m1 = FilaLeidaM1(
        hoja="Consolidado", fila=2, sede="Bilwi",
        actividad="Capacitacion Docente",
        programa=None, area_responsable=None, nombre_area=None,
    )
    p3 = ParticipacionConsolidada(
        id_actividad="act-1", id_persona="per-1", nombre_apellidos="Funcionario X",
        categoria=CategoriaParticipacion.ADMINISTRATIVO, matriz_origen="M3",
        hoja="Admin", fila=2, actividad_asociada="Capacitacion Docente",
        sede="Bilwi", programa="Innovación y Emprendimiento",
    )
    act = ActividadConsolidada(
        clave_negocio=("capacitacion docente", "bilwi", None),
        nombre_actividad="Capacitacion Docente", sede="Bilwi",
        fuente_m1=fila_m1, programa=None, area_responsable=None,
        participaciones=[p3],
    )
    cons = ResultadoConsolidacion(
        actividades=[act], fuentes={}, total_actividades=1,
        total_asistencia_bruta=1, total_personas_unicas=1,
        total_recurrencia=0, tasa_recurrencia=1.0,
        periodo=periodo_septiembre_2026,
    )

    informe = InstitutionalReportTransformer.transformar(
        consolidacion=cons,
        politica_presentacion="DECLARADO_M1",
    )
    assert informe.departamento_responsable == "Innovación y Emprendimiento"
    assert informe.trazabilidad_global.get("departamento_responsable_origen") == "NIVEL_3_CONTEXTO_NOMINAL"


# ============================================================================
# TEST 7: M2/M3 CON VALORES DIFERENTES — AMBIGÜEDAD Y FALLBACK
# ============================================================================

def test_07_m2_m3_valores_divergentes_ambiguedad(periodo_septiembre_2026):
    """
    Test 7: M2 = 'Programa Alpha', M3 = 'Programa Beta'.
    NO elegir arbitrariamente. Debe documentar discrepancia/observación
    y utilizar el fallback 'Área No Especificada'.
    """
    fila_m1 = FilaLeidaM1(
        hoja="Consolidado", fila=2, sede="Bilwi",
        actividad="Actividad Ambigua",
        programa=None, area_responsable=None, nombre_area=None,
    )
    p2 = ParticipacionConsolidada(
        id_actividad="act-1", id_persona="per-1", nombre_apellidos="Estudiante A",
        categoria=CategoriaParticipacion.ESTUDIANTE, matriz_origen="M2",
        hoja="Estudiantes", fila=2, actividad_asociada="Actividad Ambigua",
        sede="Bilwi", programa="Programa Alpha",
    )
    p3 = ParticipacionConsolidada(
        id_actividad="act-1", id_persona="per-2", nombre_apellidos="Docente B",
        categoria=CategoriaParticipacion.ADMINISTRATIVO, matriz_origen="M3",
        hoja="Admin", fila=2, actividad_asociada="Actividad Ambigua",
        sede="Bilwi", programa="Programa Beta",
    )
    act = ActividadConsolidada(
        clave_negocio=("actividad ambigua", "bilwi", None),
        nombre_actividad="Actividad Ambigua", sede="Bilwi",
        fuente_m1=fila_m1, programa=None, area_responsable=None,
        participaciones=[p2, p3],
    )
    cons = ResultadoConsolidacion(
        actividades=[act], fuentes={}, total_actividades=1,
        total_asistencia_bruta=2, total_personas_unicas=2,
        total_recurrencia=0, tasa_recurrencia=1.0,
        periodo=periodo_septiembre_2026,
    )

    informe = InstitutionalReportTransformer.transformar(
        consolidacion=cons,
        politica_presentacion="DECLARADO_M1",
    )
    # No elige arbitrariamente: usa fallback
    assert informe.departamento_responsable == "Área No Especificada"
    assert informe.trazabilidad_global.get("departamento_responsable_origen") == "NIVEL_5_FALLBACK"

    # Se registró observación de discrepancia por ambigüedad
    assert informe.observaciones is not None
    assert "DISCREPANCIA INSTITUCIONAL" in informe.observaciones
    assert "Programa Alpha" in informe.observaciones
    assert "Programa Beta" in informe.observaciones


# ============================================================================
# TEST 8: TODO VACÍO — MANTIENE FALLBACK (NIVEL 5)
# ============================================================================

def test_08_todo_vacio_mantiene_fallback(periodo_septiembre_2026):
    """
    Test 8: M1 vacío, M2 vacío, M3 vacío.
    Debe mantenerse el fallback oficial 'Área No Especificada'.
    """
    fila_m1 = FilaLeidaM1(
        hoja="Consolidado", fila=2, sede="Bilwi",
        actividad="Actividad Sin Datos",
        programa=None, area_responsable=None, nombre_area=None,
    )
    act = ActividadConsolidada(
        clave_negocio=("actividad sin datos", "bilwi", None),
        nombre_actividad="Actividad Sin Datos", sede="Bilwi",
        fuente_m1=fila_m1, programa=None, area_responsable=None,
        participaciones=[],
    )
    cons = ResultadoConsolidacion(
        actividades=[act], fuentes={}, total_actividades=1,
        total_asistencia_bruta=0, total_personas_unicas=0,
        total_recurrencia=0, tasa_recurrencia=1.0,
        periodo=periodo_septiembre_2026,
    )

    informe = InstitutionalReportTransformer.transformar(
        consolidacion=cons,
        politica_presentacion="DECLARADO_M1",
    )
    assert informe.departamento_responsable == "Área No Especificada"
    assert informe.trazabilidad_global.get("departamento_responsable_origen") == "NIVEL_5_FALLBACK"


# ============================================================================
# TEST 9: PRODUCTO A SIN REGRESIONES
# ============================================================================

def test_09_product_a_no_alterado(tmp_path, dir_output_real, periodo_septiembre_2026):
    """
    Test 9: La corrección NO debe alterar Product A.
    Verifica que el pipeline genera el Informe Técnico con éxito y estructura íntegra.
    """
    pipeline = WordConsolidationPipeline(output_dir=tmp_path)
    res = pipeline.ejecutar(
        fuentes=dir_output_real,
        periodo=periodo_septiembre_2026,
        generar_tecnico=True,
        generar_institucional=False,
    )
    assert res.status_tecnico == "SUCCESS"
    assert res.status_institucional == "NOT_REQUESTED"
    assert res.ruta_docx is not None
    p_tec = Path(res.ruta_docx)
    assert p_tec.exists()

    doc = Document(p_tec)
    textos = [p.text for p in doc.paragraphs if p.text.strip()]
    assert any("INFORME CONSOLIDADO INSTITUCIONAL DE ACTIVIDADES" in t for t in textos)
    assert res.total_actividades == 1
    assert res.total_participaciones == 18
    assert res.total_personas_unicas == 18
    assert res.total_filas_activas == 19


# ============================================================================
# TEST 10: M5 32 REGISTROS HISTÓRICOS INTACTOS
# ============================================================================

def test_10_m5_32_registros_historicos_intactos(dir_output_real, periodo_septiembre_2026):
    """
    Test 10: Los 32 registros históricos de M5 fuera de período deben mantenerse intactos.
    """
    m5_path = dir_output_real / "Matriz_5_Protagonistas_Beneficiados.xlsx"
    filas_m5, _ = MatrixReader.leer_matriz_nominal(ruta_archivo=m5_path, codigo_matriz="M5")
    assert len(filas_m5) == 32

    # Al filtrar por período septiembre 2026 semana 1, las 32 filas quedan fuera
    conjunto = ConjuntoMatricesLeidas(beneficiarios_m5=filas_m5)
    filtrado = PeriodFilter.filtrar(conjunto, periodo_septiembre_2026)
    assert len(filtrado.beneficiarios_en_periodo) == 0
    assert len(filtrado.nominales_historicos_fuera_periodo) == 32


# ============================================================================
# TEST 11: INMUTABILIDAD CRIPTOGRÁFICA DE LOS CINCO EXCEL DE ENTRADA
# ============================================================================

def test_11_inmutabilidad_cinco_excel_entrada(tmp_path, dir_output_real, periodo_septiembre_2026):
    """
    Test 11: Los 5 Excel de entrada deben permanecer byte por byte idénticos
    antes y después de la ejecución completa.
    """
    archivos = sorted(list(dir_output_real.glob("Matriz_*.xlsx")))
    assert len(archivos) == 5

    hashes_antes = {p.name: calcular_sha256(p) for p in archivos}

    pipeline = WordConsolidationPipeline(output_dir=tmp_path)
    res = pipeline.ejecutar(
        fuentes=dir_output_real,
        periodo=periodo_septiembre_2026,
        generar_tecnico=True,
        generar_institucional=True,
    )
    assert "EXITOSO" in res.estado_final

    hashes_despues = {p.name: calcular_sha256(p) for p in archivos}

    for nombre, sha_pre in hashes_antes.items():
        sha_post = hashes_despues[nombre]
        assert sha_pre == sha_post, f"Inmutabilidad violada en {nombre}: {sha_pre} != {sha_post}"


# ============================================================================
# TEST 12: DETERMINISMO ESTRICTO
# ============================================================================

def test_12_determinismo_ejecuciones_identicas(tmp_path, dir_output_real, periodo_septiembre_2026):
    """
    Test 12: Dos ejecuciones sucesivas con los mismos inputs deben producir
    exactamente los mismos resultados institucionales.
    """
    out1 = tmp_path / "run1"
    out2 = tmp_path / "run2"

    p1 = WordConsolidationPipeline(output_dir=out1)
    res1 = p1.ejecutar(
        fuentes=dir_output_real,
        periodo=periodo_septiembre_2026,
        generar_tecnico=True,
        generar_institucional=True,
    )

    p2 = WordConsolidationPipeline(output_dir=out2)
    res2 = p2.ejecutar(
        fuentes=dir_output_real,
        periodo=periodo_septiembre_2026,
        generar_tecnico=True,
        generar_institucional=True,
    )

    assert res1.total_actividades == res2.total_actividades
    assert res1.total_filas_activas == res2.total_filas_activas

    # Verificar texto idéntico en cabeceras de tablas de Product B
    doc1 = Document(res1.ruta_docx_institucional)
    doc2 = Document(res2.ruta_docx_institucional)

    t1 = doc1.tables[0].rows[0].cells[1].text
    t2 = doc2.tables[0].rows[0].cells[1].text
    assert t1 == t2
    assert "Innovación y Emprendimiento" in t1


# ============================================================================
# TESTS ADICIONALES (A, B, C, D, E) EXIGIDOS EN LA APROBACIÓN
# ============================================================================

def test_adicional_a_dos_actividades_distintas_no_se_mezclan(periodo_septiembre_2026):
    """
    Test Adicional A: Dos actividades diferentes:
    Actividad A -> Programa A
    Actividad B -> Programa B
    Verificar que el contexto nominal no se contamina entre actividades.
    """
    act_a = ActividadConsolidada(
        clave_negocio=("actividad a", "bilwi", None),
        nombre_actividad="Actividad A", sede="Bilwi",
        participaciones=[
            ParticipacionConsolidada(
                id_actividad="act-a", id_persona="p-1", nombre_apellidos="Persona 1",
                categoria=CategoriaParticipacion.ESTUDIANTE, matriz_origen="M2",
                hoja="H1", fila=2, actividad_asociada="Actividad A", sede="Bilwi",
                programa="Programa Alpha",
            )
        ],
    )
    act_b = ActividadConsolidada(
        clave_negocio=("actividad b", "bilwi", None),
        nombre_actividad="Actividad B", sede="Bilwi",
        participaciones=[
            ParticipacionConsolidada(
                id_actividad="act-b", id_persona="p-2", nombre_apellidos="Persona 2",
                categoria=CategoriaParticipacion.ESTUDIANTE, matriz_origen="M2",
                hoja="H1", fila=3, actividad_asociada="Actividad B", sede="Bilwi",
                programa="Programa Beta",
            )
        ],
    )

    progs_a = InstitutionalReportTransformer._obtener_programas_nominales_actividad(act_a)
    progs_b = InstitutionalReportTransformer._obtener_programas_nominales_actividad(act_b)

    assert progs_a == ["Programa Alpha"]
    assert progs_b == ["Programa Beta"]
    assert "Programa Beta" not in progs_a
    assert "Programa Alpha" not in progs_b


def test_adicional_b_misma_actividad_dos_periodos_aislamiento():
    """
    Test Adicional B: Misma actividad en dos períodos diferentes con programas diferentes.
    Verificar aislamiento por período sin fuga de metadatos.
    """
    p_sept = PeriodoConsolidacion(
        tipo_periodo=TipoPeriodo.MES, anio=2026, mes=9, etiqueta="Septiembre 2026"
    )
    p_oct = PeriodoConsolidacion(
        tipo_periodo=TipoPeriodo.MES, anio=2026, mes=10, etiqueta="Octubre 2026"
    )

    # 1. Aislamiento estricto de fechas por período
    assert PeriodFilter.evaluar_fecha("2026-09-10", p_sept) is True
    assert PeriodFilter.evaluar_fecha("2026-10-15", p_sept) is False
    assert PeriodFilter.evaluar_fecha("2026-09-10", p_oct) is False
    assert PeriodFilter.evaluar_fecha("2026-10-15", p_oct) is True

    # 2. Corrida Septiembre: Misma actividad 'Taller Continuo', pero contexto nominal es 'Programa Septiembre'
    part_sept = ParticipacionConsolidada(
        id_actividad="act-sept", id_persona="p-1", nombre_apellidos="Estudiante 1",
        categoria=CategoriaParticipacion.ESTUDIANTE, matriz_origen="M2",
        hoja="H1", fila=2, actividad_asociada="Taller Continuo", sede="Bilwi",
        programa="Programa Septiembre",
    )
    act_sept = ActividadConsolidada(
        clave_negocio=("taller continuo", "bilwi", "2026-09-10"),
        nombre_actividad="Taller Continuo", sede="Bilwi",
        participaciones=[part_sept],
    )
    cons_sept = ResultadoConsolidacion(
        actividades=[act_sept], fuentes={}, total_actividades=1,
        total_asistencia_bruta=1, total_personas_unicas=1,
        total_recurrencia=0, tasa_recurrencia=1.0,
        periodo=p_sept,
    )

    # 3. Corrida Octubre: Misma actividad 'Taller Continuo', pero contexto nominal es 'Programa Octubre'
    part_oct = ParticipacionConsolidada(
        id_actividad="act-oct", id_persona="p-2", nombre_apellidos="Estudiante 2",
        categoria=CategoriaParticipacion.ESTUDIANTE, matriz_origen="M2",
        hoja="H1", fila=2, actividad_asociada="Taller Continuo", sede="Bilwi",
        programa="Programa Octubre",
    )
    act_oct = ActividadConsolidada(
        clave_negocio=("taller continuo", "bilwi", "2026-10-15"),
        nombre_actividad="Taller Continuo", sede="Bilwi",
        participaciones=[part_oct],
    )
    cons_oct = ResultadoConsolidacion(
        actividades=[act_oct], fuentes={}, total_actividades=1,
        total_asistencia_bruta=1, total_personas_unicas=1,
        total_recurrencia=0, tasa_recurrencia=1.0,
        periodo=p_oct,
    )

    # 4. Verificar aislamiento en resolución de departamento (sin contaminación cruzada)
    depto_sept, origen_sept, obs_sept = InstitutionalReportTransformer.resolver_departamento_responsable(cons_sept)
    depto_oct, origen_oct, obs_oct = InstitutionalReportTransformer.resolver_departamento_responsable(cons_oct)

    assert depto_sept == "Programa Septiembre"
    assert origen_sept == "NIVEL_3_CONTEXTO_NOMINAL"
    assert "Programa Septiembre" in obs_sept

    assert depto_oct == "Programa Octubre"
    assert origen_oct == "NIVEL_3_CONTEXTO_NOMINAL"
    assert "Programa Octubre" in obs_oct

    # No hay contaminación entre períodos
    assert depto_sept != depto_oct



def test_adicional_c_m2_un_programa_m3_otro_misma_actividad_ambiguedad(periodo_septiembre_2026):
    """
    Test Adicional C: M2 tiene un programa y M3 otro para la misma actividad.
    Debe detectarse ambigüedad formalmente y reflejarse en las observaciones.
    """
    act = ActividadConsolidada(
        clave_negocio=("actividad mixta", "bilwi", None),
        nombre_actividad="Actividad Mixta", sede="Bilwi",
        participaciones=[
            ParticipacionConsolidada(
                id_actividad="act-1", id_persona="p-1", nombre_apellidos="Estudiante 1",
                categoria=CategoriaParticipacion.ESTUDIANTE, matriz_origen="M2",
                hoja="H1", fila=2, actividad_asociada="Actividad Mixta", sede="Bilwi",
                programa="Programa Ingenieria",
            ),
            ParticipacionConsolidada(
                id_actividad="act-1", id_persona="p-2", nombre_apellidos="Admin 1",
                categoria=CategoriaParticipacion.ADMINISTRATIVO, matriz_origen="M3",
                hoja="H1", fila=2, actividad_asociada="Actividad Mixta", sede="Bilwi",
                programa="Programa Administracion",
            ),
        ],
    )
    cons = ResultadoConsolidacion(
        actividades=[act], fuentes={}, total_actividades=1,
        total_asistencia_bruta=2, total_personas_unicas=2,
        total_recurrencia=0, tasa_recurrencia=1.0,
        periodo=periodo_septiembre_2026,
    )

    depto, origen, obs = InstitutionalReportTransformer.resolver_departamento_responsable(cons)
    assert depto == "Área No Especificada"
    assert origen == "NIVEL_5_FALLBACK"
    assert obs is not None
    assert "DISCREPANCIA INSTITUCIONAL" in obs
    assert "Programa Administracion" in obs
    assert "Programa Ingenieria" in obs


def test_adicional_d_m2_multiples_filas_mismo_programa_sigue_univoco(periodo_septiembre_2026):
    """
    Test Adicional D: M2 tiene 10 filas con el mismo programa.
    Debe seguir considerándose un único valor inequívoco.
    """
    parts = [
        ParticipacionConsolidada(
            id_actividad="act-1", id_persona=f"p-{i}", nombre_apellidos=f"Estudiante {i}",
            categoria=CategoriaParticipacion.ESTUDIANTE, matriz_origen="M2",
            hoja="H1", fila=i + 2, actividad_asociada="Taller Unico", sede="Bilwi",
            programa="Innovación y Emprendimiento",
        )
        for i in range(10)
    ]
    act = ActividadConsolidada(
        clave_negocio=("taller unico", "bilwi", None),
        nombre_actividad="Taller Unico", sede="Bilwi",
        participaciones=parts,
    )
    cons = ResultadoConsolidacion(
        actividades=[act], fuentes={}, total_actividades=1,
        total_asistencia_bruta=10, total_personas_unicas=10,
        total_recurrencia=0, tasa_recurrencia=1.0,
        periodo=periodo_septiembre_2026,
    )

    depto, origen, obs = InstitutionalReportTransformer.resolver_departamento_responsable(cons)
    assert depto == "Innovación y Emprendimiento"
    assert origen == "NIVEL_3_CONTEXTO_NOMINAL"
    assert obs is not None
    assert "inequívoco" in obs


def test_adicional_e_una_actividad_con_nominales_otra_sin_nominales_no_reutiliza(periodo_septiembre_2026):
    """
    Test Adicional E: Actividad 1 tiene registros nominales con 'Innovación y Emprendimiento'.
    Actividad 2 no tiene registros nominales.
    Actividad 2 no debe heredar ni reutilizar el programa de Actividad 1 en su trazabilidad propia.
    """
    act1 = ActividadConsolidada(
        clave_negocio=("actividad con nominales", "bilwi", None),
        nombre_actividad="Actividad Con Nominales", sede="Bilwi",
        participaciones=[
            ParticipacionConsolidada(
                id_actividad="act-1", id_persona="p-1", nombre_apellidos="Estudiante 1",
                categoria=CategoriaParticipacion.ESTUDIANTE, matriz_origen="M2",
                hoja="H1", fila=2, actividad_asociada="Actividad Con Nominales", sede="Bilwi",
                programa="Innovación y Emprendimiento",
            )
        ],
    )
    act2 = ActividadConsolidada(
        clave_negocio=("actividad sin nominales", "bilwi", None),
        nombre_actividad="Actividad Sin Nominales", sede="Bilwi",
        participaciones=[],
    )

    progs1 = InstitutionalReportTransformer._obtener_programas_nominales_actividad(act1)
    progs2 = InstitutionalReportTransformer._obtener_programas_nominales_actividad(act2)

    assert progs1 == ["Innovación y Emprendimiento"]
    assert progs2 == []
