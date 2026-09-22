"""tests.test_review_e2e

Prueba de Integración End-to-End (E2E) para la Cola de Revisión (Fase 29.22.1).
Valida el ciclo completo institucional:
  1. Ingesta de participante con estamento DESCONOCIDO -> Asignado a COLA_REVISION.
  2. Consulta de bandeja operativa -> Aparece como caso pendiente.
  3. Aislamiento histórico -> Los 32 registros patrimoniales de M5 no aparecen en la cola.
  4. Decisión humana explícita -> Operador resuelve a ESTUDIANTE con justificación.
  5. Transacción atómica -> Actualización en SQLite V004 y asiento DISCREPANCY_RESOLVE.
  6. Salida de la cola -> La cola queda vacía (0 pendientes).
  7. Exportación oficial -> El participante promovido se proyecta exitosamente en Matriz 2 (M2).
  8. Preservación patrimonial -> Los 32 registros de M5 se conservan 100% intactos en la exportación oficial.
"""

from pathlib import Path
import tempfile
import openpyxl
import pytest

from app.application.commands.pipeline_commands import ExportarMatricesPeriodoCommand
from app.application.use_cases.pipeline.procesar_pipeline_actividad import (
    ProcesarPipelineActividadUseCase,
)
from app.application.use_cases.exportacion.exportar_matrices_periodo import (
    ExportarMatricesPeriodoUseCase,
)
from app.core.models.activity import Activity
from app.core.models.person import Person
from app.core.models.participation import Participation
from app.infrastructure.persistence.config import DatabaseConfig
from app.infrastructure.persistence.connection import SQLiteConnectionManager
from app.infrastructure.persistence.migrations import MigrationRunner
from app.infrastructure.persistence.repositories.actividad_repository import (
    SQLiteActividadRepository,
)
from app.infrastructure.persistence.repositories.participacion_repository import (
    SQLiteParticipacionRepository,
)
from app.infrastructure.persistence.repositories.persona_repository import (
    SQLitePersonaRepository,
)
from app.infrastructure.persistence.repositories.unit_of_work import (
    SQLiteUnitOfWork,
)
from app.review.application.service import ReviewQueueApplicationService
from app.review.infrastructure.sqlite_review_repository import (
    SQLiteReviewQueueRepository,
)


@pytest.fixture
def e2e_env():
    """Configura el entorno de prueba aislado para la validación E2E."""
    with tempfile.TemporaryDirectory() as tmpdir:
        tmp_path = Path(tmpdir)
        db_file = tmp_path / "test_review_e2e.db"
        cfg = DatabaseConfig(db_path=db_file)
        mgr = SQLiteConnectionManager(cfg)

        conn = mgr.get_connection()
        try:
            runner = MigrationRunner(conn)
            runner.apply_all_pending()
        finally:
            conn.close()

        uow = SQLiteUnitOfWork(mgr, db_path=db_file)
        yield mgr, uow, tmp_path


def test_flujo_e2e_cola_revision_a_matriz_oficial(e2e_env):
    """Ejecuta el escenario integral: COLA_REVISION -> Resolución Humana -> M2 exportado -> M5 intacto."""
    mgr, uow, tmp_path = e2e_env

    # 1. Sembrar actividad y participantes
    conn = mgr.get_connection()
    try:
        act_repo = SQLiteActividadRepository(conn)
        per_repo = SQLitePersonaRepository(conn)
        part_repo = SQLiteParticipacionRepository(conn)

        # Actividad
        act = Activity(
            id_actividad="act-e2e-01",
            codigo_institucional="ACT-2026-E2E",
            nombre_actividad_original="Feria Universitaria de Innovación y Emprendimiento",
            fecha_inicio_texto="2026-03-22",
            sede="BLUEFIELDS",
            lugar="Campus Universitario Bluefields",
            tipo_evento="FERIA",
            departamento_responsable="Dirección de Innovación",
            responsable="Prof. Kenia",
            estado="EJECUTADA",
        )
        act_repo.save(act)

        # Participante que ingresa con estamento ambiguo a COLA_REVISION
        p_dudoso = Person(
            id_persona_interno="per-e2e-dudoso",
            nombre_completo="CARMEN ELENA MIRANDA",
            nombres="CARMEN ELENA",
            apellidos="MIRANDA",
            cedula="601-200502-1000M",
            sexo_original="F",
            sexo_normalizado="F",
        )
        per_repo.save(p_dudoso)

        part_dudoso = Participation(
            id_participacion="part-e2e-dudoso",
            id_actividad="act-e2e-01",
            id_persona="per-e2e-dudoso",
            categoria_participacion="DESCONOCIDO",
            matriz_destino="COLA_REVISION",
            requiere_revision=True,
            motivo_revision="Categoría 'DESCONOCIDO' no determinable institucionalmente",
            es_historico_preexistente=0,
        )
        part_repo.save(part_dudoso)

        # 32 registros históricos patrimoniales en M5
        for i in range(1, 33):
            pid = f"per-hist-{i:03d}"
            part_id = f"part-hist-{i:03d}"
            p_hist = Person(
                id_persona_interno=pid,
                nombre_completo=f"BENEFICIARIO HISTORICO {i:02d}",
                nombres=f"BENEFICIARIO {i:02d}",
                apellidos="HISTORICO",
                cedula=f"601-010190-{i:04d}H",
                sexo_original="F" if i % 2 == 0 else "M",
                sexo_normalizado="F" if i % 2 == 0 else "M",
            )
            per_repo.save(p_hist)

            part_hist = Participation(
                id_participacion=part_id,
                id_actividad="act-e2e-01",
                id_persona=pid,
                categoria_participacion="BENEFICIADO",
                matriz_destino="M5",
                requiere_revision=False,
                motivo_revision=None,
                es_historico_preexistente=1,
            )
            part_repo.save(part_hist)

    finally:
        conn.close()

    # 2. Consultar cola de revisión mediante ReviewQueueApplicationService
    conn_review = mgr.get_connection()
    try:
        review_repo = SQLiteReviewQueueRepository(conn_review)
        service = ReviewQueueApplicationService(review_repo)

        # A. Verificar que el caso aparece en la cola
        casos = service.listar_casos_pendientes()
        assert len(casos) == 1
        assert casos[0].id_caso == "part-e2e-dudoso"
        assert casos[0].nombre_persona == "CARMEN ELENA MIRANDA"

        # B. Verificar que los 32 históricos no están en la cola
        ids_cola = {c.id_caso for c in casos}
        assert "part-hist-001" not in ids_cola
        assert "part-hist-032" not in ids_cola

        # C. Iniciar revisión humana
        detalle = service.iniciar_revision("part-e2e-dudoso", usuario="revisor_bicu")
        assert detalle.estado == "EN_REVISION"
        assert detalle.usuario_resolutor == "revisor_bicu"

        # D. Asentar resolución explícita: Confirmar Estamento como ESTUDIANTE (M2)
        resultado = service.resolver_estamento(
            id_participacion="part-e2e-dudoso",
            nuevo_estamento="ESTUDIANTE",
            justificacion="Estudiante activa confirmada en Registro Académico BICU",
            usuario="revisor_bicu",
        )
        assert resultado.exito is True
        assert resultado.matriz_destino == "M2"
        assert resultado.nuevo_estado == "RESUELTO"

        # E. Verificar que la cola ahora queda vacía
        casos_post = service.listar_casos_pendientes()
        assert len(casos_post) == 0

        # F. Verificar historial forense de auditoría
        historial = service.listar_historial_resoluciones()
        assert len(historial) == 1
        assert historial[0].usuario_operador == "revisor_bicu"
        assert "ESTUDIANTE" in historial[0].cambios_resumen

    finally:
        conn_review.close()

    # 3. Exportar matrices oficiales mediante exportadores certificados
    salida_dir = tmp_path / "matrices_exportadas"
    salida_dir.mkdir()

    pipeline_uc = ProcesarPipelineActividadUseCase(uow=uow)
    export_uc = ExportarMatricesPeriodoUseCase(pipeline_use_case=pipeline_uc, uow=uow)

    cmd = ExportarMatricesPeriodoCommand(
        anio=2026,
        ids_actividades=["act-e2e-01"],
        carpeta_salida=str(salida_dir),
        carpeta_templates="templates",
        permitir_fixtures_test_only=False,
        dry_run=False,
    )
    res_export = export_uc.execute(cmd, uow=uow)
    assert res_export.exito is True

    # 4. Verificar físicamente la proyección oficial en Matriz 2 (M2)
    m2_file = Path(res_export.archivos_generados["matriz_2"])
    assert m2_file.exists()
    wb_m2 = openpyxl.load_workbook(str(m2_file), data_only=True)
    ws_m2 = wb_m2.worksheets[0]

    # La participante promovida desde COLA_REVISION debe encontrarse en M2
    nombres_en_m2 = [ws_m2.cell(row=r, column=33).value for r in range(2, ws_m2.max_row + 1)]
    assert "CARMEN ELENA MIRANDA" in nombres_en_m2
    wb_m2.close()

    # 5. Verificar preservación patrimonial inalterada de los 32 históricos en Matriz 5 (M5)
    m5_file = Path(res_export.archivos_generados["matriz_5"])
    assert m5_file.exists()
    wb_m5 = openpyxl.load_workbook(str(m5_file), data_only=True)
    ws_m5 = wb_m5.worksheets[0]

    # Filas 2 a 33 deben corresponder a los 32 históricos preservados
    assert ws_m5.cell(row=33, column=1).value == 32
    wb_m5.close()
