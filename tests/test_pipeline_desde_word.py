"""tests.test_pipeline_desde_word

Suite de pruebas para el Subbloque 3.2:
ProcesarPipelineDesdeWordUseCase (Ingesta + Pipeline para un DOCX individual).

Cobertura de pruebas:
- T-01: DOCX válido → Ingesta + Pipeline completan exitosamente.
- T-02: DOCX con participantes nominales ausentes → total_participaciones == 0 esperado (Decisión P-02).
- T-03: DOCX inválido / no compatible → ingesta rechaza, pipeline no se ejecuta.
- T-04: Documento informe semanal → rechazado con motivo explicativo, pipeline no se ejecuta.
- T-05: Excepción en pipeline institucional → capturada de forma segura, error registrado, no crash.
- T-06: Modo dry_run=True → Cero mutaciones persistentes en SQLite; pipeline se evalúa en memoria.
- T-07: Heurística GAP-3 → posible_duplicado_advertido no bloquea la ejecución (Decisión P-03).
- T-08: Conversión a BatchArchivoResultadoDTO produce el estado correcto (ACEPTADO, CON_ADVERTENCIAS, etc.).
- T-09: Aislamiento arquitectónico (AST) → procesar_pipeline_desde_word.py no importa sqlite3, docx ni openpyxl.
"""

import ast
import inspect
from pathlib import Path
from typing import Generator
import pytest

from app.application.dto.batch_dtos import (
    BatchArchivoResultadoDTO,
    PipelineDesdeWordResultDTO,
)
from app.application.dto.export_dtos import PoliticaExportacionRevision
from app.application.fakes.in_memory_uow import InMemoryUnitOfWork
from app.application.use_cases.batch.procesar_pipeline_desde_word import (
    ProcesarPipelineDesdeWordUseCase,
)
from app.application.use_cases.ingestar_actividad_desde_word import (
    IngestarActividadDesdeWordUseCase,
)
from app.application.use_cases.pipeline.procesar_pipeline_actividad import (
    ProcesarPipelineActividadUseCase,
)
from app.infrastructure.persistence.config import DatabaseConfig
from app.infrastructure.persistence.connection import SQLiteConnectionManager
from app.infrastructure.persistence.migrations import MigrationRunner
from app.infrastructure.persistence.repositories.unit_of_work import SQLiteUnitOfWork
from app.infrastructure.word_reader.word_activity_extractor import (
    WordActivityExtractor,
)
from tests.test_word_activity_extractor import (
    _crear_docx_actividad_valido,
)


def _crear_docx_incompatible(path: Path) -> Path:
    import docx
    doc = docx.Document()
    doc.add_heading("Carta de Solicitud de Materiales", level=1)
    doc.add_paragraph("Solicito la compra de papelería para la oficina.")
    doc.save(str(path))
    return path


def _crear_docx_informe_semanal(path: Path) -> Path:
    import docx
    doc = docx.Document()
    doc.add_paragraph("INFORME SEMANAL - Mecanismo institucional de monitoreo y seguimiento operativo.")
    doc.save(str(path))
    return path


# ---------------------------------------------------------------------------
# Fixtures
# ---------------------------------------------------------------------------

@pytest.fixture
def sqlite_test_env(tmp_path) -> Generator[SQLiteUnitOfWork, None, None]:
    """Crea una base de datos SQLite temporal migrada con esquema v001."""
    db_file = tmp_path / "test_pipeline_word.db"
    config = DatabaseConfig(db_path=db_file)
    manager = SQLiteConnectionManager(config)

    conn = manager.get_connection()
    try:
        runner = MigrationRunner(conn)
        runner.apply_all_pending()
    finally:
        conn.close()

    uow = SQLiteUnitOfWork(connection_manager=manager, db_path=db_file)
    yield uow


@pytest.fixture
def extractor() -> WordActivityExtractor:
    return WordActivityExtractor()


@pytest.fixture
def ingesta_use_case(extractor, sqlite_test_env) -> IngestarActividadDesdeWordUseCase:
    return IngestarActividadDesdeWordUseCase(extractor=extractor, uow=sqlite_test_env)


@pytest.fixture
def pipeline_use_case(sqlite_test_env) -> ProcesarPipelineActividadUseCase:
    return ProcesarPipelineActividadUseCase(uow=sqlite_test_env)


@pytest.fixture
def use_case(ingesta_use_case, pipeline_use_case, sqlite_test_env) -> ProcesarPipelineDesdeWordUseCase:
    return ProcesarPipelineDesdeWordUseCase(
        ingesta_use_case=ingesta_use_case,
        pipeline_use_case=pipeline_use_case,
        uow=sqlite_test_env,
    )


# ---------------------------------------------------------------------------
# Tests: Flujo individual de pipeline desde Word
# ---------------------------------------------------------------------------

class TestProcesarPipelineDesdeWordUseCase:

    def test_t01_docx_valido_ingesta_y_pipeline_exitosos(self, tmp_path, use_case, sqlite_test_env):
        """T-01: DOCX válido ejecuta ingesta y pipeline correctamente."""
        docx_path = _crear_docx_actividad_valido(
            tmp_path,
            nombre_archivo="Actividad_Conferencia_BICU.docx",
            actividad="Conferencia de Innovación Tecnológica 2026",
            lugar="Auditorio Central BICU Bluefields",
            fecha="15/09/2026",
            total=50,
            femenino=30,
            masculino=20,
            estudiantes=50,
        )

        resultado = use_case.execute(docx_path)

        assert isinstance(resultado, PipelineDesdeWordResultDTO)
        assert resultado.exitoso is True
        assert resultado.ingesta_exitosa is True
        assert resultado.pipeline_ejecutado is True
        assert resultado.id_actividad is not None
        assert resultado.tipo_documento == "INFORME_ACTIVIDAD"
        assert resultado.calidad_estado is not None
        assert len(resultado.errores) == 0

        # Verificar que la actividad quedó efectivamente en SQLite
        with sqlite_test_env:
            guardada = sqlite_test_env.actividades.get_by_id(resultado.id_actividad)
            assert guardada is not None
            assert guardada.nombre_actividad_original == "Conferencia de Innovación Tecnológica 2026"

    def test_t02_cero_participaciones_nominales_es_esperado_p02(self, tmp_path, use_case):
        """T-02: Actividades extraídas de Word tienen 0 participaciones nominales (P-02)."""
        docx_path = _crear_docx_actividad_valido(
            tmp_path,
            nombre_archivo="Taller_Word.docx",
            actividad="Taller de Liderazgo Estudiantil",
        )

        resultado = use_case.execute(docx_path)

        assert resultado.exitoso is True
        assert resultado.total_participaciones == 0  # Invariante P-02

    def test_t03_docx_incompatible_rechazado_sin_pipeline(self, tmp_path, use_case, sqlite_test_env):
        """T-03: DOCX incompatible se rechaza sin ejecutar el pipeline ni mutar SQLite."""
        docx_path = tmp_path / "Documento_Invalido.docx"
        _crear_docx_incompatible(docx_path)

        resultado = use_case.execute(docx_path)

        assert resultado.exitoso is False
        assert resultado.ingesta_exitosa is False
        assert resultado.pipeline_ejecutado is False
        assert resultado.id_actividad is None
        assert resultado.motivo_rechazo is not None

        # Verificar que SQLite está vacío
        with sqlite_test_env:
            assert sqlite_test_env.actividades.count() == 0

    def test_t04_informe_semanal_rechazado_con_motivo(self, tmp_path, use_case):
        """T-04: Informe Semanal no se procesa como actividad individual."""
        docx_path = tmp_path / "Informe_Semanal_Consolidado.docx"
        _crear_docx_informe_semanal(docx_path)

        resultado = use_case.execute(docx_path)

        assert resultado.exitoso is False
        assert resultado.ingesta_exitosa is False
        assert resultado.pipeline_ejecutado is False
        assert "Semanal" in (resultado.motivo_rechazo or "")

    def test_t05_excepcion_en_pipeline_capturada_segura(self, tmp_path, extractor, sqlite_test_env):
        """T-05: Excepción en el pipeline no crashea el use case y se reporta en errores."""
        docx_path = _crear_docx_actividad_valido(
            tmp_path,
            nombre_archivo="Actividad_Test_Pipeline_Error.docx",
            actividad="Actividad Para Error",
        )

        ingesta_uc = IngestarActividadDesdeWordUseCase(extractor=extractor, uow=sqlite_test_env)

        # Mock o pipeline use case que simule fallo
        class PipelineFallaFake:
            def execute(self, *args, **kwargs):
                raise RuntimeError("Fallo técnico simulado en el motor de calidad")

        failing_uc = ProcesarPipelineDesdeWordUseCase(
            ingesta_use_case=ingesta_uc,
            pipeline_use_case=PipelineFallaFake(),
            uow=sqlite_test_env,
        )

        resultado = failing_uc.execute(docx_path)

        assert resultado.exitoso is False
        assert resultado.ingesta_exitosa is True
        assert resultado.pipeline_ejecutado is False
        assert any("Fallo técnico simulado" in e for e in resultado.errores)

    def test_t06_dry_run_cero_escrituras_en_sqlite(self, tmp_path, use_case, sqlite_test_env):
        """T-06: dry_run=True no persiste en SQLite real pero evalúa pipeline en memoria."""
        docx_path = _crear_docx_actividad_valido(
            tmp_path,
            nombre_archivo="Simulacion_Actividad.docx",
            actividad="Feria de Ciencias - Simulación",
        )

        resultado = use_case.execute(docx_path, dry_run=True)

        assert resultado.exitoso is True
        assert resultado.ingesta_exitosa is True
        assert resultado.pipeline_ejecutado is True
        assert resultado.id_actividad is not None

        # Verificar CERO registros en el SQLite persistente real
        with sqlite_test_env:
            assert sqlite_test_env.actividades.count() == 0
            assert sqlite_test_env.actividades.get_by_id(resultado.id_actividad) is None

    def test_t07_posible_duplicado_gap3_emite_advertencia_sin_bloquear(self, tmp_path, use_case):
        """T-07: Advertencia heurística GAP-3 no bloquea la ejecución (P-03)."""
        docx_path = _crear_docx_actividad_valido(
            tmp_path,
            nombre_archivo="Actividad_Repetida.docx",
            actividad="Capacitación Docente 2026",
        )

        resultado = use_case.execute(
            docx_path,
            posible_duplicado_advertido=True,
        )

        assert resultado.posible_duplicado_advertido is True
        assert any("GAP-3" in w for w in resultado.advertencias)
        assert resultado.exitoso is True  # Continúa y completa

    def test_t08_conversion_a_batch_archivo_resultado(self, tmp_path, use_case):
        """T-08: to_batch_archivo_resultado genera el DTO con el estado correcto."""
        docx_path = _crear_docx_actividad_valido(
            tmp_path,
            nombre_archivo="Actividad_Conversion.docx",
            actividad="Charla Ambiental",
        )

        resultado = use_case.execute(docx_path)
        batch_archivo = resultado.to_batch_archivo_resultado(ruta_archivo=str(docx_path))

        assert isinstance(batch_archivo, BatchArchivoResultadoDTO)
        assert batch_archivo.nombre_archivo == "Actividad_Conversion.docx"
        assert batch_archivo.ruta_archivo == str(docx_path)
        assert batch_archivo.estado in ("ACEPTADO", "CON_ADVERTENCIAS")
        assert batch_archivo.id_actividad == resultado.id_actividad
        assert batch_archivo.total_participaciones_enrutadas == 0

    def test_t09_aislamiento_arquitectonico(self):
        """T-09: procesar_pipeline_desde_word.py no importa sqlite3, docx ni openpyxl."""
        archivo_use_case = Path("app/application/use_cases/batch/procesar_pipeline_desde_word.py")
        arbol = ast.parse(archivo_use_case.read_text(encoding="utf-8"))

        modulos_prohibidos = {"sqlite3", "docx", "openpyxl"}
        modulos_importados = set()

        for nodo in ast.walk(arbol):
            if isinstance(nodo, ast.Import):
                for alias in nodo.names:
                    modulos_importados.add(alias.name.split(".")[0])
            elif isinstance(nodo, ast.ImportFrom) and nodo.module:
                modulos_importados.add(nodo.module.split(".")[0])

        interseccion = modulos_prohibidos & modulos_importados
        assert not interseccion, f"Violación de Clean Architecture: imports prohibidos {interseccion}"
