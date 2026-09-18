"""
tests/test_exporters.py

Suite de pruebas automatizadas para la Fase 9: Motor de Exportación a Excel.
Verifica:
- Bloqueo de --mode export si las plantillas no son OFICIAL_REAL.
- Atomicidad All-or-Nothing ante fallos (cero entregables parciales).
- Verificación criptográfica SHA-256 de plantillas base (inmutabilidad 100%).
- Preservación exacta de fórmulas (identidad, coordenadas, rango, N=0 y N>1).
- Política declarativa de catálogos (TEXTO_LITERAL_NO_HOMOLOGADO vs BLOQUEAR).
- Diferenciación de SIN_FUENTE (opcional vs obligatorio).
- Autorización estricta de escritura dentro de ZonaEscribible.
- Cumplimiento de invariantes de conservación de participaciones (18 = 17 + 1 + 0 + 0).
- Generación completa de manifiesto_exportacion.json en entorno de prueba.
"""

import json
from pathlib import Path
import pytest
import openpyxl

from app.exporters.base_exporter import calcular_sha256
from app.exporters.consolidado_exporter import ConsolidadoExporter
from app.exporters.coordinator import ExportCoordinator
from app.exporters.exceptions import (
    ExportacionBloqueadaError,
    PlantillasOficialesRequeridasError,
    ViolacionInvarianteExportacionError,
)
from app.exporters.models import (
    EstatusPlantilla,
    ModoExportacion,
    ModoExportacionCatalogo,
    PoliticaCatalogo,
    TipoSinFuente,
)
from app.exporters.participantes_exporter import ParticipantesExporter
from app.exporters.reporter import ExporterReporter
from app.matching.activity_matcher import ActivityMatcher
from app.matching.identity_resolver import IdentityResolver
from app.parsers.attendance_parser import AttendanceParser
from app.parsers.base_parser import RawAttendanceSheetData, RawParticipantRow
from app.parsers.word_parser import WordReportParser
from app.routing.participant_router import ParticipantRouter
from app.statistics.statistics_engine import StatisticsEngine
from app.templates_analysis.inspector import TemplateInspector
from app.templates_analysis.mapping_manifesto import MappingManifestoEngine
from app.templates_analysis.models import EsquemaPlantilla, ItemMappingManifesto, TipoMapeoColumna
from tests.template_fixture_factory import TemplateFixtureFactory


@pytest.fixture(scope="module")
def datos_caso_real():
    """Genera los datos institucionales unificados del Caso Real Septiembre 2026 (18 participantes)."""
    datos_word = {
        "actividad": "BICU CUR Bilwi fortalece conocimientos estudiantiles en diseño de logotipos e inteligencia artificial",
        "departamento_institucional": "Innovación y Emprendimiento",
        "mes": "Septiembre",
        "semana": "1",
        "sede": "Bilwi",
        "departamento_territorial": "RACCN",
        "municipio": "Puerto Cabezas",
        "eje": "11.41.67",
        "tipo_actividad": "7. Creatividad, Ciencias, Investigación e Innovación",
        "fecha": "08/09/2026",
        "horario": "10:00 AM – 12:00 PM",
    }
    filas = [
        RawParticipantRow(numero_fila=1, hoja=1, nombre_completo="Johnny Melvin", sexo="M", cedula="601-120502-1001A", edad=24, carrera="Lic. Conta", categoria="Estudiante"),
        RawParticipantRow(numero_fila=2, hoja=1, nombre_completo="Shana Ruiz Watson", sexo="F", cedula="20-04-2008", edad=22, carrera="Contabilidad", categoria="Estudiante"),
        RawParticipantRow(numero_fila=3, hoja=1, nombre_completo="Claudia J Fernandez", sexo="F", cedula="607-221205-1002W", edad=18, carrera="contabilidad", categoria="Estudiante"),
        RawParticipantRow(numero_fila=4, hoja=1, nombre_completo="Besi Wrealy N Taylor", sexo="F", cedula="612-311004-1000B", edad=21, carrera="Lic Contabil.", categoria="Estudiante"),
        RawParticipantRow(numero_fila=5, hoja=1, nombre_completo="Yeison Johnston R.M.", sexo="M", cedula="601-140703-1003C", edad=23, carrera="Ing. sistc.", categoria="Estudiante"),
        RawParticipantRow(numero_fila=6, hoja=1, nombre_completo="Neidy Rivera G", sexo="F", cedula="601-090905-1004D", edad=21, carrera="Sistema", categoria="Estudiante"),
        RawParticipantRow(numero_fila=7, hoja=1, nombre_completo="Elba Wilson Smith", sexo="F", cedula="601-300101-1005E", edad=25, carrera="A. estudiantil", categoria="No Docente", area="ACCES"),
        RawParticipantRow(numero_fila=8, hoja=2, nombre_completo="Keyra Rolack G", sexo="F", cedula="601-180304-1006F", edad=22, carrera="Contabilidad", categoria="Estudiante"),
        RawParticipantRow(numero_fila=9, hoja=2, nombre_completo="Luz Marina Alan A", sexo="F", cedula="23/06/2008", edad=23, carrera="Contabilidad", categoria="Estudiante"),
        RawParticipantRow(numero_fila=10, hoja=2, nombre_completo="Mari Esther Zeniga", sexo="F", cedula="601-050604-1008H", edad=22, carrera="Contabilidad", categoria="Estudiante"),
        RawParticipantRow(numero_fila=11, hoja=2, nombre_completo="Jessenia Gomez Rivera", sexo="F", cedula="601-150205-1009I", edad=21, carrera="Contabilidad", categoria="Estudiante"),
        RawParticipantRow(numero_fila=12, hoja=2, nombre_completo="Esther Gabriela Pedro", sexo="F", cedula="601-111004-1010J", edad=22, carrera="Contabilidad", categoria="Estudiante"),
        RawParticipantRow(numero_fila=13, hoja=2, nombre_completo="Nelba Jhanida Pablo Zamora", sexo="F", cedula="601-040403-1011K", edad=23, carrera="Contabilidad", categoria="Estudiante"),
        RawParticipantRow(numero_fila=14, hoja=2, nombre_completo="Estrella Reymundo Peñer", sexo="F", cedula="601-290704-1012L", edad=22, carrera="Contabilidad", categoria="Estudiante"),
        RawParticipantRow(numero_fila=15, hoja=2, nombre_completo="Marcos Michell Hernandez", sexo="M", cedula="601-170802-1013M", edad=24, carrera="Ing. sistc.", categoria="Estudiante"),
        RawParticipantRow(numero_fila=16, hoja=2, nombre_completo="Jhondy Jimmy Arragón", sexo="M", cedula="601-020903-1014N", edad=23, carrera="Ing. sistc.", categoria="Estudiante"),
        RawParticipantRow(numero_fila=17, hoja=2, nombre_completo="Dadam Ossiel Mendoza", sexo="M", cedula=None, edad=24, carrera="Ing. sistc.", categoria="Estudiante"),
        RawParticipantRow(numero_fila=18, hoja=2, nombre_completo="Joyner Henr Ruiz Watson", sexo="M", cedula="601-070104-1016P", edad=22, carrera="Ing. sistc.", categoria="Estudiante"),
    ]
    sheet_data = RawAttendanceSheetData(
        nombre_actividad="Fortalecimiento estudiantil en diseño de logotipos e inteligencia artificial",
        fecha="08/09/2026",
        sede="Bilwi",
        departamento="RACCN",
        municipio="Puerto Cabezas",
        horario="10:00 AM – 12:00 PM",
        filas=filas,
        fuente_origen="asistencia_fisica"
    )

    act_word = WordReportParser.parse_from_dict(datos_word, fuente="informe_word")
    match_res = ActivityMatcher.evaluar_coincidencia(
        actividad_word=act_word,
        datos_asistencia={
            "nombre_actividad_asistencia": sheet_data.nombre_actividad,
            "fecha_evento": sheet_data.fecha,
            "sede": sheet_data.sede,
            "departamento": sheet_data.departamento,
            "municipio_evento": sheet_data.municipio,
            "horario": sheet_data.horario,
            "fuente_origen": sheet_data.fuente_origen
        }
    )
    act_unificada = match_res.actividad_unificada or act_word

    parsed_asist = AttendanceParser.parse_sheet_data(sheet_data, id_actividad=act_unificada.id_actividad)
    res_dedup = IdentityResolver.resolver_identidades(parsed_asist.personas, parsed_asist.participaciones)

    resultado_routing = ParticipantRouter.enrutar(
        actividades=[act_unificada],
        personas=res_dedup.personas,
        participaciones=res_dedup.participaciones,
        validaciones=[]
    )

    estadisticas = StatisticsEngine.calcular(
        resultado_routing=resultado_routing,
        actividades=[act_unificada],
        personas=res_dedup.personas,
        participaciones=res_dedup.participaciones,
    )

    return resultado_routing, estadisticas, act_unificada


@pytest.fixture(scope="module")
def fixtures_plantillas_test_only(tmp_path_factory):
    """Genera los 5 Golden Test Fixtures en una carpeta temporal identificada como TEST-ONLY."""
    carpeta_fixtures = tmp_path_factory.mktemp("test_only_templates")
    TemplateFixtureFactory.crear_fixtures_de_prueba(carpeta_fixtures)

    esquemas = {}
    manifestos = {}
    config_matrices = {
        "matriz_1": ("Matriz 1: Consolidado de Actividades", "fixture_test_matriz_1.xlsx"),
        "matriz_2": ("Matriz 2: Estudiantes", "fixture_test_matriz_2.xlsx"),
        "matriz_3": ("Matriz 3: Académicos y Administrativos", "fixture_test_matriz_3.xlsx"),
        "matriz_4": ("Matriz 4: Colaboradores", "fixture_test_matriz_4.xlsx"),
        "matriz_5": ("Matriz 5: Protagonistas Beneficiados", "fixture_test_matriz_5.xlsx"),
    }

    for id_m, (nombre, archivo) in config_matrices.items():
        ruta = carpeta_fixtures / archivo
        esq = TemplateInspector.inspeccionar_archivo(
            ruta_archivo=ruta,
            id_matriz=id_m,
            nombre_matriz=nombre,
            es_oficial=False,  # Marcar explícitamente como test fixture
        )
        man = MappingManifestoEngine.generar_manifesto(esq)
        esquemas[id_m] = esq
        manifestos[id_m] = man

    return carpeta_fixtures, esquemas, manifestos


class TestMotorExportacionFase9:
    """Suite integral de pruebas unitarias y de integración para Fase 9."""

    def test_01_bloqueo_modo_export_sin_plantillas_oficiales(self, datos_caso_real, fixtures_plantillas_test_only, tmp_path):
        """1. --mode export debe bloquearse si cualquiera de las cinco plantillas no es OFICIAL_REAL."""
        routing, estadisticas, _ = datos_caso_real
        _, esquemas, manifestos = fixtures_plantillas_test_only

        # Los esquemas de test fixtures tienen es_oficial = False
        coordinador = ExportCoordinator(
            resultado_routing=routing,
            informe_estadistico=estadisticas,
            esquemas=esquemas,
            manifestos=manifestos,
            carpeta_salida=tmp_path / "salida_no_permitida",
        )

        with pytest.raises(PlantillasOficialesRequeridasError) as exc_info:
            # Intentar exportación oficial estricta sin permitir fixtures
            coordinador.exportar(modo=ModoExportacion.EXPORT, permitir_fixtures_test_only=False)

        assert "MODO EXPORTACIÓN OFICIAL BLOQUEADO" in str(exc_info.value)
        # Constatar que no se creó la carpeta ni se escribieron archivos parciales
        assert not (tmp_path / "salida_no_permitida").exists()

    def test_02_modo_preview_permite_fixtures_sin_tocar_disco(self, datos_caso_real, fixtures_plantillas_test_only, tmp_path):
        """2. --mode preview puede trabajar con fixtures y jamás deja archivos en disco."""
        routing, estadisticas, _ = datos_caso_real
        _, esquemas, manifestos = fixtures_plantillas_test_only
        carpeta_salida = tmp_path / "salida_preview"

        coordinador = ExportCoordinator(
            resultado_routing=routing,
            informe_estadistico=estadisticas,
            esquemas=esquemas,
            manifestos=manifestos,
            carpeta_salida=carpeta_salida,
        )

        manifiesto = coordinador.exportar(modo=ModoExportacion.PREVIEW, permitir_fixtures_test_only=True)

        assert manifiesto.modo == ModoExportacion.PREVIEW
        assert manifiesto.total_matrices == 5
        assert manifiesto.total_asistencias_enrutadas == 18
        assert manifiesto.total_filas_detalle_escritas == 18
        assert manifiesto.invariante_filas_valida is True
        assert manifiesto.invariante_conservacion_valida is True
        # En preview no debe crearse ningún archivo en carpeta_salida
        assert not carpeta_salida.exists() or len(list(carpeta_salida.glob("*"))) == 0

    def test_03_inmutabilidad_sha256_plantillas_base(self, datos_caso_real, fixtures_plantillas_test_only, tmp_path):
        """3. Verificar SHA-256 de las cinco plantillas originales antes y después demostrando 0 modificaciones."""
        routing, estadisticas, _ = datos_caso_real
        carpeta_fixtures, esquemas, manifestos = fixtures_plantillas_test_only

        # Tomar hashes antes
        hashes_antes = {
            id_m: calcular_sha256(Path(esq.ruta_archivo))
            for id_m, esq in esquemas.items()
        }

        coordinador = ExportCoordinator(
            resultado_routing=routing,
            informe_estadistico=estadisticas,
            esquemas=esquemas,
            manifestos=manifestos,
            carpeta_salida=tmp_path / "salida_hashes",
        )
        manifiesto = coordinador.exportar(modo=ModoExportacion.EXPORT, permitir_fixtures_test_only=True)

        # Tomar hashes después
        hashes_despues = {
            id_m: calcular_sha256(Path(esq.ruta_archivo))
            for id_m, esq in esquemas.items()
        }

        assert hashes_antes == hashes_despues
        assert manifiesto.hashes_plantillas_coinciden is True
        for id_m, res in manifiesto.matrices.items():
            assert res.plantilla_inalterada is True
            assert res.sha256_plantilla_antes == res.sha256_plantilla_despues

    def test_04_preservacion_exacta_formulas_consolidado_matriz_1(self, datos_caso_real, fixtures_plantillas_test_only, tmp_path):
        """
        4. Verificar comprobación exhaustiva de fórmulas en Matriz 1:
        - Coordenada ('I10', 'J10', 'K10', 'L10')
        - Fórmula original vs fórmula posterior
        - Rango referenciado ('I5:I9', etc.)
        - Que la fórmula no haya sido reemplazada por Python
        - Que las filas escritas no invadan el rango de la fórmula
        - Diferenciación clara entre PRESERVACIÓN ESTRUCTURAL y RECALCULACIÓN DEL RESULTADO.
        """
        routing, estadisticas, _ = datos_caso_real
        _, esquemas, manifestos = fixtures_plantillas_test_only
        carpeta_salida = tmp_path / "salida_formulas"

        coordinador = ExportCoordinator(
            resultado_routing=routing,
            informe_estadistico=estadisticas,
            esquemas=esquemas,
            manifestos=manifestos,
            carpeta_salida=carpeta_salida,
        )
        manifiesto = coordinador.exportar(modo=ModoExportacion.EXPORT, permitir_fixtures_test_only=True)

        archivo_m1 = carpeta_salida / "Matriz_1_Consolidado_Actividades.xlsx"
        assert archivo_m1.exists()

        # 1. Comprobación desde el modelo de fórmulas verificadas del manifiesto
        res_m1 = manifiesto.matrices["matriz_1"]
        assert len(res_m1.celdas_con_formula_preservadas) == 4

        formulas_esperadas = {
            "I10": ("=SUM(I5:I9)", "I5:I9"),
            "J10": ("=SUM(J5:J9)", "J5:J9"),
            "K10": ("=SUM(K5:K9)", "K5:K9"),
            "L10": ("=SUM(L5:L9)", "L5:L9"),
        }

        for fv in res_m1.celdas_con_formula_preservadas:
            coord = fv.coordenada
            assert coord in formulas_esperadas
            f_esperada, r_esperado = formulas_esperadas[coord]

            assert fv.formula == f_esperada
            assert fv.formula_original == f_esperada
            assert fv.formula_posterior == f_esperada
            assert fv.rango_referenciado == r_esperado
            assert fv.preservada is True
            assert fv.no_reemplazada is True
            assert fv.filas_no_invaden_rango is True
            assert fv.tipo_preservacion == "PRESERVACION_ESTRUCTURAL"
            assert "openpyxl" in fv.recalculacion_nota

        # 2. Comprobación física con openpyxl (PRESERVACIÓN ESTRUCTURAL)
        wb = openpyxl.load_workbook(str(archivo_m1), data_only=False)
        ws = wb["Consolidado"]

        # Fila 5: Datos de la actividad (escrita dentro de la zona 5..9)
        assert ws["A5"].value == 1
        assert "BICU CUR Bilwi" in str(ws["B5"].value)
        assert ws["I5"].value == 18  # Total Asistencias
        assert ws["J5"].value == 18  # Personas Únicas
        assert ws["K5"].value == 12  # Femenino
        assert ws["L5"].value == 6   # Masculino

        # Fila 10: Fórmulas de Totales al pie preservadas estructuralmente
        assert ws["B10"].value == "TOTALES GENERALES:"
        assert ws["I10"].value == "=SUM(I5:I9)"
        assert ws["J10"].value == "=SUM(J5:J9)"
        assert ws["K10"].value == "=SUM(K5:K9)"
        assert ws["L10"].value == "=SUM(L5:L9)"
        wb.close()

        # 3. Demostración explícita de que openpyxl NO recalcula fórmulas en runtime (RECALCULACIÓN DEL RESULTADO)
        wb_data = openpyxl.load_workbook(str(archivo_m1), data_only=True)
        ws_data = wb_data["Consolidado"]
        # En openpyxl sin abrir en Excel, el valor numérico evaluado de la fórmula no existe (es None o cached)
        # demostrando que la recalculación del resultado es responsabilidad de Microsoft Excel al abrir el archivo.
        assert ws_data["I10"].value is None or not isinstance(ws_data["I10"].value, (int, float))
        wb_data.close()

    def test_05_comportamiento_matrices_con_cero_participantes_n_cero(self, datos_caso_real, fixtures_plantillas_test_only, tmp_path):
        """5. Comportamiento en N = 0 (Matrices 4 y 5): plantilla intacta con 0 filas añadidas."""
        routing, estadisticas, _ = datos_caso_real
        _, esquemas, manifestos = fixtures_plantillas_test_only
        carpeta_salida = tmp_path / "salida_n_cero"

        coordinador = ExportCoordinator(
            resultado_routing=routing,
            informe_estadistico=estadisticas,
            esquemas=esquemas,
            manifestos=manifestos,
            carpeta_salida=carpeta_salida,
        )
        manifiesto = coordinador.exportar(modo=ModoExportacion.EXPORT, permitir_fixtures_test_only=True)

        res_m4 = manifiesto.matrices["matriz_4"]
        res_m5 = manifiesto.matrices["matriz_5"]

        assert res_m4.total_filas_escritas == 0
        assert res_m5.total_filas_escritas == 0

        # Verificar que los archivos existen y conservan encabezados intactos sin filas corruptas
        archivo_m4 = carpeta_salida / "Matriz_4_Colaboradores.xlsx"
        archivo_m5 = carpeta_salida / "Matriz_5_Protagonistas_Beneficiados.xlsx"
        assert archivo_m4.exists()
        assert archivo_m5.exists()

        wb4 = openpyxl.load_workbook(str(archivo_m4), data_only=False)
        ws4 = wb4["Colaboradores"]
        assert ws4.cell(row=4, column=1).value == "No."
        assert ws4.cell(row=5, column=1).value is None  # Fila 5 limpia sin basura
        wb4.close()

    def test_06_invariante_conservacion_participaciones(self, datos_caso_real, fixtures_plantillas_test_only, tmp_path):
        """6. Invariante de conservación de participaciones: 18 = 17 (Est) + 1 (Admin) + 0 (Colab) + 0 (Benef)."""
        routing, estadisticas, _ = datos_caso_real
        _, esquemas, manifestos = fixtures_plantillas_test_only

        coordinador = ExportCoordinator(
            resultado_routing=routing,
            informe_estadistico=estadisticas,
            esquemas=esquemas,
            manifestos=manifestos,
            carpeta_salida=tmp_path / "salida_invariante",
        )
        manifiesto = coordinador.exportar(modo=ModoExportacion.EXPORT, permitir_fixtures_test_only=True)

        assert manifiesto.matrices["matriz_1"].total_filas_escritas == 1
        assert manifiesto.matrices["matriz_2"].total_filas_escritas == 17
        assert manifiesto.matrices["matriz_3"].total_filas_escritas == 1
        assert manifiesto.matrices["matriz_4"].total_filas_escritas == 0
        assert manifiesto.matrices["matriz_5"].total_filas_escritas == 0

        assert manifiesto.total_filas_detalle_escritas == 18
        assert manifiesto.total_asistencias_enrutadas == 18
        assert manifiesto.invariante_filas_valida is True
        assert manifiesto.invariante_conservacion_valida is True

    def test_07_regla_no_invencion_sin_fuente_opcional_vacia(self, datos_caso_real, fixtures_plantillas_test_only, tmp_path):
        """7. Regla de no invención: columnas SIN_FUENTE_OPCIONAL permanecen estrictamente vacías (None)."""
        routing, estadisticas, _ = datos_caso_real
        _, esquemas, manifestos = fixtures_plantillas_test_only
        carpeta_salida = tmp_path / "salida_sin_fuente"

        coordinador = ExportCoordinator(
            resultado_routing=routing,
            informe_estadistico=estadisticas,
            esquemas=esquemas,
            manifestos=manifestos,
            carpeta_salida=carpeta_salida,
        )
        coordinador.exportar(modo=ModoExportacion.EXPORT, permitir_fixtures_test_only=True)

        archivo_m2 = carpeta_salida / "Matriz_2_Estudiantes.xlsx"
        wb = openpyxl.load_workbook(str(archivo_m2), data_only=False)
        ws = wb["Estudiantes"]

        # En Matriz 2, columna K (11) es 'Código Presupuestario Externo' -> SIN_FUENTE
        assert ws.cell(row=4, column=11).value == "Código Presupuestario Externo"
        for r in range(5, 22):
            val = ws.cell(row=r, column=11).value
            assert val is None, f"Fila {r} columna K no está vacía: '{val}'"
        wb.close()

    def test_08_politica_catalogos_declarativa_texto_literal(self, datos_caso_real, fixtures_plantillas_test_only, tmp_path):
        """8. Política declarativa: CATALOGO_PENDIENTE con TEXTO_LITERAL_NO_HOMOLOGADO exporta el texto original."""
        routing, estadisticas, _ = datos_caso_real
        _, esquemas, manifestos = fixtures_plantillas_test_only
        carpeta_salida = tmp_path / "salida_politica_ok"

        politicas = {
            "carrera": PoliticaCatalogo(
                estado="CATALOGO_PENDIENTE",
                requerido=False,
                modo_exportacion=ModoExportacionCatalogo.TEXTO_LITERAL_NO_HOMOLOGADO
            )
        }

        coordinador = ExportCoordinator(
            resultado_routing=routing,
            informe_estadistico=estadisticas,
            esquemas=esquemas,
            manifestos=manifestos,
            politicas_catalogos=politicas,
            carpeta_salida=carpeta_salida,
        )
        manifiesto = coordinador.exportar(modo=ModoExportacion.EXPORT, permitir_fixtures_test_only=True)

        res_m2 = manifiesto.matrices["matriz_2"]
        assert "Carrera" in res_m2.columnas_catalogo_no_homologadas

        # Verificar que el texto original de carrera se escribió
        archivo_m2 = carpeta_salida / "Matriz_2_Estudiantes.xlsx"
        wb = openpyxl.load_workbook(str(archivo_m2), data_only=False)
        ws = wb["Estudiantes"]
        assert ws["F5"].value in ("Lic. Conta", "Contabilidad")
        wb.close()

    def test_09_politica_catalogos_declarativa_bloqueo(self, datos_caso_real, fixtures_plantillas_test_only, tmp_path):
        """9. Política declarativa: si una directiva exige homologación requerida (BLOQUEAR), la exportación se detiene."""
        routing, estadisticas, _ = datos_caso_real
        _, esquemas, manifestos = fixtures_plantillas_test_only
        carpeta_salida = tmp_path / "salida_bloqueo"

        politicas_bloqueo = {
            "carrera": PoliticaCatalogo(
                estado="CATALOGO_PENDIENTE",
                requerido=True,
                modo_exportacion=ModoExportacionCatalogo.BLOQUEAR
            )
        }

        coordinador = ExportCoordinator(
            resultado_routing=routing,
            informe_estadistico=estadisticas,
            esquemas=esquemas,
            manifestos=manifestos,
            politicas_catalogos=politicas_bloqueo,
            carpeta_salida=carpeta_salida,
        )

        with pytest.raises(ExportacionBloqueadaError) as exc_info:
            coordinador.exportar(modo=ModoExportacion.EXPORT, permitir_fixtures_test_only=True)

        assert "catálogo de carreras requerido y no homologado" in str(exc_info.value)
        # Atomicidad: nada debe haber quedado en carpeta_salida
        assert not carpeta_salida.exists() or len(list(carpeta_salida.glob("*.xlsx"))) == 0

    def test_10_atomicidad_all_or_nothing_rollback_ante_fallo(self, datos_caso_real, fixtures_plantillas_test_only, tmp_path):
        """10. Atomicidad All-or-Nothing: si una matriz falla, no se produce ningún archivo parcial en output."""
        routing, estadisticas, _ = datos_caso_real
        _, esquemas, manifestos = fixtures_plantillas_test_only
        carpeta_salida = tmp_path / "salida_atomicidad"

        # Simular esquema corrupto en matriz 5 para provocar fallo en la última matriz
        esquemas_corruptos = dict(esquemas)
        esquema_m5_roto = esquemas["matriz_5"].model_copy(update={"ruta_archivo": "archivo_inexistente_999.xlsx"})
        esquemas_corruptos["matriz_5"] = esquema_m5_roto

        coordinador = ExportCoordinator(
            resultado_routing=routing,
            informe_estadistico=estadisticas,
            esquemas=esquemas_corruptos,
            manifestos=manifestos,
            carpeta_salida=carpeta_salida,
        )

        with pytest.raises(FileNotFoundError):
            coordinador.exportar(modo=ModoExportacion.EXPORT, permitir_fixtures_test_only=True)

        # Verificar que NO se crearon las matrices 1, 2, 3 ni 4 en la carpeta final (rollback total)
        archivos_creados = list(carpeta_salida.glob("*.xlsx")) if carpeta_salida.exists() else []
        assert len(archivos_creados) == 0, f"Se encontraron archivos parciales huérfanos: {archivos_creados}"

    def test_11_generacion_y_serializacion_manifiesto_exportacion_json(self, datos_caso_real, fixtures_plantillas_test_only, tmp_path):
        """11. Generación completa del manifiesto_exportacion.json en entorno de prueba."""
        routing, estadisticas, _ = datos_caso_real
        _, esquemas, manifestos = fixtures_plantillas_test_only
        carpeta_salida = tmp_path / "salida_manifiesto"

        coordinador = ExportCoordinator(
            resultado_routing=routing,
            informe_estadistico=estadisticas,
            esquemas=esquemas,
            manifestos=manifestos,
            carpeta_salida=carpeta_salida,
        )
        manifiesto = coordinador.exportar(modo=ModoExportacion.EXPORT, permitir_fixtures_test_only=True)

        ruta_manifiesto = carpeta_salida / "manifiesto_exportacion.json"
        assert ruta_manifiesto.exists()

        with open(ruta_manifiesto, "r", encoding="utf-8") as f:
            data = json.load(f)

        assert data["total_matrices"] == 5
        assert data["total_asistencias_enrutadas"] == 18
        assert data["total_filas_detalle_escritas"] == 18
        assert data["invariante_filas_valida"] is True
        assert data["invariante_conservacion_valida"] is True
        assert data["hashes_plantillas_coinciden"] is True
        assert "matriz_1" in data["matrices"]
        assert "matriz_2" in data["matrices"]
        assert "matriz_3" in data["matrices"]
        assert "matriz_4" in data["matrices"]
        assert "matriz_5" in data["matrices"]
        # Validar que el archivo JSON en disco contiene su propia ruta y no null
        assert data["ruta_manifiesto"] == str(ruta_manifiesto.resolve())
        assert manifiesto.ruta_manifiesto == str(ruta_manifiesto.resolve())

    def test_12_secuencialidad_correlativa_en_detalle(self, datos_caso_real, fixtures_plantillas_test_only, tmp_path):
        """12. Comprobar que la columna 'No.' contiene una secuencia entera estricta 1..N sin saltos."""
        routing, estadisticas, _ = datos_caso_real
        _, esquemas, manifestos = fixtures_plantillas_test_only
        carpeta_salida = tmp_path / "salida_secuencial"

        coordinador = ExportCoordinator(
            resultado_routing=routing,
            informe_estadistico=estadisticas,
            esquemas=esquemas,
            manifestos=manifestos,
            carpeta_salida=carpeta_salida,
        )
        coordinador.exportar(modo=ModoExportacion.EXPORT, permitir_fixtures_test_only=True)

        archivo_m2 = carpeta_salida / "Matriz_2_Estudiantes.xlsx"
        wb = openpyxl.load_workbook(str(archivo_m2), data_only=False)
        ws = wb["Estudiantes"]

        # Filas 5 a 21 corresponden a los 17 estudiantes
        for idx in range(1, 18):
            r = 4 + idx
            val_sec = ws.cell(row=r, column=1).value
            assert val_sec == idx, f"Fila {r} esperaba No.={idx}, obtuvo {val_sec}"
        wb.close()

    def test_13_manifiesto_exportacion_ruta_autoconsistente(self, datos_caso_real, fixtures_plantillas_test_only, tmp_path):
        """
        13. Comprobar que el manifiesto es 100% autoconsistente:
        - En modo EXPORT: el JSON en disco contiene 'ruta_manifiesto' no nulo y coincide con la ruta real del archivo.
        - En modo PREVIEW: 'ruta_manifiesto' es None y no se crea ningún archivo en disco.
        """
        routing, estadisticas, _ = datos_caso_real
        _, esquemas, manifestos = fixtures_plantillas_test_only

        # A. En modo EXPORT
        carpeta_export = tmp_path / "salida_autoconsistente"
        coord_export = ExportCoordinator(
            resultado_routing=routing,
            informe_estadistico=estadisticas,
            esquemas=esquemas,
            manifestos=manifestos,
            carpeta_salida=carpeta_export,
        )
        manifiesto_export = coord_export.exportar(modo=ModoExportacion.EXPORT, permitir_fixtures_test_only=True)

        ruta_esperada = carpeta_export / "manifiesto_exportacion.json"
        assert ruta_esperada.exists()
        assert manifiesto_export.ruta_manifiesto == str(ruta_esperada.resolve())

        # Inspección directa del JSON en disco: NUNCA debe contener "ruta_manifiesto": null
        with open(ruta_esperada, "r", encoding="utf-8") as f:
            contenido_disco = json.load(f)

        assert contenido_disco["ruta_manifiesto"] is not None
        assert contenido_disco["ruta_manifiesto"] == str(ruta_esperada.resolve())

        # B. En modo PREVIEW
        carpeta_preview = tmp_path / "salida_preview_autoconsistente"
        coord_preview = ExportCoordinator(
            resultado_routing=routing,
            informe_estadistico=estadisticas,
            esquemas=esquemas,
            manifestos=manifestos,
            carpeta_salida=carpeta_preview,
        )
        manifiesto_preview = coord_preview.exportar(modo=ModoExportacion.PREVIEW, permitir_fixtures_test_only=True)

        assert manifiesto_preview.ruta_manifiesto is None
        assert not (carpeta_preview / "manifiesto_exportacion.json").exists()

    def test_14_preservacion_formulas_matriz_1_n_cero(self, fixtures_plantillas_test_only, tmp_path):
        """
        14. Comprobar fórmulas en Matriz 1 con N = 0 actividades:
        - 0 filas escritas
        - Fórmulas de sumatoria al pie preservadas estructuralmente
        - No hay reemplazo ni invasión de celdas
        """
        _, esquemas, manifestos = fixtures_plantillas_test_only
        carpeta_salida = tmp_path / "salida_m1_n_cero"

        exporter = ConsolidadoExporter(
            id_matriz="matriz_1",
            nombre_matriz="Matriz 1: Consolidado de Actividades",
            nombre_archivo_salida="Matriz_1_Consolidado_Actividades.xlsx",
            esquema=esquemas["matriz_1"],
            manifesto=manifestos["matriz_1"],
        )

        res = exporter.exportar(actividades=[], carpeta_destino=carpeta_salida)

        assert res.total_filas_escritas == 0
        assert res.fila_inicio is None
        assert res.fila_fin is None
        assert len(res.celdas_con_formula_preservadas) == 4

        for fv in res.celdas_con_formula_preservadas:
            assert fv.preservada is True
            assert fv.no_reemplazada is True
            assert fv.filas_no_invaden_rango is True
            assert fv.tipo_preservacion == "PRESERVACION_ESTRUCTURAL"

        # Comprobación en el archivo Excel generado
        archivo = carpeta_salida / "Matriz_1_Consolidado_Actividades.xlsx"
        wb = openpyxl.load_workbook(str(archivo), data_only=False)
        ws = wb["Consolidado"]
        assert ws["I10"].value == "=SUM(I5:I9)"
        assert ws["J10"].value == "=SUM(J5:J9)"
        assert ws["K10"].value == "=SUM(K5:K9)"
        assert ws["L10"].value == "=SUM(L5:L9)"
        # La fila 5 no tiene datos
        assert ws["A5"].value is None
        wb.close()

    def test_15_preservacion_formulas_matriz_1_n_mayor_a_uno(self, datos_caso_real, fixtures_plantillas_test_only, tmp_path):
        """
        15. Comprobar fórmulas en Matriz 1 con N > 1 actividades (N = 3):
        - 3 filas escritas dentro de la zona escribible (filas 5, 6, 7)
        - Fórmulas al pie (fila 10) preservadas estructuralmente
        - No hay invasión de rango (7 < 10)
        """
        _, estadisticas, _ = datos_caso_real
        _, esquemas, manifestos = fixtures_plantillas_test_only
        carpeta_salida = tmp_path / "salida_m1_n_tres"

        # Simular 3 actividades con métricas
        act_base = estadisticas.actividades[0]
        act_1 = act_base.model_copy(update={"id_actividad": "act_001", "nombre_actividad": "Actividad Alpha"})
        act_2 = act_base.model_copy(update={"id_actividad": "act_002", "nombre_actividad": "Actividad Beta"})
        act_3 = act_base.model_copy(update={"id_actividad": "act_003", "nombre_actividad": "Actividad Gamma"})

        exporter = ConsolidadoExporter(
            id_matriz="matriz_1",
            nombre_matriz="Matriz 1: Consolidado de Actividades",
            nombre_archivo_salida="Matriz_1_Consolidado_Actividades.xlsx",
            esquema=esquemas["matriz_1"],
            manifesto=manifestos["matriz_1"],
        )

        res = exporter.exportar(actividades=[act_1, act_2, act_3], carpeta_destino=carpeta_salida)

        assert res.total_filas_escritas == 3
        assert res.fila_inicio == 5
        assert res.fila_fin == 7
        assert len(res.celdas_con_formula_preservadas) == 4

        for fv in res.celdas_con_formula_preservadas:
            assert fv.preservada is True
            assert fv.no_reemplazada is True
            assert fv.filas_no_invaden_rango is True
            assert fv.tipo_preservacion == "PRESERVACION_ESTRUCTURAL"

        archivo = carpeta_salida / "Matriz_1_Consolidado_Actividades.xlsx"
        wb = openpyxl.load_workbook(str(archivo), data_only=False)
        ws = wb["Consolidado"]

        # Filas 5, 6, 7 tienen las 3 actividades
        assert ws["A5"].value == 1
        assert ws["B5"].value == "Actividad Alpha"
        assert ws["A6"].value == 2
        assert ws["B6"].value == "Actividad Beta"
        assert ws["A7"].value == 3
        assert ws["B7"].value == "Actividad Gamma"

        # Fórmulas en fila 10 intactas
        assert ws["I10"].value == "=SUM(I5:I9)"
        assert ws["J10"].value == "=SUM(J5:J9)"
        assert ws["K10"].value == "=SUM(K5:K9)"
        assert ws["L10"].value == "=SUM(L5:L9)"
        wb.close()

    def test_16_matriz_estudiantes_n_cero_sin_filas_ficticias(self, fixtures_plantillas_test_only, tmp_path):
        """
        16. Comprobar Matriz 2 (Estudiantes) con N = 0:
        - 0 filas escritas
        - Plantilla limpia sin filas ficticias ni caracteres vacíos
        """
        _, esquemas, manifestos = fixtures_plantillas_test_only
        carpeta_salida = tmp_path / "salida_m2_n_cero"

        exporter = ParticipantesExporter(
            id_matriz="matriz_2",
            nombre_matriz="Matriz 2: Estudiantes",
            nombre_archivo_salida="Matriz_2_Estudiantes.xlsx",
            esquema=esquemas["matriz_2"],
            manifesto=manifestos["matriz_2"],
        )

        res = exporter.exportar(registros=[], carpeta_destino=carpeta_salida)

        assert res.total_filas_escritas == 0
        assert res.fila_inicio is None
        assert res.fila_fin is None

        archivo = carpeta_salida / "Matriz_2_Estudiantes.xlsx"
        wb = openpyxl.load_workbook(str(archivo), data_only=False)
        ws = wb["Estudiantes"]
        # Encabezados preservados en fila 4
        assert ws["A4"].value == "No."
        # Fila 5 debe estar vacía (sin participantes ficticios)
        assert ws["A5"].value is None
        wb.close()


class TestExportadorProduccionRealFase93:
    """
    Suite de pruebas oficiales para FASE 9.3: Corrección del Exportador para Producción Real.
    Verifica compatibilidad estricta con las 5 plantillas oficiales reales de BICU:
    - M1: 42 cols, header 1, data 2, Tabla1 A1:AP...
    - M2: 57 cols, header 1, data 2, Tabla1 A1:BE...
    - M3: 58 cols, header 1, data 2, Tabla1 A1:BF...
    - M4: 57 cols, header 1, data 2, Tabla1 A1:BE...
    - M5: 53 cols, header 1, data 2, Tabla1 A1:BA...
    - Cero destrucción y conservación histórica en M5 (Opción A: 32 registros persistentes en filas 2..33).
    - Preservación estricta de fórmulas nativas (=DATEDIF..., =+A{r-1}+1) sin convertir a enteros estáticos.
    - Sincronización segura de Tabla1.ref y autoFilter.ref para N=0, N=1, N=18, N=100.
    - Exportación atómica del Caso Real Septiembre 2026 (18 participaciones: 17 M2 + 1 M3 + 0 M4 + 0 M5).
    - Inmutabilidad SHA-256 de las plantillas base oficiales en disco.
    """

    @pytest.fixture(scope="class")
    def esquemas_y_manifestos_oficiales(self):
        """Carga esquemas y manifestos de las 5 plantillas oficiales reales en templates/."""
        from app.templates_analysis.reporter import TemplateReporter
        carpeta_templates = Path("templates")
        rep = TemplateReporter.ejecutar_inspeccion_global(carpeta_templates=carpeta_templates)
        return rep.esquemas, rep.manifestos

    def test_01_estructura_oficial_cinco_matrices(self, esquemas_y_manifestos_oficiales):
        """
        Verifica dimensiones y estructura auditada de las 5 plantillas oficiales reales:
        M1=42, M2=57, M3=58, M4=57, M5=53 columnas.
        Encabezados en fila 1, inicio de datos en fila 2.
        204 fórmulas nativas preexistentes en total.
        """
        esquemas, _ = esquemas_y_manifestos_oficiales

        assert esquemas["matriz_1"].max_columnas == 42
        assert esquemas["matriz_1"].fila_encabezados == 1
        assert esquemas["matriz_1"].fila_inicio_datos == 2

        assert esquemas["matriz_2"].max_columnas == 57
        assert esquemas["matriz_2"].fila_encabezados == 1
        assert esquemas["matriz_2"].fila_inicio_datos == 2

        assert esquemas["matriz_3"].max_columnas == 58
        assert esquemas["matriz_3"].fila_encabezados == 1
        assert esquemas["matriz_3"].fila_inicio_datos == 2

        assert esquemas["matriz_4"].max_columnas == 57
        assert esquemas["matriz_4"].fila_encabezados == 1
        assert esquemas["matriz_4"].fila_inicio_datos == 2

        assert esquemas["matriz_5"].max_columnas == 53
        assert esquemas["matriz_5"].fila_encabezados == 1
        assert esquemas["matriz_5"].fila_inicio_datos == 2

        # Conteo exacto de fórmulas institucionales (20 + 41 + 41 + 41 + 61 = 204)
        total_formulas = sum(
            len([cp for cp in esq.celdas_protegidas if cp.tipo_proteccion == "FORMULA"])
            for esq in esquemas.values()
        )
        assert total_formulas == 204

    def test_02_preservacion_formula_edad_datedif_no_estatico(self, esquemas_y_manifestos_oficiales, tmp_path):
        """
        Verifica que la columna Edad conserva estrictamente la fórmula =DATEDIF(...)
        y NUNCA escribe person.edad como un entero estático.
        """
        from app.core.models.person import Person
        from app.core.models.activity import Activity
        from app.core.models.participation import Participation
        from app.routing.models import EnrutamientoRegistro
        from app.routing.enums import MatrizDestino, SubtipoInstitucional

        esquemas, manifestos = esquemas_y_manifestos_oficiales
        carpeta_salida = tmp_path / "salida_test_datedif"

        p = Person(
            id_persona="p_test",
            nombre_completo="Ana Maria Gomez",
            cedula="601-150500-1000A",
            nombres="Ana Maria",
            apellidos="Gomez",
            fecha_nacimiento="2000-05-15",
            edad=26,  # Entero estático que NO debe ser escrito en la celda de Edad
        )
        act = Activity(
            id_actividad="act_1",
            nombre_actividad_original="Capacitacion IA",
            sede="Bilwi",
            nombre_actividad_oficial="Capacitacion Oficial IA",
        )
        reg = EnrutamientoRegistro(
            persona=p,
            actividad=act,
            participacion=Participation(
                id_participacion="part_1",
                id_persona="p_test",
                id_actividad="act_1",
                rol="Estudiante",
                categoria_participacion="Estudiante",
            ),
            matriz_destino=MatrizDestino.ESTUDIANTES,
            subtipo_institucional=SubtipoInstitucional.NO_APLICA,
            estado_operativo="APTO",
            motivo_enrutamiento="test_datedif",
        )

        exp_m2 = ParticipantesExporter(
            id_matriz="matriz_2",
            nombre_matriz="Matriz 2: Estudiantes",
            nombre_archivo_salida="Matriz_2_Estudiantes.xlsx",
            esquema=esquemas["matriz_2"],
            manifesto=manifestos["matriz_2"],
        )
        res = exp_m2.exportar([reg], carpeta_destino=carpeta_salida)

        archivo_m2 = carpeta_salida / "Matriz_2_Estudiantes.xlsx"
        wb = openpyxl.load_workbook(str(archivo_m2), data_only=False)
        ws = wb.active

        # Columna 40 es 'Edad' en M2
        val_edad = ws.cell(row=2, column=40).value
        assert val_edad == '=DATEDIF(Tabla1[[#This Row],[Fecha_nacimiento]], TODAY(), "Y")'
        assert val_edad != 26  # NO es entero estático

        # Columna 39 es 'Fecha_nacimiento'
        val_fn = ws.cell(row=2, column=39).value
        assert str(val_fn).startswith("2000-05-15")

        wb.close()

    def test_03_m5_conservacion_historica_cero_destruccion_n_cero(self, esquemas_y_manifestos_oficiales, tmp_path):
        """
        M5 con N=0 nuevos registros:
        - 32 registros históricos preservados intactos en filas 2..33.
        - 0 filas nuevas escritas.
        - Tabla1.ref permanece en A1:BA33.
        - Verificación campo por campo: 100% de las 1696 celdas históricas idénticas al original.
        - Valores manuales A15=14 y A33=32 inalterados.
        """
        esquemas, manifestos = esquemas_y_manifestos_oficiales
        carpeta_salida = tmp_path / "salida_m5_n0"

        exp_m5 = ParticipantesExporter(
            id_matriz="matriz_5",
            nombre_matriz="Matriz 5: Protagonistas Beneficiados",
            nombre_archivo_salida="Matriz_5_Protagonistas_Beneficiados.xlsx",
            esquema=esquemas["matriz_5"],
            manifesto=manifestos["matriz_5"],
        )
        res = exp_m5.exportar(registros=[], carpeta_destino=carpeta_salida)

        assert res.total_registros_historicos_preservados == 32
        assert res.total_filas_escritas == 0

        # Cargar libro original y generado
        wb_orig = openpyxl.load_workbook("templates/Matriz_5_Protagonistas_Beneficiados.xlsx", data_only=False)
        ws_orig = wb_orig.active

        archivo_gen = carpeta_salida / "Matriz_5_Protagonistas_Beneficiados.xlsx"
        wb_gen = openpyxl.load_workbook(str(archivo_gen), data_only=False)
        ws_gen = wb_gen.active

        assert ws_gen.tables["Tabla1"].ref == "A1:BA33"

        # Comparación celda por celda de los 32 registros históricos (filas 2 a 33, cols 1 a 53)
        celdas_comparadas = 0
        diferencias = []
        for r in range(2, 34):
            for c in range(1, 54):
                celdas_comparadas += 1
                vo = ws_orig.cell(row=r, column=c).value
                vg = ws_gen.cell(row=r, column=c).value
                if vo != vg:
                    diferencias.append((r, c, vo, vg))

        assert celdas_comparadas == 1696
        assert len(diferencias) == 0, f"Diferencias en históricos M5: {diferencias[:5]}"

        # Particularidad institucional de M5: A15 y A33 son enteros manuales
        assert ws_gen["A15"].value == 14
        assert ws_gen["A33"].value == 32

        wb_orig.close()
        wb_gen.close()

    def test_04_m5_append_n_uno(self, esquemas_y_manifestos_oficiales, tmp_path):
        """
        M5 con N=1 nuevo registro:
        - Ocupa fila 34 (después de los 32 históricos).
        - Correlativo en fila 34 es '=+A33+1'.
        - Edad en fila 34 es '=DATEDIF(...)'.
        - Tabla1.ref se expande a A1:BA34.
        - Filas 2..33 permanecen 100% idénticas.
        """
        from app.core.models.person import Person
        from app.core.models.activity import Activity
        from app.core.models.participation import Participation
        from app.routing.models import EnrutamientoRegistro
        from app.routing.enums import MatrizDestino, SubtipoInstitucional

        esquemas, manifestos = esquemas_y_manifestos_oficiales
        carpeta_salida = tmp_path / "salida_m5_n1"

        p = Person(
            id_persona="p_ben_1",
            nombre_completo="Carlos Benitez",
            cedula="601-201090-1001K",
            nombres="Carlos",
            apellidos="Benitez",
            fecha_nacimiento="1990-10-20",
        )
        act = Activity(
            id_actividad="act_1",
            nombre_actividad_original="Taller Productivo",
            sede="Bluefields",
            nombre_actividad_oficial="Taller Oficial",
        )
        reg = EnrutamientoRegistro(
            persona=p,
            actividad=act,
            participacion=Participation(
                id_participacion="part_b1",
                id_persona="p_ben_1",
                id_actividad="act_1",
                rol="Protagonista",
                categoria_participacion="Protagonista",
            ),
            matriz_destino=MatrizDestino.BENEFICIADOS,
            subtipo_institucional=SubtipoInstitucional.NO_APLICA,
            estado_operativo="APTO",
            motivo_enrutamiento="test_m5_n1",
        )

        exp_m5 = ParticipantesExporter(
            id_matriz="matriz_5",
            nombre_matriz="Matriz 5: Protagonistas Beneficiados",
            nombre_archivo_salida="Matriz_5_Protagonistas_Beneficiados.xlsx",
            esquema=esquemas["matriz_5"],
            manifesto=manifestos["matriz_5"],
        )
        res = exp_m5.exportar([reg], carpeta_destino=carpeta_salida)

        assert res.total_registros_historicos_preservados == 32
        assert res.total_filas_escritas == 1
        assert res.fila_inicio == 34
        assert res.fila_fin == 34

        archivo_gen = carpeta_salida / "Matriz_5_Protagonistas_Beneficiados.xlsx"
        wb_gen = openpyxl.load_workbook(str(archivo_gen), data_only=False)
        ws_gen = wb_gen.active

        assert ws_gen.tables["Tabla1"].ref == "A1:BA34"
        assert ws_gen.cell(34, 1).value == "=+A33+1"
        assert ws_gen.cell(34, 41).value == '=DATEDIF(Tabla1[[#This Row],[Fecha_nacimiento]], TODAY(), "Y")'
        assert ws_gen.cell(34, 34).value == "Carlos Benitez"
        assert ws_gen.cell(34, 36).value == "601-201090-1001K"
        assert ws_gen.cell(34, 2).value == "Bluefields"

        # Históricos en fila 33 inalterados
        assert ws_gen.cell(33, 1).value == 32

        wb_gen.close()

    def test_05_m5_append_n_dieciocho(self, esquemas_y_manifestos_oficiales, tmp_path):
        """
        M5 con N=18 nuevos registros:
        - 32 históricos (filas 2..33) + 18 nuevos = 50 registros totales.
        - Filas nuevas ocupan 34..51.
        - Tabla1.ref se expande a A1:BA51.
        - Correlativos: fila 34 es '=+A33+1', fila 51 es '=+A50+1'.
        """
        from app.core.models.person import Person
        from app.core.models.activity import Activity
        from app.core.models.participation import Participation
        from app.routing.models import EnrutamientoRegistro
        from app.routing.enums import MatrizDestino, SubtipoInstitucional

        esquemas, manifestos = esquemas_y_manifestos_oficiales
        carpeta_salida = tmp_path / "salida_m5_n18"

        registros = []
        act = Activity(
            id_actividad="act_18",
            nombre_actividad_original="Capacitacion Comunitaria",
            sede="Rama",
            nombre_actividad_oficial="Capacitacion Oficial",
        )
        for i in range(1, 19):
            p = Person(
                id_persona=f"p_ben_{i}",
                nombre_completo=f"Beneficiario {i}",
                cedula=f"601-010195-{i:04d}Z",
                nombres=f"Beneficiario",
                apellidos=f"{i}",
                fecha_nacimiento="1995-01-01",
            )
            reg = EnrutamientoRegistro(
                persona=p,
                actividad=act,
                participacion=Participation(
                    id_participacion=f"part_b_{i}",
                    id_persona=f"p_ben_{i}",
                    id_actividad="act_18",
                    rol="Protagonista",
                    categoria_participacion="Protagonista",
                ),
                matriz_destino=MatrizDestino.BENEFICIADOS,
                subtipo_institucional=SubtipoInstitucional.NO_APLICA,
                estado_operativo="APTO",
                motivo_enrutamiento="test_m5_n18",
            )
            registros.append(reg)

        exp_m5 = ParticipantesExporter(
            id_matriz="matriz_5",
            nombre_matriz="Matriz 5: Protagonistas Beneficiados",
            nombre_archivo_salida="Matriz_5_Protagonistas_Beneficiados.xlsx",
            esquema=esquemas["matriz_5"],
            manifesto=manifestos["matriz_5"],
        )
        res = exp_m5.exportar(registros, carpeta_destino=carpeta_salida)

        assert res.total_registros_historicos_preservados == 32
        assert res.total_filas_escritas == 18
        assert res.fila_inicio == 34
        assert res.fila_fin == 51

        archivo_gen = carpeta_salida / "Matriz_5_Protagonistas_Beneficiados.xlsx"
        wb_gen = openpyxl.load_workbook(str(archivo_gen), data_only=False)
        ws_gen = wb_gen.active

        assert ws_gen.tables["Tabla1"].ref == "A1:BA51"
        assert ws_gen.cell(34, 1).value == "=+A33+1"
        assert ws_gen.cell(51, 1).value == "=+A50+1"
        assert ws_gen.cell(34, 41).value == '=DATEDIF(Tabla1[[#This Row],[Fecha_nacimiento]], TODAY(), "Y")'
        assert ws_gen.cell(51, 41).value == '=DATEDIF(Tabla1[[#This Row],[Fecha_nacimiento]], TODAY(), "Y")'

        wb_gen.close()

    def test_06_m5_append_n_cien(self, esquemas_y_manifestos_oficiales, tmp_path):
        """
        M5 con N=100 nuevos registros:
        - Filas nuevas ocupan 34..133.
        - Tabla1.ref se expande a A1:BA133.
        - 32 históricos preservados sin corrupción de tabla.
        """
        from app.core.models.person import Person
        from app.core.models.activity import Activity
        from app.core.models.participation import Participation
        from app.routing.models import EnrutamientoRegistro
        from app.routing.enums import MatrizDestino, SubtipoInstitucional

        esquemas, manifestos = esquemas_y_manifestos_oficiales
        carpeta_salida = tmp_path / "salida_m5_n100"

        act = Activity(
            id_actividad="act_100",
            nombre_actividad_original="Feria Tecnologica",
            sede="Bilwi",
            nombre_actividad_oficial="Feria Oficial",
        )
        registros = []
        for i in range(1, 101):
            p = Person(
                id_persona=f"p_{i}",
                nombre_completo=f"Protagonista {i}",
                cedula=f"601-010190-{i:04d}A",
                nombres=f"Protagonista",
                apellidos=f"{i}",
                fecha_nacimiento="1990-01-01",
            )
            reg = EnrutamientoRegistro(
                persona=p,
                actividad=act,
                participacion=Participation(
                    id_participacion=f"part_{i}",
                    id_persona=f"p_{i}",
                    id_actividad="act_100",
                    rol="Protagonista",
                    categoria_participacion="Protagonista",
                ),
                matriz_destino=MatrizDestino.BENEFICIADOS,
                subtipo_institucional=SubtipoInstitucional.NO_APLICA,
                estado_operativo="APTO",
                motivo_enrutamiento="test_m5_n100",
            )
            registros.append(reg)

        exp_m5 = ParticipantesExporter(
            id_matriz="matriz_5",
            nombre_matriz="Matriz 5: Protagonistas Beneficiados",
            nombre_archivo_salida="Matriz_5_Protagonistas_Beneficiados.xlsx",
            esquema=esquemas["matriz_5"],
            manifesto=manifestos["matriz_5"],
        )
        res = exp_m5.exportar(registros, carpeta_destino=carpeta_salida)

        assert res.total_registros_historicos_preservados == 32
        assert res.total_filas_escritas == 100
        assert res.fila_inicio == 34
        assert res.fila_fin == 133

        archivo_gen = carpeta_salida / "Matriz_5_Protagonistas_Beneficiados.xlsx"
        wb_gen = openpyxl.load_workbook(str(archivo_gen), data_only=False)
        ws_gen = wb_gen.active

        assert ws_gen.tables["Tabla1"].ref == "A1:BA133"
        assert ws_gen.cell(133, 1).value == "=+A132+1"
        assert ws_gen.cell(133, 41).value == '=DATEDIF(Tabla1[[#This Row],[Fecha_nacimiento]], TODAY(), "Y")'

        wb_gen.close()

    def test_07_inmutabilidad_criptografica_sha256_plantillas_base(self, esquemas_y_manifestos_oficiales, tmp_path):
        """
        Verifica que las plantillas base en templates/ NO sufren alteración alguna durante
        la ejecución del exportador (SHA-256 idéntico antes y después de exportar).
        """
        from app.exporters.base_exporter import calcular_sha256

        carpeta_templates = Path("templates")
        plantillas = [
            carpeta_templates / "Matriz_1_Consolidado_Actividades.xlsx",
            carpeta_templates / "Matriz_2_Estudiantes.xlsx",
            carpeta_templates / "Matriz_3_Academicos_Administrativos.xlsx",
            carpeta_templates / "Matriz_4_Colaboradores.xlsx",
            carpeta_templates / "Matriz_5_Protagonistas_Beneficiados.xlsx",
        ]

        hashes_antes = {p.name: calcular_sha256(p) for p in plantillas}

        # Ejecutar exportaciones de prueba en tmp_path
        esquemas, manifestos = esquemas_y_manifestos_oficiales
        for id_m, nom_archivo in [
            ("matriz_1", "Matriz_1_Consolidado_Actividades.xlsx"),
            ("matriz_2", "Matriz_2_Estudiantes.xlsx"),
            ("matriz_3", "Matriz_3_Academicos_Administrativos.xlsx"),
            ("matriz_4", "Matriz_4_Colaboradores.xlsx"),
            ("matriz_5", "Matriz_5_Protagonistas_Beneficiados.xlsx"),
        ]:
            if id_m == "matriz_1":
                exp = ConsolidadoExporter(id_m, id_m, nom_archivo, esquemas[id_m], manifestos[id_m])
                exp.exportar([], tmp_path / "test_sha")
            else:
                exp = ParticipantesExporter(id_m, id_m, nom_archivo, esquemas[id_m], manifestos[id_m])
                exp.exportar([], tmp_path / "test_sha")

        hashes_despues = {p.name: calcular_sha256(p) for p in plantillas}

        for nom, h_ant in hashes_antes.items():
            assert h_ant == hashes_despues[nom], f"Plantilla {nom} fue modificada en disco!"

    def test_08_caso_real_septiembre_2026_coordinador_atomico_oficial(self, tmp_path):
        """
        Verifica la exportación integral del Caso Real Septiembre 2026 usando las 5 plantillas oficiales:
        - 18 participaciones de entrada
        - 18 participaciones enrutadas (17 M2 + 1 M3 + 0 M4 + 0 M5)
        - 1 actividad consolidada en M1
        - 32 registros históricos preservados en M5
        - 204 fórmulas nativas preservadas
        - Invariantes de conservación y filas cumplidas
        - Generación atómica con manifiesto_exportacion.json
        """
        import yaml
        from app.main import _procesar_pipeline_datos
        from app.exporters.coordinator import ExportCoordinator
        from app.exporters.models import ModoExportacion

        with open("config/settings.yaml", "r", encoding="utf-8") as f:
            settings = yaml.safe_load(f)

        routing, stats, esquemas, manifestos, politicas = _procesar_pipeline_datos(settings)
        carpeta_salida = tmp_path / "salida_caso_real_oficial"

        coordinador = ExportCoordinator(
            resultado_routing=routing,
            informe_estadistico=stats,
            esquemas=esquemas,
            manifestos=manifestos,
            politicas_catalogos=politicas,
            carpeta_salida=carpeta_salida,
        )

        manifiesto = coordinador.exportar(modo=ModoExportacion.EXPORT, permitir_fixtures_test_only=False)

        assert manifiesto.es_oficial is True
        assert manifiesto.plantillas_todas_oficiales is True
        assert manifiesto.invariante_conservacion_valida is True
        assert manifiesto.invariante_filas_valida is True
        assert manifiesto.total_asistencias_enrutadas == 18
        assert manifiesto.total_filas_detalle_escritas == 18

        # Verificar cada matriz resultante
        assert manifiesto.matrices["matriz_1"].total_filas_escritas == 1
        assert manifiesto.matrices["matriz_2"].total_filas_escritas == 17
        assert manifiesto.matrices["matriz_3"].total_filas_escritas == 1
        assert manifiesto.matrices["matriz_4"].total_filas_escritas == 0
        assert manifiesto.matrices["matriz_5"].total_filas_escritas == 0
        assert manifiesto.matrices["matriz_5"].total_registros_historicos_preservados == 32

        # Verificar total de fórmulas preservadas (204)
        total_form_pres = sum(
            len(res.celdas_con_formula_preservadas)
            for res in manifiesto.matrices.values()
        )
        assert total_form_pres == 204

        # Verificar manifiesto JSON en disco
        ruta_man_json = carpeta_salida / "manifiesto_exportacion.json"
        assert ruta_man_json.exists()
        with open(ruta_man_json, "r", encoding="utf-8") as f:
            data_man = json.load(f)
        assert data_man["es_oficial"] is True
        assert data_man["total_filas_detalle_escritas"] == 18

