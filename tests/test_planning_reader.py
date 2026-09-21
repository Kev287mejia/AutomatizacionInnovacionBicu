"""Suite de Pruebas Automatizadas — Planning Reader Institucional y Servicio de Ingestión.

Fase 29.7 — Implementación Controlada del Planning Reader Institucional.
Cubre:
  - Pruebas unitarias de ExcelPlanningReader (Formatos A y B, errores estructurales, tipos, vacíos).
  - Pruebas de integración de PlannedActivityIngestionService con PlanningUnitOfWork y SQLite V003.
  - Idempotencia y protección estricta de la regla R-08 sobre diseños en estado APPROVED.
  - Auditoría AST de aislamiento arquitectónico (cero dependencias prohibidas).
  - Prueba de cero contaminación hacia el pipeline de ejecución Word → M1–M5.
"""
from __future__ import annotations

import ast
from datetime import date, datetime
import os
from pathlib import Path
import sqlite3
import tempfile
import uuid

import openpyxl
import pytest

from app.planning.application.services import PlannedActivityIngestionService
from app.planning.domain.dtos import (
    PlanningIngestionReportDTO,
    PlanningSourceActivityDTO,
    PlanningSourceReadResultDTO,
)
from app.planning.domain.entities import MethodologicalDesign, PlannedActivity
from app.planning.domain.ports import PlanningSourceReaderPort, PlanningUnitOfWorkPort
from app.planning.domain.value_objects import DesignStatus, ParticipantGoals
from app.planning.infrastructure.persistence.schema_v003 import V003_SCHEMA_DDL_STATEMENTS
from app.planning.infrastructure.persistence.unit_of_work import PlanningUnitOfWork
from app.planning.infrastructure.readers.excel_planning_reader import ExcelPlanningReader


# ---------------------------------------------------------------------------
# FIXTURES
# ---------------------------------------------------------------------------
@pytest.fixture
def temp_db_path(tmp_path: Path) -> str:
    """Crea una base de datos SQLite temporal con el esquema V003 inicializado."""
    db_file = str(tmp_path / "test_planning_reader.sqlite")
    conn = sqlite3.connect(db_file)
    try:
        conn.execute("PRAGMA foreign_keys = ON;")
        for stmt in V003_SCHEMA_DDL_STATEMENTS:
            conn.execute(stmt)
        conn.commit()
    finally:
        conn.close()
    return db_file


@pytest.fixture
def uow(temp_db_path: str) -> PlanningUnitOfWork:
    """Instancia de PlanningUnitOfWork apuntando a la base de datos temporal."""
    return PlanningUnitOfWork(db_path=temp_db_path)


@pytest.fixture
def reader() -> ExcelPlanningReader:
    """Instancia del lector ExcelPlanningReader."""
    return ExcelPlanningReader()


# ---------------------------------------------------------------------------
# 1. PRUEBAS DEL LECTOR (ExcelPlanningReader)
# ---------------------------------------------------------------------------
class TestExcelPlanningReader:
    """Pruebas unitarias para el lector de matrices institucionales Excel."""

    def test_reader_nonexistent_file(self, reader: ExcelPlanningReader):
        """Verifica que un archivo inexistente retorne error estructural sin lanzar excepciones."""
        result = reader.read_planning_source("archivo_fantasma_inexistente.xlsx")
        assert len(result.structural_errors) > 0
        assert "no existe" in result.structural_errors[0]
        assert len(result.activities) == 0

    def test_reader_invalid_sheet_name(self, reader: ExcelPlanningReader, tmp_path: Path):
        """Verifica que solicitar una hoja inexistente retorne error estructural."""
        fpath = str(tmp_path / "test_hoja_invalida.xlsx")
        wb = openpyxl.Workbook()
        wb.active.title = "HojaValida"
        wb.save(fpath)
        wb.close()

        result = reader.read_planning_source(fpath, sheet_name="HojaInexistente")
        assert len(result.structural_errors) > 0
        assert "no existe" in result.structural_errors[0]
        assert len(result.activities) == 0

    def test_reader_missing_required_headers(self, reader: ExcelPlanningReader, tmp_path: Path):
        """Verifica que un Excel sin columnas 'Actividad' o 'Sede' retorne error estructural."""
        fpath = str(tmp_path / "test_sin_encabezados.xlsx")
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Programas, proyectos y act"
        ws.cell(1, 1, "Columna1")
        ws.cell(1, 2, "Columna2")
        ws.cell(2, 1, "Dato1")
        wb.save(fpath)
        wb.close()

        result = reader.read_planning_source(fpath)
        assert len(result.structural_errors) > 0
        assert "No se encontraron los encabezados institucionales mínimos" in result.structural_errors[0]

    def test_reader_real_format_a_documento_kevds_2(self, reader: ExcelPlanningReader):
        """Lee el archivo institucional real Documento_Kevds_2.xlsx (Formato A: Agregado)."""
        real_file = "input/matrices_reales_uat/Documento_Kevds_2.xlsx"
        if not os.path.isfile(real_file):
            pytest.skip(f"Archivo real '{real_file}' no disponible en este entorno.")

        result = reader.read_planning_source(real_file)
        assert len(result.structural_errors) == 0
        assert result.sheet_name == "Programas, proyectos y act"
        assert len(result.activities) >= 1

        act = result.activities[0]
        assert act.raw_no == "1"
        assert act.sede == "Bilwi"
        assert "logotipos e inteligencia artificial" in (act.activity_name or "")
        assert act.dep_sede == "RACCN"
        assert act.mun_sede == "Puerto Cabezas"
        # Metas cuantitativas agregadas verificadas en archivo real
        assert act.est_grado_m == 6
        assert act.est_grado_f == 11
        assert act.docentes_m == 0
        assert act.docentes_f == 0
        assert act.administrativos_m == 0
        assert act.administrativos_f == 1

    def test_reader_real_format_b_documento_kevds_1(self, reader: ExcelPlanningReader):
        """Lee el archivo institucional real Documento_Kevds_1.xlsx (Formato B: Detallado)."""
        real_file = "input/matrices_reales_uat/Documento_Kevds_1.xlsx"
        if not os.path.isfile(real_file):
            pytest.skip(f"Archivo real '{real_file}' no disponible en este entorno.")

        result = reader.read_planning_source(real_file)
        assert len(result.structural_errors) == 0
        assert result.sheet_name == "Pregrado_Grado en Actividades"
        assert len(result.activities) >= 1

        act = result.activities[0]
        assert act.sede == "Bilwi"
        assert act.programa == "Innovación y Emprendimiento"
        assert "logotipos e inteligencia artificial" in (act.activity_name or "")
        assert act.codigo_presupuestario == "11.41.67"

    def test_reader_real_format_b_documento_kevds_3(self, reader: ExcelPlanningReader):
        """Lee el archivo institucional real Documento_Kevds_3.xlsx (Formato B: Múltiples actividades)."""
        real_file = "input/matrices_reales_uat/Documento_Kevds_3.xlsx"
        if not os.path.isfile(real_file):
            pytest.skip(f"Archivo real '{real_file}' no disponible en este entorno.")

        result = reader.read_planning_source(real_file)
        assert len(result.structural_errors) == 0
        assert len(result.activities) > 1
        # Verifica que se extraigan actividades con sede El Rama
        assert any(act.sede == "BICU CUR El Rama" for act in result.activities)

    def test_reader_empty_rows_skipped(self, reader: ExcelPlanningReader, tmp_path: Path):
        """Verifica que las filas completamente en blanco sean omitidas sin error."""
        fpath = str(tmp_path / "test_filas_vacias.xlsx")
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Pregrado_Grado en Actividades"
        # Encabezados
        ws.cell(1, 1, "No")
        ws.cell(1, 2, "Sede")
        ws.cell(1, 3, "Actividad")
        # Fila 2: válida
        ws.cell(2, 1, "1")
        ws.cell(2, 2, "Bilwi")
        ws.cell(2, 3, "Actividad 1")
        # Fila 3 y 4: vacías
        # Fila 5: válida
        ws.cell(5, 1, "2")
        ws.cell(5, 2, "Bluefields")
        ws.cell(5, 3, "Actividad 2")
        wb.save(fpath)
        wb.close()

        result = reader.read_planning_source(fpath)
        assert len(result.structural_errors) == 0
        assert len(result.activities) == 2
        assert result.activities[0].activity_name == "Actividad 1"
        assert result.activities[1].activity_name == "Actividad 2"

    def test_reader_invalid_numeric_generates_warning(self, reader: ExcelPlanningReader, tmp_path: Path):
        """Verifica que un valor no numérico en columna cuantitativa genere advertencia."""
        fpath = str(tmp_path / "test_numerico_invalido.xlsx")
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Programas, proyectos y act"
        ws.cell(1, 1, "Sede")
        ws.cell(1, 2, "Actividad")
        ws.cell(1, 3, "Total, estud_M_grado")
        ws.cell(2, 1, "Bilwi")
        ws.cell(2, 2, "Taller de Robótica")
        ws.cell(2, 3, "CINCO")  # Inválido como entero
        wb.save(fpath)
        wb.close()

        result = reader.read_planning_source(fpath)
        assert len(result.structural_errors) == 0
        assert len(result.activities) == 1
        assert result.activities[0].est_grado_m is None
        assert any("no es un entero válido" in w for w in result.reading_warnings)

    def test_reader_date_formatting(self, reader: ExcelPlanningReader, tmp_path: Path):
        """Verifica que celdas con fecha se formateen como cadena ISO YYYY-MM-DD."""
        fpath = str(tmp_path / "test_fechas.xlsx")
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Programas, proyectos y act"
        ws.cell(1, 1, "Sede")
        ws.cell(1, 2, "Actividad")
        ws.cell(1, 3, "Fecha_evento")
        ws.cell(2, 1, "Bilwi")
        ws.cell(2, 2, "Feria Vocacional")
        ws.cell(2, 3, date(2026, 10, 15))
        wb.save(fpath)
        wb.close()

        result = reader.read_planning_source(fpath)
        assert len(result.structural_errors) == 0
        assert result.activities[0].fecha_evento == "2026-10-15"


# ---------------------------------------------------------------------------
# 2. PRUEBAS DEL SERVICIO DE INGESTIÓN (PlannedActivityIngestionService)
# ---------------------------------------------------------------------------
class TestPlannedActivityIngestionService:
    """Pruebas de la capa de aplicación: ingestión, validación institucional e idempotencia."""

    def test_ingestion_service_creates_planned_activity(
        self,
        reader: ExcelPlanningReader,
        uow: PlanningUnitOfWork,
        tmp_path: Path,
    ):
        """Verifica la creación y persistencia exitosa de una actividad en SQLite V003."""
        fpath = str(tmp_path / "test_ingesta_valida.xlsx")
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Programas, proyectos y act"
        ws.cell(1, 1, "No")
        ws.cell(1, 2, "Sede")
        ws.cell(1, 3, "Actividad")
        ws.cell(1, 4, "Área_responsable")
        ws.cell(1, 5, "Eje_estrategia")
        ws.cell(1, 6, "Programa")
        ws.cell(1, 7, "Tipo_Evento")
        ws.cell(1, 8, "Total, estud_M_grado")
        ws.cell(1, 9, "Total, estud_F_grado")

        ws.cell(2, 1, "101")
        ws.cell(2, 2, "Bilwi")
        ws.cell(2, 3, "Taller de Inteligencia Artificial")
        ws.cell(2, 4, "Coordinación de Innovación")
        ws.cell(2, 5, "EJE_11")
        ws.cell(2, 6, "PGM_07")
        ws.cell(2, 7, "EVT_CAPACITACION")
        ws.cell(2, 8, 15)
        ws.cell(2, 9, 20)
        wb.save(fpath)
        wb.close()

        service = PlannedActivityIngestionService(reader=reader, uow=uow)
        report = service.ingest_planning_source(fpath)

        assert report.total_rows_examined == 1
        assert report.rows_accepted == 1
        assert report.rows_rejected == 0
        assert len(report.created_activity_ids) == 1
        expected_pid = "POA-Bilwi-101"
        assert report.created_activity_ids[0] == expected_pid

        # Verificar persistencia en base de datos real
        with uow:
            persisted = uow.planned_activities.get_by_planning_id(expected_pid)
            assert persisted is not None
            assert persisted.activity_name == "Taller de Inteligencia Artificial"
            assert persisted.sede == "Bilwi"
            assert persisted.participant_goals.est_grado_m == 15
            assert persisted.participant_goals.est_grado_f == 20
            assert persisted.participant_goals.total() == 35

    def test_ingestion_service_mandatory_fields_validation(
        self,
        reader: ExcelPlanningReader,
        uow: PlanningUnitOfWork,
        tmp_path: Path,
    ):
        """Verifica que filas con campos obligatorios ausentes sean rechazadas sin fabricar datos."""
        fpath = str(tmp_path / "test_ingesta_incompleta.xlsx")
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Programas, proyectos y act"
        ws.cell(1, 1, "Sede")
        ws.cell(1, 2, "Actividad")
        ws.cell(1, 3, "Área_responsable")

        # Fila 2: falta programa, eje y evento
        ws.cell(2, 1, "Bilwi")
        ws.cell(2, 2, "Actividad Incompleta")
        ws.cell(2, 3, "Coordinación de Ciencias")
        wb.save(fpath)
        wb.close()

        service = PlannedActivityIngestionService(reader=reader, uow=uow)
        report = service.ingest_planning_source(fpath)

        assert report.total_rows_examined == 1
        assert report.rows_accepted == 0
        assert report.rows_rejected == 1
        assert len(report.rejection_reasons) == 1
        assert "Faltan campos" in report.rejection_reasons[0][1]

        # Verificar que no se insertó nada en la BD
        with uow:
            assert len(uow.planned_activities.list_all()) == 0

    def test_ingestion_service_with_defaults(
        self,
        reader: ExcelPlanningReader,
        uow: PlanningUnitOfWork,
        tmp_path: Path,
    ):
        """Verifica que los defaults institucionales se apliquen cuando la celda es nula en el Excel."""
        fpath = str(tmp_path / "test_defaults.xlsx")
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Programas, proyectos y act"
        ws.cell(1, 1, "No")
        ws.cell(1, 2, "Sede")
        ws.cell(1, 3, "Actividad")
        ws.cell(2, 1, "50")
        ws.cell(2, 2, "Bluefields")
        ws.cell(2, 3, "Jornada de Reforestación")
        wb.save(fpath)
        wb.close()

        service = PlannedActivityIngestionService(reader=reader, uow=uow)
        report = service.ingest_planning_source(
            file_path=fpath,
            default_area_responsable="Dirección de Ambiente",
            default_eje_estrategia="EJE_7",
            default_programa="PGM_05",
            default_tipo_evento="EVT_CAMPANA",
        )

        assert report.rows_accepted == 1
        with uow:
            act = uow.planned_activities.get_by_planning_id("POA-Bluefields-50")
            assert act is not None
            assert act.area_responsable == "Dirección de Ambiente"
            assert act.eje_estrategia == "EJE_7"
            assert act.programa == "PGM_05"
            assert act.tipo_evento == "EVT_CAMPANA"

    def test_ingestion_service_negative_goals_rejected(
        self,
        reader: ExcelPlanningReader,
        uow: PlanningUnitOfWork,
        tmp_path: Path,
    ):
        """Verifica que filas con metas numéricas negativas sean rechazadas."""
        fpath = str(tmp_path / "test_metas_negativas.xlsx")
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Programas, proyectos y act"
        ws.cell(1, 1, "Sede")
        ws.cell(1, 2, "Actividad")
        ws.cell(1, 3, "Área_responsable")
        ws.cell(1, 4, "Eje_estrategia")
        ws.cell(1, 5, "Programa")
        ws.cell(1, 6, "Tipo_Evento")
        ws.cell(1, 7, "Total, estud_M_grado")

        ws.cell(2, 1, "Bilwi")
        ws.cell(2, 2, "Actividad Negativa")
        ws.cell(2, 3, "Área X")
        ws.cell(2, 4, "EJE_1")
        ws.cell(2, 5, "PGM_01")
        ws.cell(2, 6, "EVT_CHARLA")
        ws.cell(2, 7, -5)  # Negativo
        wb.save(fpath)
        wb.close()

        service = PlannedActivityIngestionService(reader=reader, uow=uow)
        report = service.ingest_planning_source(fpath)

        assert report.rows_rejected == 1
        assert any("no pueden ser negativas" in r[1] for r in report.rejection_reasons)

    def test_ingestion_service_idempotence_and_update(
        self,
        reader: ExcelPlanningReader,
        uow: PlanningUnitOfWork,
        tmp_path: Path,
    ):
        """Verifica que procesar el mismo archivo dos veces actualice in-place sin duplicar registros."""
        fpath = str(tmp_path / "test_idempotencia.xlsx")
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Programas, proyectos y act"
        ws.cell(1, 1, "No")
        ws.cell(1, 2, "Sede")
        ws.cell(1, 3, "Actividad")
        ws.cell(1, 4, "Área_responsable")
        ws.cell(1, 5, "Eje_estrategia")
        ws.cell(1, 6, "Programa")
        ws.cell(1, 7, "Tipo_Evento")

        ws.cell(2, 1, "77")
        ws.cell(2, 2, "Bilwi")
        ws.cell(2, 3, "Conferencia de Innovación")
        ws.cell(2, 4, "Coordinación Innovación")
        ws.cell(2, 5, "EJE_11")
        ws.cell(2, 6, "PGM_07")
        ws.cell(2, 7, "EVT_CONFERENCIA")
        wb.save(fpath)
        wb.close()

        service = PlannedActivityIngestionService(reader=reader, uow=uow)

        # Primera corrida: crea
        report1 = service.ingest_planning_source(fpath)
        assert report1.rows_accepted == 1
        assert len(report1.created_activity_ids) == 1
        assert len(report1.updated_activity_ids) == 0

        # Modificar archivo: cambiar nombre de actividad
        wb = openpyxl.load_workbook(fpath)
        ws = wb["Programas, proyectos y act"]
        ws.cell(2, 3, "Conferencia de Innovación Actualizada")
        wb.save(fpath)
        wb.close()

        # Segunda corrida: actualiza in-place
        report2 = service.ingest_planning_source(fpath)
        assert report2.rows_accepted == 1
        assert len(report2.created_activity_ids) == 0
        assert len(report2.updated_activity_ids) == 1
        assert len(report2.duplicates_detected) == 1

        with uow:
            all_acts = uow.planned_activities.list_all()
            assert len(all_acts) == 1
            assert all_acts[0].activity_name == "Conferencia de Innovación Actualizada"

    def test_ingestion_service_protects_approved_design(
        self,
        reader: ExcelPlanningReader,
        uow: PlanningUnitOfWork,
        tmp_path: Path,
    ):
        """Verifica que la regla R-08 proteja actividades con diseños en estado APPROVED contra sobreescritura."""
        fpath = str(tmp_path / "test_proteccion_approved.xlsx")
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Programas, proyectos y act"
        ws.cell(1, 1, "No")
        ws.cell(1, 2, "Sede")
        ws.cell(1, 3, "Actividad")
        ws.cell(1, 4, "Área_responsable")
        ws.cell(1, 5, "Eje_estrategia")
        ws.cell(1, 6, "Programa")
        ws.cell(1, 7, "Tipo_Evento")

        ws.cell(2, 1, "99")
        ws.cell(2, 2, "Bilwi")
        ws.cell(2, 3, "Actividad Aprobada Protegida")
        ws.cell(2, 4, "Área Protegida")
        ws.cell(2, 5, "EJE_1")
        ws.cell(2, 6, "PGM_01")
        ws.cell(2, 7, "EVT_TALLER")
        wb.save(fpath)
        wb.close()

        service = PlannedActivityIngestionService(reader=reader, uow=uow)
        report1 = service.ingest_planning_source(fpath)
        assert report1.rows_accepted == 1

        # Crear y aprobar un diseño metodológico sobre esta actividad
        with uow:
            act = uow.planned_activities.get_by_planning_id("POA-Bilwi-99")
            assert act is not None
            design = MethodologicalDesign.create_draft(
                planned_activity_ref=act.planning_id,
                created_by="Responsable UAT",
            )
            design.planned_activity_internal_id = act.activity_internal_id
            design.activity_name = act.activity_name
            design.objectives = ["Obj 1", "Obj 2"]
            design.approve(approved_by="Decano BICU")
            uow.methodological_designs.save(design)
            uow.commit()

        # Intentar modificar la actividad con un segundo archivo
        wb = openpyxl.load_workbook(fpath)
        ws = wb["Programas, proyectos y act"]
        ws.cell(2, 3, "Actividad Que Intenta Cambiar")
        wb.save(fpath)
        wb.close()

        report2 = service.ingest_planning_source(fpath)
        assert report2.rows_accepted == 1
        assert len(report2.updated_activity_ids) == 0  # No debe haberse actualizado
        assert any("inmutable R-08" in d for d in report2.duplicates_detected)

        # Comprobar que el valor original permanece intacto
        with uow:
            act_db = uow.planned_activities.get_by_planning_id("POA-Bilwi-99")
            assert act_db is not None
            assert act_db.activity_name == "Actividad Aprobada Protegida"


# ---------------------------------------------------------------------------
# 3. AUDITORÍA AST DE AISLAMIENTO ARQUITECTÓNICO
# ---------------------------------------------------------------------------
class TestPlanningReaderArchitectureIsolation:
    """Auditorías periciales de AST para garantizar la estricta separación de responsabilidades."""

    def test_reader_infrastructure_ast_isolation(self):
        """Verifica que el lector Excel no importe sqlite3, app.word_consolidator ni IA."""
        reader_file = Path("app/planning/infrastructure/readers/excel_planning_reader.py")
        assert reader_file.is_file()

        tree = ast.parse(reader_file.read_text(encoding="utf-8"))
        forbidden_modules = {
            "sqlite3",
            "app.word_consolidator",
            "app.consolidation",
            "app.infrastructure.persistence",
            "openai",
            "google.generativeai",
        }

        for node in ast.walk(tree):
            if isinstance(node, ast.Import):
                for alias in node.names:
                    for forbidden in forbidden_modules:
                        assert not alias.name.startswith(forbidden), (
                            f"Violación de aislamiento en {reader_file}: importa '{alias.name}'"
                        )
            elif isinstance(node, ast.ImportFrom):
                mod = node.module or ""
                for forbidden in forbidden_modules:
                    assert not mod.startswith(forbidden), (
                        f"Violación de aislamiento en {reader_file}: importa de '{mod}'"
                    )

    def test_ingestion_service_ast_isolation(self):
        """Verifica que el servicio de aplicación no importe openpyxl, sqlite3 ni app.word_consolidator."""
        services_file = Path("app/planning/application/services.py")
        assert services_file.is_file()

        tree = ast.parse(services_file.read_text(encoding="utf-8"))
        forbidden_modules = {
            "openpyxl",
            "sqlite3",
            "app.word_consolidator",
            "app.consolidation",
        }

        for node in ast.walk(tree):
            if isinstance(node, ast.Import):
                for alias in node.names:
                    for forbidden in forbidden_modules:
                        assert not alias.name.startswith(forbidden), (
                            f"Violación de aislamiento en {services_file}: importa '{alias.name}'"
                        )
            elif isinstance(node, ast.ImportFrom):
                mod = node.module or ""
                for forbidden in forbidden_modules:
                    assert not mod.startswith(forbidden), (
                        f"Violación de aislamiento en {services_file}: importa de '{mod}'"
                    )


# ---------------------------------------------------------------------------
# 4. PRUEBA DE NO CONTAMINACIÓN HACIA EJECUCIÓN (PLANIFICADO ≠ EJECUTADO)
# ---------------------------------------------------------------------------
class TestPlanningReaderNoContamination:
    """Verifica que la ingestión no contamine ni altere el subsistema de ejecución real."""

    def test_zero_contamination_execution(self, reader: ExcelPlanningReader, uow: PlanningUnitOfWork, tmp_path: Path):
        """Comprueba que la ingestión de actividades en V003 no cree registros en tablas de ejecución."""
        fpath = str(tmp_path / "test_no_contaminacion.xlsx")
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Programas, proyectos y act"
        ws.cell(1, 1, "No")
        ws.cell(1, 2, "Sede")
        ws.cell(1, 3, "Actividad")
        ws.cell(2, 1, "1")
        ws.cell(2, 2, "Bilwi")
        ws.cell(2, 3, "Actividad de Prueba No Contaminación")
        wb.save(fpath)
        wb.close()

        service = PlannedActivityIngestionService(reader=reader, uow=uow)
        report = service.ingest_planning_source(
            file_path=fpath,
            default_area_responsable="Área Test",
            default_eje_estrategia="EJE_1",
            default_programa="PGM_01",
            default_tipo_evento="EVT_TALLER",
        )
        assert report.rows_accepted == 1

        # Verificar que en la base de datos de test SOLO existen tablas planning_*
        with uow:
            cursor = uow.connection.cursor()
            cursor.execute("SELECT name FROM sqlite_master WHERE type='table' AND name NOT LIKE 'sqlite_%';")
            tables = [r[0] for r in cursor.fetchall()]
            for t in tables:
                assert t.startswith("planning_"), f"Tabla inesperada detectada: {t}"
