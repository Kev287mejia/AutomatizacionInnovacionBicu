"""
tests/template_fixture_factory.py

Generador de Golden Test Fixtures exclusivo para pruebas unitarias de Fase 8.
ADVERTENCIA: Estos archivos son artefactos sintéticos de prueba ("NO OFICIALES").
Bajo ninguna circunstancia sustituyen las plantillas oficiales institucionales.
"""

from pathlib import Path
import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation


class TemplateFixtureFactory:
    """
    Genera libros de Excel con características complejas (fórmulas, combinaciones,
    validaciones de datos, hojas ocultas) para validar exhaustivamente el TemplateInspector.
    """

    @classmethod
    def crear_fixtures_de_prueba(cls, carpeta_destino: Path) -> Path:
        """
        Genera los 5 fixtures sintéticos de prueba en carpeta_destino.
        """
        carpeta_destino.mkdir(parents=True, exist_ok=True)

        # Definiciones estructurales de prueba
        cls._crear_fixture_matriz_1(carpeta_destino / "fixture_test_matriz_1.xlsx")
        cls._crear_fixture_matriz_2(carpeta_destino / "fixture_test_matriz_2.xlsx")
        cls._crear_fixture_matriz_3(carpeta_destino / "fixture_test_matriz_3.xlsx")
        cls._crear_fixture_matriz_4(carpeta_destino / "fixture_test_matriz_4.xlsx")
        cls._crear_fixture_matriz_5(carpeta_destino / "fixture_test_matriz_5.xlsx")

        return carpeta_destino

    @classmethod
    def _crear_fixture_matriz_1(cls, ruta: Path) -> None:
        """Matriz 1: Consolidado de Actividades (Prueba sintética)."""
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Consolidado"

        # Banner institucional simulado con celdas combinadas
        ws.merge_cells("A1:Q2")
        ws["A1"] = "GOLDEN TEST FIXTURE (NO OFICIAL) - MATRIZ 1 CONSOLIDADO"
        ws["A1"].font = Font(bold=True, size=14, color="FFFFFF")
        ws["A1"].fill = PatternFill(start_color="1F497D", end_color="1F497D", fill_type="solid")
        ws["A1"].alignment = Alignment(horizontal="center", vertical="center")

        ws["A3"] = "Sede Bilwi / Período 2026"

        # Fila 4: Encabezados
        encabezados = [
            "No.", "Nombre de la Actividad", "Sede", "Municipio", "Departamento", "Fecha",
            "Eje", "Tipo de Evento", "Total Asistencias", "Personas Unicas", "Femenino",
            "Masculino", "Estudiantes", "Docentes", "Administrativos", "Beneficiados", "Colaboradores"
        ]
        for col_idx, enc in enumerate(encabezados, 1):
            cell = ws.cell(row=4, column=col_idx, value=enc)
            cell.font = Font(bold=True)

        # Fila 10: Fórmulas de Totales al pie
        ws.cell(row=10, column=2, value="TOTALES GENERALES:")
        ws.cell(row=10, column=9, value="=SUM(I5:I9)")
        ws.cell(row=10, column=10, value="=SUM(J5:J9)")
        ws.cell(row=10, column=11, value="=SUM(K5:K9)")
        ws.cell(row=10, column=12, value="=SUM(L5:L9)")

        # Hoja oculta de configuración
        ws_oculta = wb.create_sheet(title="Config_Oculta")
        ws_oculta["A1"] = "Metadatos técnicos de prueba"
        ws_oculta.sheet_state = "hidden"

        wb.save(str(ruta))
        wb.close()

    @classmethod
    def _crear_fixture_matriz_2(cls, ruta: Path) -> None:
        """Matriz 2: Estudiantes (Prueba sintética)."""
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Estudiantes"

        ws.merge_cells("A1:K2")
        ws["A1"] = "GOLDEN TEST FIXTURE (NO OFICIAL) - MATRIZ 2 ESTUDIANTES"
        ws["A1"].font = Font(bold=True, size=12)

        # Fila 4: Encabezados (incluyendo una columna sin fuente para probar SIN_FUENTE)
        encabezados = [
            "No.", "Cedula", "Nombres y Apellidos", "Sexo", "Edad", "Carrera",
            "Sede", "Nombre de la Actividad", "Fecha", "Observaciones", "Código Presupuestario Externo"
        ]
        for col_idx, enc in enumerate(encabezados, 1):
            ws.cell(row=4, column=col_idx, value=enc).font = Font(bold=True)

        # Validación de datos en columna Sexo (columna D, filas 5 a 50)
        dv = DataValidation(type="list", formula1='"FEMENINO,MASCULINO"', allow_blank=True)
        ws.add_data_validation(dv)
        dv.add("D5:D50")

        wb.save(str(ruta))
        wb.close()

    @classmethod
    def _crear_fixture_matriz_3(cls, ruta: Path) -> None:
        """Matriz 3: Académicos y Administrativos (Prueba sintética)."""
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Academicos_Administrativos"

        ws.merge_cells("A1:I2")
        ws["A1"] = "GOLDEN TEST FIXTURE (NO OFICIAL) - MATRIZ 3 ACADÉMICOS Y ADMINISTRATIVOS"

        encabezados = [
            "No.", "Cedula", "Nombres y Apellidos", "Sexo", "Edad", "Rol",
            "Area", "Actividad", "Fecha"
        ]
        for col_idx, enc in enumerate(encabezados, 1):
            ws.cell(row=4, column=col_idx, value=enc).font = Font(bold=True)

        wb.save(str(ruta))
        wb.close()

    @classmethod
    def _crear_fixture_matriz_4(cls, ruta: Path) -> None:
        """Matriz 4: Colaboradores (Prueba sintética)."""
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Colaboradores"

        ws.merge_cells("A1:H2")
        ws["A1"] = "GOLDEN TEST FIXTURE (NO OFICIAL) - MATRIZ 4 COLABORADORES"

        encabezados = [
            "No.", "Cedula", "Nombres y Apellidos", "Sexo", "Edad", "Institucion",
            "Actividad", "Fecha"
        ]
        for col_idx, enc in enumerate(encabezados, 1):
            ws.cell(row=4, column=col_idx, value=enc).font = Font(bold=True)

        wb.save(str(ruta))
        wb.close()

    @classmethod
    def _crear_fixture_matriz_5(cls, ruta: Path) -> None:
        """Matriz 5: Protagonistas Beneficiados (Prueba sintética)."""
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Beneficiados"

        ws.merge_cells("A1:H2")
        ws["A1"] = "GOLDEN TEST FIXTURE (NO OFICIAL) - MATRIZ 5 BENEFICIADOS"

        encabezados = [
            "No.", "Cedula", "Nombres y Apellidos", "Sexo", "Edad", "Comunidad",
            "Actividad", "Fecha"
        ]
        for col_idx, enc in enumerate(encabezados, 1):
            ws.cell(row=4, column=col_idx, value=enc).font = Font(bold=True)

        wb.save(str(ruta))
        wb.close()
